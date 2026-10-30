"""
The shared part of the builders of the parts files: what every robot type taken apart is written with.

**Copyright notice:** Copyright © 2026 Empa, Matthias Roesslein
**Licence:** CC BY 4.0, Creative Commons Attribution 4.0 International; the full text is in `LICENSE`
at the root of the repository.

Matthias, 2026-10-30: "Yes, build C" (documentation/claude/HANDOVER.md, C: "a parts file and builder per type, 02
drawing each type the same way"; Q1: each type in a folder of its own, the shared machinery here; "One copy of the
code").

WHAT IT HOLDS
-------------
The one copy of what the builders of every robot type share: the sheets' first rows and column widths; the sheets
that are one set for all robots (MagnetGrades, the magnets' grade classes by heat; JointParts, the G1's joint
weighed part by part, which every joint is scaled from; Shares, the materials of the parts and components that hold
more than one); the shared numbers of the sheet Constants and the sources they cite; the template's sizes and
masses where a source names a board or sensor but gives none (EVIDENCE.md E5.14); and the functions that write a
robot's block into its rows and the rows into a file. Each robot type's builder, tools/<n>_<name>/
make_parts_taken_apart.py, holds its own robots and calls these.

WHAT IT READS AND WRITES
------------------------
It reads nothing but 00's settings (src/params_schema.py, for where a file goes). It writes a file only when a
type's builder calls build(): `data.parts_taken_apart` with that type's folder (data/researched/<n>_<name>/
parts_taken_apart.xlsx, git-ignored).

WHERE IT SITS
-------------
Not a stage and not run by itself: the builders in tools/<n>_<name>/ import it. Importing it writes nothing.
"""
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# The project's root, so that its settings (src/params_schema.py) can be imported when this file is run from tools/.
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)
from src.params_schema import Params, absolute  # noqa: E402

# The grey of every sheet's first row, so that the headings stand apart from the values.
HEADING_FILL = PatternFill('solid', fgColor='FFEEEEEE')
# The width of the About sheet's one column, in characters.
ABOUT_WIDTH = 150


# --- The sheets' first rows and column widths --------------------------------------------------------------------
# The first row of every sheet and its columns' widths, the same in every robot type's file (01_schemas.py reads
# them by these headings, APART_SHEETS).
ROBOTS_HEADINGS = ['Robot', 'Maker', 'Type', 'Height (m)', 'Stated mass (kg)', 'Sources']
ROBOTS_WIDTHS = [14, 10, 6, 10, 14, 60]
JOINTS_HEADINGS = ['Robot', 'Number', 'Joint', 'Joints in the robot', 'Module', 'Torque limit (N·m)',
                   'Module mass, min (kg)', 'mode (kg)', 'max (kg)', 'Share of the joint modules', 'Cooling',
                   'Heat group', 'Body region', 'Sources', 'Basis']
JOINTS_WIDTHS = [12, 8, 22, 10, 9, 10, 10, 9, 9, 12, 30, 18, 11, 50, 70]
PARTS_HEADINGS = ['Robot', 'Part', 'Number in the robot', "Kind (the figure's row)", 'Body region',
                  'Mass of one, min (kg)', 'mode (kg)', 'max (kg)', 'Material (one material only)', 'Grade',
                  'Remainder of, or computed from', 'Label', 'Sources']
PARTS_WIDTHS = [12, 44, 10, 20, 16, 11, 9, 9, 18, 14, 52, 7, 56]
BOARDS_HEADINGS = ['Robot', 'Board', 'Number in the robot', "Kind (the figure's row)", 'Body region',
                   'Area of one, min (cm2)', 'mode (cm2)', 'max (cm2)', 'Mass from',
                   'Car board category (RAWCLICVehicleElectronics 04_)', 'Label', 'Sources']
BOARDS_WIDTHS = [12, 50, 10, 20, 16, 11, 9, 9, 52, 18, 7, 46]
SENSORS_HEADINGS = ['Robot', 'Sensor', 'Number in the robot', "Kind (the figure's row)", 'Body region',
                    'Mass of one, min (g)', 'mode (g)', 'max (g)', 'Label', 'Sources']
SENSORS_WIDTHS = [12, 40, 10, 20, 16, 11, 9, 9, 7, 52]
SENSORELEMENTS_HEADINGS = ['Sensor', 'Car sensor type (RAWCLICVehicleElectronics 07_, sheet Sensor Details)',
                           'Units in one sensor']
SENSORELEMENTS_WIDTHS = [40, 50, 12]
CONSTANTS_HEADINGS = ['Robot', 'Quantity', 'Min', 'Mode', 'Max', 'Unit', 'Label', 'Source']
CONSTANTS_WIDTHS = [12, 48, 9, 9, 9, 8, 7, 70]
SOURCES_HEADINGS = ['Code', 'Source', 'Link', 'What it is']
SOURCES_WIDTHS = [12, 70, 90, 50]
MAGNETGRADES_WIDTHS = [18, 12, 12, 7, 90]
JOINTPARTS_WIDTHS = [44, 12, 12, 9, 9, 16, 12, 10, 9, 9, 60, 14, 60]
SHARES_WIDTHS = [52, 18, 22, 10, 10, 10, 10, 7, 90]

# --- MagnetGrades ---------------------------------------------------------------------------------------------------
# The magnets' grade classes by heat group: each group a distribution over grade classes, drawn in each draw, tied
# within a draw (Matthias, 2026-10-28: "Heat has to matter and so the distribution will be different depending on the
# joint", "Accept as proposed"). Every joint judged as passively cooled ("Fans is not a sustainable solution"). No
# maker states the grade of a joint motor: the probabilities are estimates on the order of heat the sources give
# (EVIDENCE.md E2.11).
MAGNETGRADES_HEADINGS = ['Heat group', 'Grade class', 'Probability', 'Label', 'Basis']
MAGNETGRADES_ROWS = [
    ('Knees', 'SH', 0.3, 'E',
     'E2.11: the largest load, passive, derating to protect the magnets (OFweek; T1 slide 42; T3)'),
    ('Knees', 'UH', 0.5, 'E',
     'E2.11: the largest load, passive, derating to protect the magnets (OFweek; T1 slide 42; T3)'),
    ('Knees', 'EH', 0.2, 'E',
     'E2.11: the largest load, passive, derating to protect the magnets (OFweek; T1 slide 42; T3)'),
    ('Hips', 'SH', 0.5, 'E', 'E2.11: carry the body; passive as judged (T3, OFweek)'),
    ('Hips', 'UH', 0.4, 'E', 'E2.11: carry the body; passive as judged (T3, OFweek)'),
    ('Hips', 'EH', 0.1, 'E', 'E2.11: carry the body; passive as judged (T3, OFweek)'),
    ('Waist', 'SH', 0.6, 'E', 'E2.11: moves often (OFweek); passive as judged'),
    ('Waist', 'UH', 0.4, 'E', 'E2.11: moves often (OFweek); passive as judged'),
    ('Ankles', 'SH', 0.6, 'E', 'E2.11: carry the robot; small joints (T1)'),
    ('Ankles', 'UH', 0.4, 'E', 'E2.11: carry the robot; small joints (T1)'),
    ('Shoulders', 'H', 0.5, 'E', 'E2.11: hold the arm; low load'),
    ('Shoulders', 'SH', 0.5, 'E', 'E2.11: hold the arm; low load'),
    ('Elbows and wrists', 'H', 0.7, 'E', 'E2.11: the lowest load'),
    ('Elbows and wrists', 'SH', 0.3, 'E', 'E2.11: the lowest load'),
]

# --- JointParts -----------------------------------------------------------------------------------------------------
# The G1's small joint part by part as China Post Securities weighed it (slide 43), the small joints not
# weighed at 0.5 / 1 / 1.5 times, and the large joint as the small one times a factor (EVIDENCE.md E2.8).
# One set for all robots, scaled to each joint's module mass ("similar within given ranges"). The column Grade:
# the steel grade whose element recipe a steel component follows, each by its use (EVIDENCE.md E5.18, agreed
# 2026-10-29): the gears 20CrMnTi, the small bearings GCr15, the output flange 42CrMo4; a component of shares has its
# grades in the sheet Shares.
JOINTPARTS_HEADINGS = ['Component', 'Small joint, weighed (g)', 'Small joints not weighed: times min', 'times mode',
                       'times max', 'Material (one material only)', 'Grade', 'Large joint: factor min',
                       'factor mode', 'factor max', 'Why the factor', 'Label', 'Sources']
JOINTPARTS_ROWS = [
    ('Motor stator', 94.3, 0.5, 1, 1.5, None, None, 1, 1.78, 2.67, "the motor's cross-section, 80 against 60 mm across",
     'W; E for the spread and the factor',
     'T1 slide 43 (figure 37); factors EVIDENCE.md E2.8; Matthias, 2026-10-27, step 1: '
     '"Accept all" (EVIDENCE.md E2.8)'),
    ('Motor rotor', 42.9, 0.5, 1, 1.5, None, None, 1, 1.78, 2.67, "the motor's cross-section",
     'W; E for the spread and the factor',
     'T1 slide 43 (figure 37); factors EVIDENCE.md E2.8; Matthias, 2026-10-27, step 1: '
     '"Accept all" (EVIDENCE.md E2.8)'),
    ('Magnet carrier', 16.9, 0.5, 1, 1.5, None, None, 1, 1.78, 2.67, "the motor's cross-section",
     'W; E for the spread and the factor',
     'T1 slide 43 (figure 37); factors EVIDENCE.md E2.8; Matthias, 2026-10-27, step 1: '
     '"Accept all" (EVIDENCE.md E2.8)'),
    ('Gear stage 1, sun gear', 8, 0.5, 1, 1.5, 'Steel', '20CrMnTi', 1.17, 2.33, 3.56,
     'from the cross-section alone to growing with the torque (99.3 against 27.9 N·m); mode '
     '3.56^(2/3)', 'W; E for the spread and the factor',
     'T1 slide 43 (figure 37); factors EVIDENCE.md E2.8; Matthias, 2026-10-27, step 1: '
     '"Accept all" (EVIDENCE.md E2.8)'),
    ('Gear stage 1, planets and carrier with stage 2 sun gear', 25.3, 0.5, 1, 1.5, 'Steel', '20CrMnTi', 1.17, 2.33,
     3.56, 'as the sun gear', 'W; E for the spread and the factor',
     'T1 slide 43 (figure 37); factors EVIDENCE.md E2.8; Matthias, 2026-10-27, step 1: '
     '"Accept all" (EVIDENCE.md E2.8)'),
    ('Gear stage 1, ring gear', 24.1, 0.5, 1, 1.5, 'Steel', '20CrMnTi', 1.17, 2.33, 3.56, 'as the sun gear',
     'W; E for the spread and the factor',
     'T1 slide 43 (figure 37); factors EVIDENCE.md E2.8; Matthias, 2026-10-27, step 1: '
     '"Accept all" (EVIDENCE.md E2.8)'),
    ('Gear stage 2, planets, carrier and crossed roller bearing', 74, 0.5, 1, 1.5, None, None, 1.17, 2.33, 3.56,
     'as the sun gear', 'W; E for the spread and the factor',
     'T1 slide 43 (figure 37); factors EVIDENCE.md E2.8; Matthias, 2026-10-27, step 1: '
     '"Accept all" (EVIDENCE.md E2.8)'),
    ('Gear stage 2, ring gear', 23.8, 0.5, 1, 1.5, 'Steel', '20CrMnTi', 1.17, 2.33, 3.56, 'as the sun gear',
     'W; E for the spread and the factor',
     'T1 slide 43 (figure 37); factors EVIDENCE.md E2.8; Matthias, 2026-10-27, step 1: '
     '"Accept all" (EVIDENCE.md E2.8)'),
    ('Driver board', 12.8, 0.5, 1, 1.5, 'Electronics', None, 1.78, 3.56, 5.33,
     'one, two or three boards of the larger size (T3k: "a stack")', 'W; E for the spread and the factor',
     'T1 slide 43 (figure 37); factors EVIDENCE.md E2.8; Matthias, 2026-10-27, step 1: '
     '"Accept all" (EVIDENCE.md E2.8)'),
    ('Two small bearings', 11, 0.5, 1, 1.5, 'Steel', 'GCr15', 1, 1.78, 2.67, 'the area',
     'W; E for the spread and the factor',
     'T1 slide 43 (figure 37); factors EVIDENCE.md E2.8; Matthias, 2026-10-27, step 1: '
     '"Accept all" (EVIDENCE.md E2.8)'),
    ('Housing', 79.3, 0.5, 1, 1.5, 'Aluminium', None, 1, 1.39, 2.09, "the joint's surface",
     'W; E for the spread and the factor',
     'T1 slide 43 (figure 37); factors EVIDENCE.md E2.8; Matthias, 2026-10-27, step 1: '
     '"Accept all" (EVIDENCE.md E2.8)'),
    ('Two end caps', 23.7, 0.5, 1, 1.5, 'Aluminium', None, 1, 1.78, 2.67, 'the area',
     'W; E for the spread and the factor',
     'T1 slide 43 (figure 37); factors EVIDENCE.md E2.8; Matthias, 2026-10-27, step 1: '
     '"Accept all" (EVIDENCE.md E2.8)'),
    ('Output flange', 50.8, 0.5, 1, 1.5, 'Steel', '42CrMo4', 1, 1.78, 2.67, 'the area',
     'W; E for the spread and the factor',
     'T1 slide 43 (figure 37); factors EVIDENCE.md E2.8; Matthias, 2026-10-27, step 1: '
     '"Accept all" (EVIDENCE.md E2.8)'),
    ('Other', 38.1, 0.5, 1, 1.5, None, None, 1, 1.78, 2.67, 'the area', 'W; E for the spread and the factor',
     'T1 slide 43 (figure 37); factors EVIDENCE.md E2.8; Matthias, 2026-10-27, step 1: '
     '"Accept all" (EVIDENCE.md E2.8)'),
]

# --- Shares ---------------------------------------------------------------------------------------------------------
# The materials of the parts and components that hold more than one; each group has one named remainder and
# nothing is rescaled (no normalisation). EVIDENCE.md E2.8, E2.10. The column Grade: the grade whose element recipe the
# material follows, as for a part (E5.18, E5.19): 'A or B' is either, each as often (a discrete state, drawn); empty
# for aluminium and plastics (type 1's recipes) and for magnets (their heat group). What the share is of the part
# opens its Basis. A material written 'A or B' is either material, each as often (the battery housing, E5.19 point 1).
SHARES_HEADINGS = ['Part or component', 'Material', 'Grade', 'Share min', 'Share mode', 'Share max', 'Remainder',
                   'Label', 'Basis']
SHARES_ROWS = [
    ('Motor stator', 'Copper', None, 0.23, 0.36, 0.45, None, 'S',
     'The winding: measured car traction motors 23-29 % (Drexler et al. 2025), car auxiliary motors 27 / 36 / 45 %; '
     'E2.10; Matthias: "Yes use this numberrs"'),
    ('Motor stator', 'Plastics', None, 0.02, 0.05, 0.1, None, 'E', "The winding's coating: E2.8"),
    ('Motor stator', 'Steel', 'M19 or M27', None, None, None, 'Yes', 'E',
     'The electrical steel: Matthias, 2026-10-27, step 1: "Accept all" (EVIDENCE.md E2.8); its grade E5.18'),
    ('Motor rotor', 'Magnets (NdFeB)', None, 0.45, 0.9, 1, None, 'E',
     'T1 calls it "磁钢"; Matthias, 2026-10-27, step 1: "Accept all" (EVIDENCE.md E2.8)'),
    ('Motor rotor', 'Steel', 'M19 or M27', None, None, None, 'Yes', 'E',
     'Matthias, 2026-10-27, step 1: "Accept all" (EVIDENCE.md E2.8); its grade E5.18'),
    ('Magnet carrier', 'Plastics', None, 0.05, 0.15, 0.3, None, 'E', 'The plastic bearing: T1 "铝+塑料轴承"'),
    ('Magnet carrier', 'Aluminium', None, None, None, None, 'Yes', 'E',
     'Matthias, 2026-10-27, step 1: "Accept all" (EVIDENCE.md E2.8)'),
    ('Gear stage 2, planets, carrier and crossed roller bearing', 'Aluminium', None, 0.125, 0.25, 0.4, None, 'E',
     'The carrier: T1 "保持架为铝，其余为钢"'),
    ('Gear stage 2, planets, carrier and crossed roller bearing', 'Steel', 'Gear stage 2 steel: 20CrMnTi and GCr15',
     None, None, None, 'Yes', 'E',
     'The planets and the bearing: Matthias, 2026-10-27, step 1: "Accept all" (EVIDENCE.md E2.8); 20CrMnTi 30 / 50 / '
     '70 % of the steel, GCr15 the remainder (E5.18)'),
    ('Other', 'Copper', None, 0.1, 0.25, 0.4, None, 'E', 'Wires: no basis'),
    ('Other', 'Steel', '42CrMo4', 0.2, 0.4, 0.6, None, 'E',
     'Screws, encoder parts: no basis; the grade of class 12.9 screws (E5.18)'),
    ('Other', 'Plastics', None, None, None, None, 'Yes', 'E',
     'Matthias, 2026-10-27, step 1: "Accept all" (EVIDENCE.md E2.8)'),
    ('Hip fan', 'Copper', None, 0.05, 0.1, 0.2, None, 'E', 'The motor winding: E2.8'),
    ('Hip fan', 'Plastics', None, None, None, None, 'Yes', 'E', 'E2.8'),
    ('Waist fan', 'Copper', None, 0.05, 0.1, 0.2, None, 'E', 'The motor winding: E2.8'),
    ('Waist fan', 'Plastics', None, None, None, None, 'Yes', 'E', 'E2.8'),
    ('Main board fan', 'Copper', None, 0.05, 0.1, 0.2, None, 'E', 'The motor winding: E2.8'),
    ('Main board fan', 'Plastics', None, None, None, None, 'Yes', 'E', 'E2.8'),
    ('XT30 (2+2) connectors (two per joint)', 'Copper', 'Brass', 0.175, 0.35, 0.525, None, 'E',
     'The contacts: E2.8; brass, gold-plated (Amass, through LCSC; E5.22)'),
    ('XT30 (2+2) connectors (two per joint)', 'Plastics', None, None, None, None, 'Yes', 'E', 'The housing: E2.8'),
    ('XT60 connector, mated pair', 'Copper', 'Brass', 0.175, 0.35, 0.525, None, 'E',
     "The contacts, brass, as the XT30's (E5.22)"),
    ('XT60 connector, mated pair', 'Plastics', None, None, None, None, 'Yes', 'E', "The housing, as the XT30's"),
    ('Battery housing, busbars and wires', 'Copper', 'Nickel strip', 0.01, 0.03, 0.06, None, 'E',
     'The nickel strips between the cells, 0.2 mm thick (sevarg.net, measured), counted with the copper (Matthias, '
     '2026-10-29: "With Copper"); E5.19 point 1'),
    ('Battery housing, busbars and wires', 'Copper', None, 0.05, 0.1, 0.15, None, 'E',
     'The copper wires and busbars: E5.19 point 1'),
    ('Battery housing, busbars and wires', 'Plastics', None, 0.05, 0.1, 0.2, None, 'E',
     'The cell holders, insulation and tape: E5.19 point 1'),
    ('Battery housing, busbars and wires', 'Aluminium or Plastics', None, None, None, None, 'Yes', 'E',
     'The housing: "aluminum or high-impact plastics" (teswaybike.com, a claim), each as often (E5.19 point 1)'),
    ('Battery housing (aluminium), busbars and wires', 'Copper', 'Nickel strip', 0.01, 0.03, 0.06, None, 'E',
     'The nickel strips between the cells, counted with the copper (Matthias, 2026-10-29: "With Copper"); E5.19 '
     'point 1'),
    ('Battery housing (aluminium), busbars and wires', 'Copper', None, 0.05, 0.1, 0.15, None, 'E',
     'The copper wires and busbars: E5.19 point 1'),
    ('Battery housing (aluminium), busbars and wires', 'Plastics', None, 0.05, 0.1, 0.2, None, 'E',
     'The cell holders, insulation and tape: E5.19 point 1'),
    ('Battery housing (aluminium), busbars and wires', 'Aluminium', None, None, None, None, 'Yes', 'E',
     'The housing: "Aluminum Alloy" (EngineAI, the PM01; E5.6, E5.19 point 1)'),
    ('Frame: aluminium, with steel 45 and P20 parts', 'Steel', 'RP1 frame steel: C45 and P20', 0.1, 0.2, 0.3, None,
     'E', 'The steel parts: 28 of its 112 CNC parts are steel, 24 of steel 45 or AISI 1045 and 4 of P20 (RP1a), by '
     'count only (E5.14 point 9, E5.18)'),
    ('Frame: aluminium, with steel 45 and P20 parts', 'Aluminium', None, None, None, None, 'Yes', 'E',
     'The aluminium parts: AL6061-T6, AL7075, AL5052 (RP1a); E5.14 point 9'),
]

# --- Constants: the shared numbers ----------------------------------------------------------------------------
# The numbers every robot of a file uses, once under the robot "all" (EVIDENCE.md E2.8, E2.8b, E5.14, E5.21, E5.22):
# the first six before a robot type's own shared numbers (type 1's for the robots whose sources stop, E5.17), the
# last five after them, in the order the files have always had.
SHARED_CONSTANTS_FIRST = [
    ('all', 'Board mass per cm2 (boards not weighed)', 0.28, 0.56, 0.84, 'g/cm2', 'E',
     'the weighed small driver board, 12.8 g on 22.9 cm2, ± 50 %; E2.8'),
    ('all', "Copper share of a board's mass", 0.125, 0.2, 0.287, 'share', 'S',
     'Bizzo et al. 2014, Table 3: 14 analyses combined (min / mean / max); E2.8b'),
    ('all', 'Power conductor cross-section (2 conductors)', 0.52, 0.82, 1.31, 'mm2', 'S',
     "20 / 18 / 16 AWG, the XT30's wire range (TME, second-hand)"),
    ('all', 'Signal conductor cross-section (2 conductors)', 0.065, 0.13, 0.2, 'mm2', 'E', 'E2.8'),
    ('all', 'Copper density', 8.96, 8.96, 8.96, 'g/cm3', 'S', 'copper'),
    ('all', 'Insulation and jackets, per mass of copper', 0.5, 1, 1.5, 'times', 'E', 'E2.8'),
]
SHARED_CONSTANTS_LAST = [
    ('all', 'Battery pack, Wh per kg (packs not weighed)', 109.3, 164, 232.5, 'Wh/kg', 'S, E',
     "the median of the G1's, X1's, N1's and PM01's packs (159 to 168); the pack at most 1.5 times its mode's mass, "
     "at least its cells at the NMC cells' lowest energy per kg (E5.14 point 3)"),
    ('all', 'Shell, share of the robot mass', 0.0575, 0.115, 0.1725, 'share', 'E',
     "the median of the G1's 12.9, the N1's 11.5 and the X1's 9.1 %; half to 1.5 times (E5.14 point 6)"),
    ('all', 'Sensor housing, aluminium share', 0.25, 0.5, 0.75, 'share', 'E',
     'a sensor module beyond its car sensor elements (the car file: "not the complete module housing"): its housing, '
     "aluminium as the LiDAR's and the depth camera's, plastics the remainder; half to 1.5 times (2026-10-29, point "
     '3c, its 30 / 50 / 70 % widened by the rule for estimates)'),
    ('all', 'Board rest, glass fibre share', 0.4, 0.55, 0.7, 'share', 'S, E',
     "the rest of a board beyond its metals: Isola's FR-4 laminates 35 to 60 % glass by weight, the populated boards' "
     'non-metals 46 to 72 %; the mode between a 1.6 mm board of 7628 cores and a dense board of thin plies '
     '(EVIDENCE.md E5.21; 2026-10-29, "Yes, enter them", point 3b)'),
    ('all', 'Board rest, bromine share', 0.044, 0.088, 0.132, 'share', 'E',
     'about 15 % TBBPA in an FR-4 laminate (Danish EPA 1999), 58.8 % of it bromine; half to 1.5 times (EVIDENCE.md '
     'E5.21; point 3b)'),
]

# --- Sources: those the shared rows cite ---------------------------------------------------------------------
# Each a row of the sheet Sources, by its code (EVIDENCE.md E2.1); a robot type's builder places them among its own.
SHARED_SOURCES = {
    'L1': ('L1', 'Livox, Mid-360 specifications', 'https://www.livoxtech.com/mid-360/specs',
           "Manufacturer's datasheet"),
    'T1': ('T1',
           'China Post Securities (Su, Sheng, Yang), Unitree G1 Humanoid Robot Teardown Report, '
           '2026-03-30', 'https://pdf.dfcfw.com/pdf/H3_AP202603311820899944_1.pdf',
           "A securities firm's own teardown, parts weighed"),
    'T3': ('T3', 'Munro & Associates, Munro Live, Unitree G1 teardown, 2026, through Humanoid.guide',
           'https://humanoid.guide/unitree-g1-teardown-reveals-actuator-and-cooling-tradeoffs/', 'Second-hand'),
    'T3k': ('T3k', "Munro's knee-actuator report (Grant), through SMZDM", 'https://post.smzdm.com/p/awwo7o5g/',
           'Second-hand'),
    'Bizzo 2014': ('Bizzo 2014', 'Bizzo, Figueiredo, de Andrade, Materials 7 (2014) 4555-4566',
           'https://pmc.ncbi.nlm.nih.gov/articles/PMC5455934/', 'University analyses, 14 combined'),
    'Drexler 2025': ('Drexler 2025', 'Drexler, Kampker, Born et al., Elektrotech. Inftech. 142 (2025) 312-345',
           'doi:10.1007/s00502-025-01331-3', 'Measured traction motors, through RAWCLICVehicleTractionMotor'),
}


# --- The template where a source stops ---------------------------------------------------------------------
# A figure a maker states once: ± 5 %, the spread between units (E5.12 point 3; the N1's accepted production
# tolerance of its battery, E5.14 point 2).
STATED_SPREAD = 0.05
# A joint's class: small below this peak torque, N·m, large from it (E5.12 point 1: the geometric middle of the G1's
# small joint, 35 N·m, and its large one, 88 N·m). Each block's joints are written with the class this gives.
LARGE_FROM_NM = 55

# The template's sizes and masses where a robot's sources name a board or sensor but give no size or mass (E5.14
# points 4 and 8, as the G1's, X1's and N1's).
# cm2: the G1's small joint's driver board, weighed at 12.8 g (T1 slides 24, 43), its area at most the 60 mm joint
# (E2.8): the area every joint's driver board is scaled from, in every robot type (a G1's row of type 1's sheet Boards;
# in another type's file the shared number DRIVER_AREA).
DRIVER_BOARD_AREA = (11.5, 22.9, 28.3)
DRIVER_AREA = "Driver board area, small joint (the G1's, weighed)"
COMPUTER = (60, 120, 180)        # cm2: a computer board, or the carrier of a compute module (as the X1's and N1's)
POWER = (30, 60, 100)            # cm2: a power distribution board (as the G1's)
BATTERY_BOARD = (30, 60, 96)     # cm2: a battery management board (as the G1's)
SMALL_BOARD = (5, 15, 30)        # cm2: a head, IMU or other small board (as the G1's head control board)
AGX_ORIN = (82.65, 87, 91.35)    # cm2: NVIDIA Jetson AGX Orin, 100 × 87 mm (search results, E5.13), ± 5 %
ORIN_NX = (29.74, 31.3, 32.87)   # cm2: NVIDIA Jetson Orin NX, 69.6 × 45 mm (search results, E5.13), ± 5 %
RADXA_CM5 = (20.9, 22, 23.1)     # cm2: Radxa CM5, 55 × 40 mm (search results, E5.13), ± 5 %
IMU_G = (1, 2, 5)                # g: a motion sensor (as the G1's)
MICROPHONES_G = (5, 10, 20)      # g: a microphone array (as the G1's)
SPEAKER_G = (17.5, 35, 60)       # g: a speaker (as the G1's)
D435_G = (72, 75, 80)            # g: an Intel RealSense D435 or D435i (Intel's datasheet, second-hand; as the G1's)
# kg: one XT30 (2+2) connector, a mated pair: Amass's cable female XT30(2+2)-F.G.B 1.1224 g and board male
# XT30PB(2+2)-M 1 g (LCSC, gross), ± 10 % (2026-10-29, "Yes, enter all"; EVIDENCE.md E5.22; 3 g before).
CONNECTOR_KG = (0.0019, 0.00212, 0.0023)
# The computer's cooler of a robot whose sources do not weigh it (Matthias, 2026-10-29, "Yes, build them", point 4;
# EVIDENCE.md E5.20, "Not entered": until then in its frame's remainder): a fan as the G1's main board fan (its copper
# winding by the sheet Shares), and for an ARM board (a Jetson, Rockchip, Radxa, Qualcomm or RDK, or a board not named)
# the G1's aluminium heat sink, for an Intel or AMD laptop-class processor the X1's and N1's copper heat pipes and fins
# (E2.8e). Each an entry of a block's `cooler`, as one of its `parts`.
FAN_KG = (0.005, 0.01, 0.02)
COOLER_FAN = ('Main board fan', 1, 'Torso boards', 'Torso', FAN_KG, None, None, 'E',
              "as the G1's main board fan (E2.8e; 2026-10-29, point 4)")
ARM_SINK = ('Main board heat sink', 1, 'Torso boards', 'Torso', (0.05, 0.1, 0.2), 'Aluminium', None, 'E',
            "an ARM board's: as the G1's (E2.8e; 2026-10-29, point 4)")
X86_SINK = ('Main board heat sink: heat pipes and fins', 1, 'Torso boards', 'Torso', (0.055, 0.12, 0.19), 'Copper',
            None, 'E', "an Intel or AMD laptop-class processor's: as the X1's and N1's (E2.8e; 2026-10-29, point 4)")
ARM_COOLER = [COOLER_FAN, ARM_SINK]
X86_COOLER = [COOLER_FAN, X86_SINK]
SCREWS_KG = (0.2, 0.5, 1)        # kg: a robot's screws (as the G1's; E5.14 point 7)
CABLES_M = (5, 10, 15)           # m of four conductors where no length is listed (as the N1's; E5.14 point 5)
# A robot's cables in the sheet Constants: four conductors (two for power, two for signals: type 1's, E2.8), or two
# for power and an Ethernet line of eight (the Fourier GR-1's and GR-2's "46V电源线和以太网线", EVIDENCE.md
# E6.8 point 9). 01 reads the same two names (CABLE_LENGTHS).
CABLE_FOUR = 'Cable length (four conductors)'
CABLE_ETHERNET = 'Cable length (two power conductors and an Ethernet line)'
# The car sensor elements of a sensor whose maker names its kind but not its model (E5.14 point 8, as type 1's).
DEPTH_CAMERA = [('laser diode array', 1), ('IR camera sensor', 2), ('CMOS image sensor', 1)]   # as the D435's
LIDAR = [('laser diode array', 1), ('photodetector array', 1), ('IMU', 1)]                      # as the Mid-360's
CAMERA = [('CMOS image sensor', 1)]
MICROPHONES = [('MEMS microphone array', 4)]   # as the G1's four (T1 slide 33)
ENCODER = [('angle sensor', 1)]                # a joint's encoder: the car's angle sensor (as the G1's, E2.8b)
# A power distribution board where a robot's sources name none (as the G1's; E5.14 point 4).
POWER_BOARD = ('Power distribution board', 1, 'Torso boards', 'Torso', POWER, 'PE_HVS', 'E',
               'as the G1 (E5.14 point 4)')

# The texts of the sheets that 02 reads as what a row is computed from (its column "Remainder of, or computed from"),
# and the basis every joint's magnets share.
JOINT_MASSES = "the Joints' module masses; inside as JointParts' joint of its size, scaled to its mass"
JOINT_SHARES = ("the Joints' shares of the joint modules, a share of the robot's mass (Constants); inside as "
                "JointParts' joint of its size, scaled to its mass")
# A size of joint whose rows are partly shares, partly stated masses: a template robot's hands or grippers (outside the
# template, which counts body joints only: EVIDENCE.md E7.4), or its stated joints in place of a group of the template
# (the JingChu's knees, E7.5 point 2).
JOINT_MIXED = ("the Joints' module masses where given, else their shares of the joint modules, a share of the robot's "
               "mass (Constants); inside as JointParts' joint of its size, scaled to its mass")
DRIVER_FROM = "JointParts' driver board of its size, scaled to its joint's module mass; area likewise"
PACK_STATED = 'its cells, its management board (Boards) and its housing (the rows below)'
PACK_FROM_ENERGY = ('its cells, its management board (Boards) and its housing (the rows below); its mass the energy '
                    "over Constants' Battery pack, Wh per kg")
PACK_CELLS_ONLY = ("its cells, its management board (Boards) and no housing: the maker's pack is lighter than its "
                   'cells (E5.14 point 3)')
CELLS_FROM_CHARGE = "from the energy: Constants voltage × capacity over the NMC cell's Wh/kg"
CELLS_FROM_MASS = ("from the energy: Constants' Battery energy per kg of robot times the robot's mass, over the NMC "
                   "cell's Wh/kg")
HOUSING = 'the battery pack, less its cells and its management board'
SHELL = "a share of the robot's mass: Constants' Shell, share of the robot mass"
FRAME = 'the robot (Constants), less every other part'
MAGNETS_BASIS = 'Its magnets: the grade classes of its heat group, sheet MagnetGrades (EVIDENCE.md E2.11)'
# The housing of a battery pack: its shares in the sheet Shares, the housing aluminium or plastics (each as often), or
# aluminium where the maker says so (the PM01).
HOUSING_MIXED = 'Battery housing, busbars and wires'
HOUSING_ALUMINIUM = 'Battery housing (aluminium), busbars and wires'
# The joints of a robot whose sources give no joint torques (E5.17 point 1): AVERAGE in its block. The text is also
# type 1's name of what its template's rows average; another type names its own (average_joints()).
AVERAGE = 'as batch 1 on average'


def spread(mode: float, low: float = 1 - STATED_SPREAD, high: float = 1 + STATED_SPREAD) -> tuple:
    """Min, mode and max: `low` and `high` times the mode, rounded to four decimals. By default ± 5 %, a figure stated
    once."""
    return round(low * mode, 4), round(mode, 4), round(high * mode, 4)


def estimate(mode: float) -> tuple:
    """An estimate (E): half to 1.5 times its mode (Matthias, 2026-10-27: "Estimates should be larger than known
    numbers")."""
    return spread(mode, 0.5, 1.5)


def size_of(torque: float) -> str:
    """A joint's class by its actuator's peak torque (E5.12 point 1)."""
    return 'large' if torque >= LARGE_FROM_NM else 'small'


def joint(name: str, count: int, torque: float | None, masses: tuple, heat_group: str, region: str,
          size: str | None = None) -> tuple:
    """One joint kind of a block: its name, joints, class (by its torque, or `size` where no torque is stated), torque
    (N·m), module mass (kg), heat group and body region."""
    return name, count, size or size_of(torque), torque, masses, heat_group, region


def block(name: str, maker: str, height: float, stated: float | None, sources: str, mass: tuple, joints,
          **more) -> dict:
    """One robot's block: what its sheet Robots row says (`maker`, `height` m, `stated` mass kg, `sources`), its mass
    for the sheet Constants (`mass`: (min, mode, max), label, source), its joints (joint() each, or AVERAGE), and in
    `more` what differs from the template (E5.14): voltage and capacity ((min, mode, max), label, source; none: the
    energy per kg of robot), pack ((min, mode, max), label, source; None: from its energy; PACK_CELLS_ONLY), housing,
    boards, sensors, cables (and cable_quantity, CABLE_FOUR or CABLE_ETHERNET), screws, parts, cooler (ARM_COOLER
    unless its computer is an Intel or AMD processor's), shell, frame, and the labels and sources of its joints,
    connectors, encoders and cells. A block whose joints are AVERAGE may also have `stated_joints` (joint() each), which
    take the place of their groups' rows of the template, and `added_joints` (its hands or grippers), which the
    template leaves out (it counts body joints only: EVIDENCE.md E7.4, E7.5 points 2 and 4)."""
    b = dict(name=name, maker=maker, height=height, stated=stated, sources=sources, mass=mass, joints=joints,
             stated_joints=[], added_joints=[],
             joints_stated=None, joint_label={'small': 'E', 'large': 'E'}, joint_sources='', voltage=None,
             capacity=None, pack=None, pack_source='', housing=HOUSING_MIXED,
             cells=('E', 'NMC cylindrical, E (E5.14 point 2)'),
             boards=[], sensors=[], cables=(CABLES_M, 'E', 'as the N1 (E5.14 point 5)'), cable_quantity=CABLE_FOUR,
             connectors=('S, E', 'two per joint, as type 1 (E5.14 point 5); the mass of a mated pair, Amass '
                                 'through LCSC (E5.22)'),
             encoders=('E', 'two per joint, as type 1 (E5.14)'),
             screws=[('Screws', 1, SCREWS_KG, '42CrMo4', 'E', 'as the G1 (E5.14 point 7); grade E5.18')],
             parts=[], cooler=ARM_COOLER, shell=('Shell: plastics', 'not stated; plastics as type 1 (E5.14 point 6)'),
             frame=('Frame', 'Aluminium', 'not stated; aluminium as type 1 (E5.14 point 9)'))
    b.update(more)
    return b


def joint_group(j: tuple) -> tuple:
    """The group of a joint for the average split: its name in the sheet Joints, heat group and body region."""
    if j[6] == 'Head':
        return 'Head joints', 'Elbows and wrists', 'Head'
    return {'Hips': ('Hip joints', 'Hips', 'Legs'), 'Knees': ('Knee joints', 'Knees', 'Legs'),
            'Ankles': ('Ankle joints', 'Ankles', 'Legs'), 'Waist': ('Waist joints', 'Waist', 'Torso'),
            'Shoulders': ('Shoulder joints', 'Shoulders', 'Arms'),
            'Elbows and wrists': ('Elbow and wrist joints', 'Elbows and wrists', 'Arms')}[j[5]]


def average_joints(blocks: list, label: str = AVERAGE) -> list:
    """The joints of a robot whose sources give no joint torques (E5.17 point 1: "split into small and large joints as
    batch 1's are on average"): for each group of joints (hips, knees, ankles, waist, shoulders, elbows and wrists,
    head) and class, its share of the joint modules' mass at the modes, averaged over `blocks` (type 1: its batch 1;
    type 2: its seven robots with every joint's torque, without their hands, E7.4); one row each, the joints of the
    group and class together, named with `label`, what they average. The shares add up to 1, as each robot's do;
    nothing is rescaled."""
    shares: dict = {}
    for b in blocks:
        total = sum(j[1] * j[4][1] for j in b['joints'])
        for j in b['joints']:
            key = joint_group(j) + (j[2],)
            shares[key] = shares.get(key, 0) + j[1] * j[4][1] / total / len(blocks)
    order = ['Hip joints', 'Knee joints', 'Ankle joints', 'Waist joints', 'Shoulder joints', 'Elbow and wrist joints',
             'Head joints']
    return [(f'{name}, {size} ({label})', 1, size, None, None, heat, region, share)
            for (name, heat, region, size), share in sorted(shares.items(), key=lambda kv: (order.index(kv[0][0]),
                                                                                          kv[0][3] == 'large'))]


def regions(joints: list) -> str:
    """The body regions of joints with their numbers, as the sheet Parts writes them ('Legs 12, Torso 1, Arms 10')."""
    counts: dict = {}
    for j in joints:
        counts[j[6]] = counts.get(j[6], 0) + j[1]
    return ', '.join(f'{region} {n}' for region, n in counts.items())


def template_joints(b: dict, averaged_joints: list) -> list:
    """The joints of a block whose joints are AVERAGE: the template's rows in their order, each group of the template
    for which the block states joints (`stated_joints`, the JingChu's knees) replaced by them, then its hands or
    grippers (`added_joints`). Nothing is rescaled: the rows kept keep the template's shares (EVIDENCE.md E7.5 point 2,
    open point A)."""
    groups = [a[0].split(',')[0] for a in averaged_joints]
    stated = {}
    for j in b['stated_joints']:
        group = joint_group(j)[0]
        if group not in groups:
            raise ValueError(f"{b['name']}: its stated joint {j[0]} is of no group of the template ({group}).")
        stated.setdefault(group, []).append(j)
    joints, placed = [], set()
    for a, group in zip(averaged_joints, groups):
        if group not in stated:
            joints.append(a)
        elif group not in placed:
            joints += stated[group]
            placed.add(group)
    return joints + list(b['added_joints'])


def template_block(name: str, maker: str, height: float, stated: float | None, mass: tuple, joints_stated: int,
                   computer: str, *, computer_source: str, joint_sources: str, default_sources: str, **more) -> dict:
    """A robot whose sources give no joint torques (type 1: E5.17; type 2: E7.4, E7.5): its joints a share of its
    mass, split as its type's robots with torques are on average (its block's joints AVERAGE, `joint_sources` saying
    so); its battery from the energy per kg of robot unless a voltage or capacity is given; a computer board as the
    template's, named `computer` as its sources name it (`computer_source`), and the template's power distribution
    board; the rest as E5.14 unless `more` says otherwise (its `sources` else `default_sources`). One copy for every
    type (Matthias, 2026-10-30, "Yes, build it", open point C): each type's builder binds its own texts."""
    defaults = dict(joints_stated=joints_stated,
                    boards=[(computer, 1, 'Torso boards', 'Torso', COMPUTER, 'ITC', 'E', computer_source),
                            ('Power distribution board', 1, 'Torso boards', 'Torso', POWER, 'PE_HVS', 'E',
                             'as the G1 (E5.14 point 4)')],
                    joint_sources=joint_sources)
    defaults.update(more)
    return block(name, maker, height, stated, defaults.pop('sources', default_sources), mass, AVERAGE, **defaults)


def robot_rows(b: dict, robot_type: int, averaged_joints: list | None = None) -> dict:
    """Every row of one robot's block on the sheets Robots, Joints, Parts, Boards, Sensors and Constants, and the car
    sensor elements of its sensors (sheet SensorElements, by the sensor's name). `robot_type` is the type of its file;
    `averaged_joints` (average_joints() of the type's batch 1) are the joints of a block whose joints are AVERAGE."""
    name = b['name']
    out = {sheet: [] for sheet in ('Robots', 'Joints', 'Parts', 'Boards', 'Sensors', 'Constants', 'SensorElements')}
    out['Robots'].append((name, b['maker'], robot_type, b['height'], b['stated'], b['sources']))

    # Its joints, their modules and driver boards, and two encoders and two connectors per joint. A robot of the
    # template: the template's rows, its stated joints in place of their groups' rows, and its hands or grippers.
    averaged = b['joints'] == AVERAGE
    joints = template_joints(b, averaged_joints) if averaged else b['joints']
    for number, (joint_name, count, size, torque, masses, heat, region, *share) in enumerate(joints, start=1):
        low, mode, high = masses or (None, None, None)
        out['Joints'].append((name, number, joint_name, count, size, torque, low, mode, high,
                              share[0] if share else None, 'not stated', heat, region, b['joint_sources'],
                              MAGNETS_BASIS))
    in_robot = b['joints_stated'] or sum(j[1] for j in joints)
    for size in ('small', 'large'):
        mine = [j for j in joints if j[2] == size]
        if mine:
            count, where = sum(j[1] for j in mine), regions(mine)
            shares = [len(j) > 7 for j in mine]   # a row of the template carries its share as an eighth item
            basis = JOINT_SHARES if all(shares) else JOINT_MIXED if any(shares) else JOINT_MASSES
            out['Parts'].append((name, f'Joint module, {size}', count, 'In the joints', where, None, None, None, None,
                                 None, basis, b['joint_label'][size], b['joint_sources']))
            out['Boards'].append((name, f'Driver board, {size} joint', count, 'In the joints', where, None, None, None,
                                  DRIVER_FROM, 'PE_HVS', 'E', 'as the G1 (E5.12 point 2)'))
    label, source = b['encoders']
    out['Parts'].append((name, 'Encoders (two per joint)', 2 * in_robot, 'In the joints', 'as their joints', None,
                         None, None, None, None, 'Sensors and SensorElements (Encoder)', label, source))
    out['Sensors'].append((name, 'Encoder', 2 * in_robot, 'In the joints', 'as their joints', None, None, None, label,
                           f'{source}; elements only'))
    for car_type, per_sensor in ENCODER:
        out['SensorElements'].append(('Encoder', car_type, per_sensor))

    # Its battery: the pack, its cells from their energy, its housing the pack's remainder, its management board.
    if b['pack'] == PACK_CELLS_ONLY:
        out['Parts'].append((name, 'Battery pack', 1, 'Battery pack', 'Torso', None, None, None, None, None,
                             PACK_CELLS_ONLY, 'S', b['pack_source']))
    elif b['pack'] is None:
        out['Parts'].append((name, 'Battery pack', 1, 'Battery pack', 'Torso', None, None, None, None, None,
                             PACK_FROM_ENERGY, 'S, E', 'its mass not stated (E5.14 point 3)'))
    else:
        (low, mode, high), label, source = b['pack']
        out['Parts'].append((name, 'Battery pack', 1, 'Battery pack', 'Torso', low, mode, high, None, None,
                             PACK_STATED, label, source))
    label, source = b['cells']
    out['Parts'].append((name, 'Battery cells', 1, 'Battery pack', 'Torso', None, None, None, 'Battery materials', None,
                         CELLS_FROM_CHARGE if b['voltage'] else CELLS_FROM_MASS, label, source))
    if b['housing']:
        out['Parts'].append((name, b['housing'], 1, 'Battery pack', 'Torso', None, None, None, None, None, HOUSING,
                             'E', 'E5.19 point 1'))
    out['Boards'].append((name, 'Battery management board', 1, 'Battery pack', 'Torso', *BATTERY_BOARD,
                          'area times Constants board mass per cm2', 'BMS', 'E', 'as the G1 (E5.14 point 4)'))

    # Its other boards and its sensors, as its sources name them.
    for board, count, kind, region, area, category, label, source in b['boards']:
        out['Boards'].append((name, board, count, kind, region, *area, 'area times Constants board mass per cm2',
                              category, label, source))
    for sensor, count, kind, region, grams, label, source, units in b['sensors']:
        out['Sensors'].append((name, sensor, count, kind, region, *(grams or (None, None, None)), label, source))
        for car_type, per_sensor in units or []:
            out['SensorElements'].append((sensor, car_type, per_sensor))

    # Its cables and connectors, screws and other parts, shell and frame.
    out['Parts'].append((name, 'Cable conductors and insulation', 1, 'Cables', 'Whole body', None, None, None, None,
                         None, 'from Constants: length, cross-sections, insulation', b['cables'][1], b['cables'][2]))
    label, source = b['connectors']
    out['Parts'].append((name, 'XT30 (2+2) connectors (two per joint)', 2 * in_robot, 'Cables', 'Whole body',
                         *CONNECTOR_KG, None, None, None, label, source))
    for part, count, kg, grade, label, source in b['screws']:
        out['Parts'].append((name, part, count, 'Frame and shell', 'Whole body', *kg, 'Steel', grade, None, label,
                             source))
    for part, count, kind, region, kg, material, grade, label, source in b['cooler'] + b['parts']:
        out['Parts'].append((name, part, count, kind, region, *kg, material, grade, None, label, source))
    shell, source = b['shell']
    out['Parts'].append((name, shell, 1, 'Frame and shell', 'Whole body', None, None, None, 'Plastics', None, SHELL,
                         'E', source))
    frame, material, source = b['frame']
    out['Parts'].append((name, frame, 1, 'Frame and shell', 'Whole body', None, None, None, material, None, FRAME, 'E',
                         source))

    # Its own constants: its mass, its battery's voltage and capacity (or none: its energy per kg of robot), its cables.
    values, label, source = b['mass']
    out['Constants'].append((name, 'Robot mass, with battery', *values, 'kg', label, source))
    if b['voltage']:
        for quantity, unit, (values, label, source) in (('Battery voltage', 'V', b['voltage']),
                                                         ('Battery capacity', 'Ah', b['capacity'])):
            out['Constants'].append((name, quantity, *values, unit, label, source))
    values, label, source = b['cables']
    out['Constants'].append((name, b['cable_quantity'], *values, 'm', label, source))
    return out


def write_sheet(book: openpyxl.Workbook, name: str, headings: list, rows: list, widths: list) -> None:
    """One sheet: its first row in bold on grey, then its rows, its column widths, the first row kept in view."""
    sheet = book.create_sheet(name)
    sheet.append(headings)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADING_FILL
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    for row in rows:
        sheet.append(list(row))
    for i, width in enumerate(widths):
        sheet.column_dimensions[get_column_letter(i + 1)].width = width
    sheet.freeze_panes = 'A2'


def sensor_elements(given: list, rows: list) -> list:
    """The sheet SensorElements: the rows `given` (a builder's own, by the sensor's name), then the car sensor elements
    of the generated robots' sensors (`rows`, robot_rows()'s 'SensorElements'), each sensor once; a sensor named twice
    must be made of the same elements."""
    known = {r[0]: [] for r in given}
    for r in given:
        known[r[0]].append(r[1:])
    new: dict = {}
    for sensor, car_type, units in rows:
        new.setdefault(sensor, [])
        if (car_type, units) not in new[sensor]:
            new[sensor].append((car_type, units))
    for sensor, units in new.items():
        if sensor in known and sorted(known[sensor]) != sorted(units):
            raise ValueError(f'{sensor}: named with other elements than in the rows given.')
    return list(given) + [(sensor, *u) for sensor, units in new.items() if sensor not in known for u in units]


def sheets(robots: list, joints: list, parts: list, boards: list, sensors: list, elements: list, constants: list,
           sources: list) -> list:
    """Every sheet after About, in the order of the file: its name, first row, rows and column widths. The robots' own
    rows are a builder's; MagnetGrades, JointParts and Shares are the one set for all, Shares only the rows of a
    component of JointParts or a part of `parts` (01 faults a share of a part the file does not have)."""
    named = {r[0] for r in JOINTPARTS_ROWS} | {r[1] for r in parts}
    return [
        ('Robots', ROBOTS_HEADINGS, robots, ROBOTS_WIDTHS),
        ('Joints', JOINTS_HEADINGS, joints, JOINTS_WIDTHS),
        ('MagnetGrades', MAGNETGRADES_HEADINGS, MAGNETGRADES_ROWS, MAGNETGRADES_WIDTHS),
        ('JointParts', JOINTPARTS_HEADINGS, JOINTPARTS_ROWS, JOINTPARTS_WIDTHS),
        ('Parts', PARTS_HEADINGS, parts, PARTS_WIDTHS),
        ('Shares', SHARES_HEADINGS, [r for r in SHARES_ROWS if r[0] in named], SHARES_WIDTHS),
        ('Boards', BOARDS_HEADINGS, boards, BOARDS_WIDTHS),
        ('Sensors', SENSORS_HEADINGS, sensors, SENSORS_WIDTHS),
        ('SensorElements', SENSORELEMENTS_HEADINGS, elements, SENSORELEMENTS_WIDTHS),
        ('Constants', CONSTANTS_HEADINGS, constants, CONSTANTS_WIDTHS),
        ('Sources', SOURCES_HEADINGS, sources, SOURCES_WIDTHS),
    ]


def type_output_path(folder: str) -> str:
    """Where a robot type's file goes: 00's setting `data.parts_taken_apart` with the type's folder ('<n>_<name>', the
    folder its builder sits in), as an absolute path."""
    return absolute(Params().data.parts_taken_apart.format(folder=folder))


def write_file(path: str, about: list, all_sheets: list) -> None:
    """Write the whole file to `path` (its folder made where it is not there): the About sheet first, its lines
    `about`, then every sheet of `all_sheets` (sheets()) in its order."""
    os.makedirs(os.path.dirname(path), exist_ok=True)
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'About'
    for row, line in enumerate(about, start=1):
        if line is not None:
            sheet.cell(row=row, column=1, value=line)
    sheet['A1'].font = Font(bold=True)
    sheet.column_dimensions['A'].width = ABOUT_WIDTH
    for name, headings, rows, widths in all_sheets:
        write_sheet(book, name, headings, rows, widths)
    book.save(path)
