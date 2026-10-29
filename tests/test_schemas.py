"""
Tests of 01_schemas.py, the schema of each category of humanoid robot.

**Copyright notice:** Copyright © 2026 Empa, Matthias Roesslein
**Licence:** CC BY 4.0, Creative Commons Attribution 4.0 International; the full text is in `LICENSE`
at the root of the repository.

Matthias, 2026-10-26: "Draw a repective schema and the we will go through and do separate code for each
of the different categories."

Most tests build a small workbook of their own, in the layout of
data/raw/humanoid_robots_2040_workbook.xlsx (two categories, three parts), break one thing in it and
check that 01 reports exactly that fault: a check is only trusted once it has been seen to catch what it
guards. One test reads the real workbook and is skipped without it. Everything is written into temporary
folders: press Run, nothing in the project is changed.
"""

import glob
import importlib.util
import os
import re
import sys
import tempfile
import types
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
_spec = importlib.util.spec_from_file_location('schemas', os.path.join(ROOT, '01_schemas.py'))
schemas = importlib.util.module_from_spec(_spec)
sys.modules['schemas'] = schemas  # the dataclasses of 01 look their module up here while it loads
_spec.loader.exec_module(schemas)

from src.params_schema import Params, absolute  # noqa: E402

HEADER = ['Category', 'BOM Anchor Robot', 'Subsystem', 'Component', 'Location / Joint', 'Qty',
          'Per-Unit Min (kg)', 'Per-Unit Mode (kg)', 'Per-Unit Max (kg)', 'Total Min (kg)', 'Total Mode (kg)',
          'Total Max (kg)', 'Primary Material Grade', 'Secondary / Mixed Material', 'EU Critical Raw Material',
          'Evidence', 'Source Reference']
D = '—'  # the workbook's dash for "not applicable"

# The small workbook: (category, anchor, subsystem, component, location, quantity, mass of one min/mode/max,
# mass of all min/mode/max, material, evidence). Category 1: two parts; category 2: one part, and an anchor
# robot that is 1 kg lighter than its parts. The header is row 3, so the rows below are rows 4 to 12.
ROWS = [
    ('1. Small', 'Robot A', 'Actuation', 'Knee actuator', 'Knees', 2, 1, 2, 3, 2, 4, 6, 'Aluminium', 'Estimate'),
    ('1. Small', 'Robot A', 'Structure', 'Frame', 'Torso', 1, 0.5, 1, 1.5, 0.5, 1, 1.5, 'Steel', 'Direct'),
    ('1. Small', 'Robot A', 'CATEGORY SUBTOTAL', 'Subtotal', 'All', 3, D, D, D, 2.5, 5, 7.5, 'Sum', 'Derived'),
    ('1. Small', 'Robot A', 'REPORTED ANCHOR MASS', 'Stated mass', 'Robot', 1, D, 5, D, D, 5, D, 'Spec', 'Direct'),
    ('1. Small', 'Robot A', 'RECONCILIATION CHECK', 'Difference', 'Check', D, D, D, D, D, 0, 0, 'K, L', 'Derived'),
    ('2. Large', "Notional 'B'", 'Energy', 'Battery', 'Torso', 4, 1, 1.5, 2, 4, 6, 8, 'Lithium cells', 'Derived'),
    ('2. Large', "Notional 'B'", 'CATEGORY SUBTOTAL', 'Subtotal', 'All', 4, D, D, D, 4, 6, 8, 'Sum', 'Derived'),
    ('2. Large', "Notional 'B'", 'REPORTED ANCHOR MASS', 'Stated mass', 'Robot', 1, D, 5, D, D, 5, D, 'Spec',
     'Estimate'),
    ('2. Large', "Notional 'B'", 'RECONCILIATION CHECK', 'Difference', 'Check', D, D, D, D, D, 1, 0.2, 'K, L',
     'Derived'),
]
GRAND = ('ALL 2 CATEGORIES', '2 robots', 'GRAND TOTAL (SUM OF 2 ANCHORS)', 'Sum', 'All', 7, D, D, D, 6.5, 11,
         15.5, 'Sum', 'Derived')
FILLS = {'Derived': 'FFFFF2CC', 'Estimate': 'FFFCE4D6'}   # Direct has no fill, as in the workbook's key
COLUMN = {name: i + 1 for i, name in enumerate(HEADER)}

# The materials of the small workbook, in the layout of 'Material Composition per Cat': (category, anchor, family,
# grade, minimum, mode, maximum, share, evidence). Category 1: 5 kg at the mode, all of its parts list; category
# 2: 4 kg, which leaves 2 kg of its 6 kg parts list and 1 kg of its 5 kg stated mass unallocated. The header is
# row 3, so the rows below are rows 4 to 10.
MATERIAL_HEADER = ['Category', 'BOM Anchor Robot', 'Material Family', 'Proposed Grade / Specification',
                   'Principal Allocation / Subsystem Role', 'Min (kg)', 'Mode (kg)', 'Max (kg)',
                   'Share of Category Material Mode (%)', 'Evidence', 'Source Reference']
MATERIAL_ROWS = [
    ('1. Small', 'Robot A', 'Aluminum alloys', 'Al 6061', 2, 3, 4.5, 0.6, 'Estimate'),
    ('1. Small', 'Robot A', 'Copper conductors', 'Copper', 1, 2, 3, 0.4, 'Estimate'),
    ('1. Small', 'Robot A', 'CATEGORY MATERIAL SUBTOTAL', 'Sum', 3, 5, 7.5, 1, 'Derived'),
    ('1. Small', 'Robot A', 'UNALLOCATED MATERIAL RESIDUAL', 'Gap', D, 0, 0, 1, 'Derived'),
    ('2. Large', "Notional 'B'", 'Battery materials', 'NMC', 3, 4, 5, 1, 'Estimate'),
    ('2. Large', "Notional 'B'", 'CATEGORY MATERIAL SUBTOTAL', 'Sum', 3, 4, 5, 1, 'Derived'),
    ('2. Large', "Notional 'B'", 'UNALLOCATED MATERIAL RESIDUAL', 'Gap', D, 2, 1, 1, 'Derived'),
]
MATERIAL_COLUMN = {name: i + 1 for i, name in enumerate(MATERIAL_HEADER)}

# The shared lists of the small workbook, in the layout of data/researched/shared_lists.xlsx.
LISTS = [('Part group', 'Joint drives'), ('Part group', 'Frame'), ('Part group', 'Battery'),
         ('Part group', 'Mixed'), ('Body region', 'Legs'), ('Body region', 'Torso'), ('Body region', 'Whole body'),
         ('Material', 'Aluminium'), ('Material', 'Copper'), ('Material', 'Battery materials'),
         ('Material', 'Mixed'), ('Material', 'Not allocated')]
PLACED_PARTS = [(1, 4, 'Knee actuator', 'Joint drives', 'Legs', '', 2, 'The location names one region.'),
                (1, 5, 'Frame', 'Frame', 'Whole body', '', 0, 'Spread over the body.'),
                (2, 9, 'Battery', 'Battery', 'Torso', '', 0, 'The location names one region.')]
PLACED_MATERIALS = [('Aluminum alloys', 'Aluminium', 'The name says the material.'),
                    ('Copper conductors', 'Copper', 'The name says the material.'),
                    ('Battery materials', 'Battery materials', 'The name says the material.')]


BATTERY_HEADINGS = ['Robot type', 'Battery row in the workbook', 'Kind', 'Cell or material', 'Cells in series',
                    'Cells in parallel', 'Packs', 'Pack voltage (V)', 'Pack capacity (Ah)', 'Pack energy (kWh)',
                    'Pack mass (kg)', 'Share min', 'Share mode', 'Share max', 'Remainder', 'Label', 'Basis', 'Decided']
# The battery of the small workbook's robot type 2 ('Battery materials', 3 / 4 / 5 kg), counted from its cells: a
# 36 V, 4 Ah pack of 'Cell A' (10 in series, 2 in parallel: 20 × 50 g = 1.0 kg) or of 'Cell B' (10 × 1: 10 × 90 g =
# 0.9 kg). The rest of the row goes to Not allocated.
BATTERY = [(2, 'Battery materials', 'Build', 'Cell A', 10, 2, 1, 36, 4, None, None, None, None, None, '', 'E', '', ''),
           (2, 'Battery materials', 'Build', 'Cell B', 10, 1, 1, 36, 4, None, None, None, None, None, '', 'E', '', ''),
           (2, 'Battery materials', 'Rest', 'Not allocated', None, None, None, None, None, None, None, None, None,
            None, 'Yes', 'E', '', '')]


def battery_row(kind, name, series=None, parallel=None, packs=None, volts=None, amp_hours=None, kwh=None,
                pack_kg=None, share=(None, None, None), remainder='', label='E'):
    """One row of the sheet Battery cells for robot type 2 of the small workbook."""
    return (2, 'Battery materials', kind, name, series, parallel, packs, volts, amp_hours, kwh, pack_kg, *share,
            remainder, label, '', '')


def build_lists(path, parts=None, materials=None, lists=None, splits=None, battery=None):
    """Write small shared lists to `path`: the rows of LISTS, PLACED_PARTS and PLACED_MATERIALS, or the ones
    given instead, a sheet Splits with the rows of `splits` and a sheet Battery cells with the rows of `battery`,
    each when it is given."""
    book = Workbook()
    sheet = book.active
    sheet.title = 'Lists'
    sheet.append(['List', 'Name', 'What it holds'])
    for kind, name in LISTS if lists is None else lists:
        sheet.append([kind, name, 'For the tests.'])
    sheet = book.create_sheet('Parts')
    sheet.append(['Robot type', 'Row', 'Part', 'Part group', 'Body region', 'Split by units', 'Motors', 'Why'])
    for row in PLACED_PARTS if parts is None else parts:
        sheet.append(list(row))
    sheet = book.create_sheet('Materials')
    sheet.append(['Material name in the workbook', 'Shared material', 'Why'])
    for row in PLACED_MATERIALS if materials is None else materials:
        sheet.append(list(row))
    if splits is not None:
        sheet = book.create_sheet('Splits')
        sheet.append(['Robot type', 'Material name in the workbook', 'Shared material', 'Share min', 'Share mode',
                      'Share max', 'Remainder', 'Label', 'Basis', 'Decided'])
        for row in splits:
            sheet.append(list(row))
    if battery is not None:
        sheet = book.create_sheet('Battery cells')
        sheet.append(BATTERY_HEADINGS)
        for row in battery:
            sheet.append(list(row))
    book.save(path)


# The element recipes of the tests, in the layout of data/researched/element_recipes.xlsx: two cells, a few
# compounds and one small recipe. 'Test recipe': a can of 20 %, either plain iron or nickel-plated (10 % nickel,
# iron the remainder), half and half; 10 % salt (LiPF6); the rest carbon.
CELL_HEADINGS = ['Cell', 'Chemistry', 'Format', 'Voltage (V)', 'Capacity (Ah)', 'Mass (g)', 'Energy (Wh)', 'Recipe',
                 'Source', 'Kind']
CELLS = [('Cell A', 'NMC', '18650', 3.6, 2.0, 50, None, 'Test recipe', 'For the tests.', 'Test'),
         ('Cell B', 'NMC', '21700', 3.6, 4.0, 90, None, 'Test recipe', 'For the tests.', 'Test')]
COMPOUND_HEADINGS = ['Compound', 'Composition', 'Basis']
COMPOUNDS = [('Iron', 'Fe', ''), ('Nickel', 'Ni', ''), ('Carbon', 'C', ''), ('LiPF6', 'LiPF6', ''),
             ('Half and half', 'C 0.5; H 0.5', '')]
RECIPE_HEADINGS = ['Recipe', 'Within', 'Item', 'Share min', 'Share mode', 'Share max', 'Remainder',
                   'Probability', 'Compound', 'Follows the energy', 'Label', 'Basis']
RECIPE = [('Test recipe', '', 'Can', 0.2, 0.2, 0.2, '', None, '', '', 'S', ''),
          ('Test recipe', 'Can', 'Iron can', None, None, None, '', 0.5, 'Iron', '', 'E', ''),
          ('Test recipe', 'Can', 'Plated can', None, None, None, '', 0.5, '', '', 'E', ''),
          ('Test recipe', 'Plated can', 'Plating', 0.05, 0.1, 0.15, '', None, 'Nickel', '', 'E', ''),
          ('Test recipe', 'Plated can', 'Steel', None, None, None, 'Yes', None, 'Iron', '', 'E', ''),
          ('Test recipe', '', 'Salt', 0.05, 0.1, 0.15, '', None, 'LiPF6', 'Yes', 'E', ''),
          ('Test recipe', '', 'Rest', None, None, None, 'Yes', None, 'Carbon', '', 'E', '')]


# A small recipe within a recipe: 'Inner' is half carbon, half 'Half and half'; 'Outer' holds 40 % iron and 'Inner' as
# its remainder.
NESTED = [('Inner', '', 'Coal', 0.5, 0.5, 0.5, '', None, 'Carbon', '', 'E', ''),
          ('Inner', '', 'Mix', None, None, None, 'Yes', None, 'Half and half', '', 'E', ''),
          ('Outer', '', 'Iron part', 0.3, 0.4, 0.5, '', None, 'Iron', '', 'E', ''),
          ('Outer', '', 'Inner part', None, None, None, 'Yes', None, 'Inner', '', 'E', '')]
ROW_HEADINGS = ['Robot type', 'Material name in the workbook', 'Shared material', 'Recipe', 'Basis']


def build_recipes(path, cells=None, compounds=None, recipe=None, rows=None):
    """Write small element recipes to `path`: the rows of CELLS, COMPOUNDS and RECIPE, or the ones given, and a
    sheet Rows with `rows` when it is given."""
    book = Workbook()
    book.active.title = 'About'
    book.active.append(['For the tests.'])
    for name, headings, rows in (('Cells', CELL_HEADINGS, CELLS if cells is None else cells),
                                 ('Compounds', COMPOUND_HEADINGS, COMPOUNDS if compounds is None else compounds),
                                 ('Recipe', RECIPE_HEADINGS, RECIPE if recipe is None else recipe)) + \
            ((('Rows', ROW_HEADINGS, rows),) if rows is not None else ()):
        sheet = book.create_sheet(name)
        sheet.append(headings)
        for row in rows:
            sheet.append(list(row))
    book.save(path)


def build(path, changes=None, taxonomy=None, coverage=3, material_changes=None):
    """Write the small workbook to `path`; `changes` maps (row, heading) of the parts to a value that replaces
    the one of ROWS, or to ('fill', ARGB) to recolour the row; `material_changes` does the same for
    MATERIAL_ROWS."""
    changes = changes or {}
    book = Workbook()
    sheet = book.active
    sheet.title = 'Taxonomy'
    sheet.append(['Category', 'Dominant Requirement & Definition', 'Working Boundary (Height & Mass)',
                  'Articulation (DoF) & Payload', 'BOM Anchor Robot'])
    for row in taxonomy or [('1. Small', 'For the tests.', 'Height: 1 m', 'DoF: 20', 'Robot A'),
                            ('2. Large', 'For the tests.', 'Height: 2 m', 'DoF: 40', "Notional 'B'")]:
        sheet.append(list(row))
    parts = book.create_sheet('Components & Materials')
    parts.append(['COMPONENT & MATERIAL BREAKDOWN'])
    parts.append([f'COVERAGE: {coverage} leaf component assemblies'])
    parts.append(HEADER)
    for values in ROWS + [GRAND]:
        cells = list(values[:13]) + ['', 'Yes', values[13], '[S01]']
        parts.append(cells)
        r = parts.max_row
        fill = FILLS.get(values[13])
        for (row, heading), value in changes.items():
            if row != r:
                continue
            if isinstance(value, tuple) and value[0] == 'fill':
                fill = value[1]
            else:
                parts.cell(row=r, column=COLUMN[heading], value=value)
        if fill:
            parts.cell(row=r, column=COLUMN['Component']).fill = PatternFill('solid', fgColor=fill)

    sheet = book.create_sheet('Material Composition per Cat')
    sheet.append(['MATERIAL COMPOSITION PER CATEGORY'])
    sheet.append(['NOTE: for the tests.'])
    sheet.append(MATERIAL_HEADER)
    for values in MATERIAL_ROWS:
        sheet.append(list(values[:4]) + ['For the tests.'] + list(values[4:]) + ['[S01]'])
        r = sheet.max_row
        fill = FILLS.get(values[8])
        for (row, heading), value in (material_changes or {}).items():
            if row != r:
                continue
            if isinstance(value, tuple) and value[0] == 'fill':
                fill = value[1]
            else:
                sheet.cell(row=r, column=MATERIAL_COLUMN[heading], value=value)
        sheet.cell(row=r, column=MATERIAL_COLUMN['Material Family']).fill = PatternFill('solid', fgColor=fill)
    book.save(path)


class Small(unittest.TestCase):
    """Builds, reads and checks the small workbook; each test changes one thing."""

    def faults(self, changes=None, **options):
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, 'small.xlsx')
            build(path, changes, **options)
            self.book = schemas.read_workbook(path, Params())
        return schemas.check_workbook(self.book, Params())

    def assertOneFault(self, faults, row, words):
        self.assertEqual(len(faults), 1, [f.text for f in faults])
        self.assertEqual(faults[0].row, row)
        self.assertIn(words, faults[0].text)


class TestReading(Small):

    def test_the_small_workbook_is_read_without_fault(self):
        self.assertEqual(self.faults(), [])
        self.assertEqual([c.number for c in self.book.categories], [1, 2])
        self.assertEqual([len(pl.parts) for pl in self.book.parts_lists.values()], [2, 1])
        small = self.book.parts_lists[1]
        self.assertEqual((small.sum_row, small.stated_row, small.difference_row), (6, 7, 8))
        self.assertEqual(small.parts[0].total_kg, (2, 4, 6))
        self.assertEqual(self.book.parts_lists[2].difference_share, 0.2)

    def test_the_notional_anchor_is_recognised(self):
        self.faults()
        self.assertTrue(schemas.is_notional(self.book.categories[1].anchor))
        self.assertFalse(schemas.is_notional(self.book.categories[0].anchor))


class TestFaults(Small):
    """Each check, seen to catch the fault it guards and nothing else."""

    def test_a_minimum_above_the_mode(self):
        faults = self.faults({(4, 'Per-Unit Min (kg)'): 2.5, (4, 'Total Min (kg)'): 5,
                              (6, 'Total Min (kg)'): 5.5, (13, 'Total Min (kg)'): 9.5})
        self.assertEqual(len(faults), 2, [f.text for f in faults])
        self.assertTrue(all(f.row == 4 and 'not in order' in f.text for f in faults))

    def test_a_mass_of_all_units_that_is_not_the_number_times_one(self):
        faults = self.faults({(4, 'Total Max (kg)'): 7, (6, 'Total Max (kg)'): 8.5, (13, 'Total Max (kg)'): 16.5})
        self.assertOneFault(faults, 4, 'is not the number 2 times the mass of one, 3 kg')

    def test_a_missing_mass(self):
        faults = self.faults({(5, 'Per-Unit Mode (kg)'): D})
        self.assertOneFault(faults, 5, 'the mode of the mass of one unit is missing')

    def test_a_sum_row_that_disagrees_with_its_parts(self):
        # The materials' unallocated row is measured against the sum row, so it moves with it.
        faults = self.faults({(6, 'Total Mode (kg)'): 5.5, (8, 'Total Mode (kg)'): 0.5, (13, 'Total Mode (kg)'): 11.5,
                              (8, 'Total Max (kg)'): 0.1}, material_changes={(7, 'Mode (kg)'): 0.5})
        self.assertOneFault(faults, 6, 'the mode of the parts adds up to 5 kg, but the sum row says 5.5 kg')

    def test_a_sum_row_with_the_wrong_number_of_units(self):
        faults = self.faults({(6, 'Qty'): 4, (13, 'Qty'): 8})
        self.assertOneFault(faults, 6, 'the parts add up to 3 units, but the sum row says 4')

    def test_a_difference_row_that_disagrees(self):
        faults = self.faults({(12, 'Total Mode (kg)'): 2})
        self.assertOneFault(faults, 12, 'the difference row says 2 kg')

    def test_a_difference_share_that_disagrees(self):
        faults = self.faults({(12, 'Total Max (kg)'): 0.25})
        self.assertOneFault(faults, 12, 'the difference row says 25.00% of the stated mass')

    def test_a_row_coloured_against_its_evidence(self):
        faults = self.faults({(5, 'Component'): ('fill', 'FFFCE4D6')})
        self.assertOneFault(faults, 5, 'its evidence, Direct, is coloured without fill')

    def test_an_evidence_that_is_not_in_the_key(self):
        faults = self.faults({(5, 'Evidence'): 'Measured'})
        self.assertOneFault(faults, 5, "the evidence 'Measured' is not one of")

    def test_a_last_row_that_disagrees_with_the_sum_rows(self):
        faults = self.faults({(13, 'Total Max (kg)'): 15})
        self.assertOneFault(faults, 13, 'the maximum of the sum rows adds up to 15.5 kg')

    def test_a_stated_number_of_parts_that_disagrees(self):
        faults = self.faults(coverage=4)
        self.assertOneFault(faults, 0, 'states 4 parts, but the sheet lists 3')

    def test_a_category_named_differently_in_the_two_sheets(self):
        # Once for the parts and once for the materials: both sheets name the category.
        faults = self.faults(taxonomy=[('1. Smaller', 'x', 'x', 'x', 'Robot A'),
                                       ('2. Large', 'x', 'x', 'x', "Notional 'B'")])
        self.assertEqual(sorted(f.sheet for f in faults), ['Components & Materials', 'Material Composition per Cat'])
        self.assertTrue(all("category 1 is called 'Smaller'" in f.text for f in faults))

    def test_a_category_without_a_parts_list(self):
        faults = self.faults(taxonomy=[('1. Small', 'x', 'x', 'x', 'A'), ('2. Large', 'x', 'x', 'x', 'B'),
                                       ('3. Extra', 'x', 'x', 'x', 'C')])
        self.assertEqual(len(faults), 2, [f.text for f in faults])
        self.assertIn('category 3 is in this sheet but has no parts list', faults[0].text)
        self.assertIn('category 3 is in this sheet but has no materials', faults[1].text)


class TestPages(Small):

    def test_every_page_is_drawn_also_with_a_fault(self):
        faults = self.faults({(4, 'Total Max (kg)'): 7, (6, 'Total Max (kg)'): 8.5, (13, 'Total Max (kg)'): 16.5})
        params = Params()
        with tempfile.TemporaryDirectory() as folder:
            reports = {1: '01_category_small_components.md', 2: None}
            pages = [('overview_1.png', schemas.overview_flow(None, None, params)),
                     ('overview_2.png', schemas.overview_types(self.book, faults, reports, None, params)),
                     ('notes.png', schemas.notes_page())]
            pages += [(f'faults_{k}.png', f) for k, f in enumerate(schemas.fault_pages(faults), 1)]
            schemas.draw_screen_pages(pages, folder, params)
            for category in self.book.categories:
                schemas.draw_category(category, self.book.parts_lists[category.number], reports[category.number],
                                      faults, os.path.join(folder, 'categories', schemas.file_name(category)),
                                      params)
            drawn = sorted(os.path.relpath(p, folder) for p in glob.glob(os.path.join(folder, '**', '*.png'),
                                                                         recursive=True))
        self.assertEqual(drawn[:2], ['categories/1_small.png', 'categories/2_large.png'])
        self.assertTrue({'faults_1.png', 'notes.png', 'overview_1.png', 'overview_2.png'} <= set(drawn))

    def test_the_faults_continue_on_the_next_page_and_each_page_is_a_screen(self):
        faults = [schemas.Fault('Sheet', k, f'a fault of the tests, number {k}, ' + 'long ' * 30) for k in range(40)]
        pages = schemas.fault_pages(faults)
        self.assertGreater(len(pages), 2)
        for figure in pages:
            self.assertEqual(tuple(figure.get_size_inches()), schemas.SCREEN_PAGE)
            schemas.plt.close(figure)

    def test_a_subsystem_split_in_the_sheet_is_shown_once(self):
        part = lambda row, subsystem: schemas.Part(row, subsystem, f'part {row}', '', 1, (1, 1, 1), (1, 1, 1),  # noqa: E731
                                                   '', 'Direct', None, '')
        groups = schemas.by_subsystem([part(1, 'Structure'), part(2, 'Energy'), part(3, 'Structure')])
        self.assertEqual([(name, [p.row for p in members]) for name, members in groups],
                         [('Structure', [1, 3]), ('Energy', [2])])

    def test_a_long_file_name_is_broken_after_underscores(self):
        lines = schemas.wrap_name('Report: ', '01_category_compact_research_developer_components_2040.md', 30)
        self.assertEqual(''.join(lines).replace('Report: ', ''), '01_category_compact_research_developer_'
                                                                  'components_2040.md')
        self.assertTrue(all(len(line) <= 30 for line in lines))
        self.assertTrue(all(line.endswith('_') for line in lines[:-1]))

    def test_the_overview_names_every_stage_of_the_project(self):
        files = sorted(os.path.basename(p) for p in glob.glob(os.path.join(ROOT, '[0-9][0-9]_*.py')))
        self.assertEqual([name for name, _ in schemas.STAGES], files)


class TestScreenPages(unittest.TestCase):
    """Matthias, 2026-10-29: "I told you the 01 figures are shit", "Also the overview acn be done nicer!!", then "Yes,
    build it": every page of 01 a screen's shape with its smallest text at least 18 pixels tall fit to a 1920 × 1080
    screen; plain words; robot type 1 as 02 draws it, every robot of the parts file in its table."""

    SCREEN, SMALLEST_PX = (1920, 1080), 18
    # Words of the project that a page must not show (his rule: no code names, no project jargon).
    PROJECT_WORDS = schemas.PROJECT_WORDS
    FAULT_PAGE_WORDS = tuple(w for w in PROJECT_WORDS if w != 'sheet ')   # the faults pages say where a fault is

    @classmethod
    def setUpClass(cls):
        params = Params()
        paths = [absolute(params.data.parts_taken_apart), absolute(params.data.register),
                 absolute(params.data.workbook)]
        if not all(os.path.exists(p) for p in paths):
            raise unittest.SkipTest('the parts file, the register or the workbook is not on disk')
        cls.params = params
        cls.apart = schemas.read_parts_taken_apart(paths[0])
        cls.register = schemas.read_register(paths[1])
        cls.book = schemas.read_workbook(paths[2], params)

    def pages(self) -> list:
        reports = {c.number: None for c in self.book.categories}
        return ([schemas.overview_flow(self.register, self.apart, self.params),
                 schemas.overview_types(self.book, [], reports, self.apart, self.params, self.register),
                 schemas.notes_page(),
                 schemas.type_one_robot_page(self.apart)] + schemas.type_one_robots_pages(self.apart))

    def test_every_page_is_readable_fit_to_a_screen(self):
        import matplotlib.image
        from matplotlib.text import Text
        with tempfile.TemporaryDirectory() as folder:
            for k, figure in enumerate(self.pages() + schemas.fault_pages([])):
                figure.canvas.draw()
                sizes = [t.get_fontsize() for t in figure.findobj(Text) if t.get_visible() and t.get_text().strip()]
                path = os.path.join(folder, f'{k}.png')
                figure.savefig(path, dpi=self.params.output.dpi, facecolor='white')
                schemas.plt.close(figure)
                height, width = matplotlib.image.imread(path).shape[:2]
                scale = min(self.SCREEN[0] / width, self.SCREEN[1] / height)
                self.assertEqual((width, height), (16 * self.params.output.dpi, 9 * self.params.output.dpi))
                self.assertGreaterEqual(min(sizes) * self.params.output.dpi / 72 * scale, self.SMALLEST_PX, msg=k)

    def test_no_project_words_on_the_overview_the_notes_and_robot_type_1s_pages(self):
        from matplotlib.text import Text
        for figure in self.pages():
            words = ' '.join(t.get_text() for t in figure.findobj(Text))
            schemas.plt.close(figure)
            for word in self.PROJECT_WORDS:
                self.assertNotIn(word, words)

    def test_no_project_words_on_the_faults_pages_but_sheet_names(self):
        from matplotlib.text import Text
        for figure in schemas.fault_pages([]):
            words = ' '.join(' '.join(t.get_text() for t in figure.findobj(Text)).split())
            schemas.plt.close(figure)
            for word in self.FAULT_PAGE_WORDS:
                self.assertNotIn(word, words)

    def test_the_table_of_robot_type_1_lists_every_robot_of_the_parts_file(self):
        rows = schemas.type_one_robots_rows(self.apart)
        self.assertEqual(sorted(r.name for r in rows), sorted(self.apart.robots))
        by_name = {r.name: r for r in rows}
        g1 = by_name['Unitree G1']
        self.assertEqual((g1.height, g1.mass, g1.joints, round(g1.battery_wh)), (1.32, (34, 35, 36), 23, 421))
        self.assertEqual(round(by_name['Noetix E1'].battery_wh), 480)   # 12 Wh per kg of its 40 kg
        self.assertEqual(by_name['PrimeBOT Q1'].mass, (9.75, 19.5, 29.25))

    def test_the_table_groups_the_robots_by_what_their_joints_rest_on_lightest_first(self):
        rows = schemas.type_one_robots_rows(self.apart)
        groups: dict = {}
        for r in rows:
            groups.setdefault(r.group, []).append(r)
        self.assertEqual(list(groups), list(schemas.JOINT_GROUPS))
        self.assertEqual([r.name for r in groups['weighed']], ['Unitree G1', 'Unitree G1+'])
        self.assertEqual({r.name for r in groups['stated']},
                         {'AgiBot X1', 'ROBOTIS AI Sapiens K1', 'ROBOTIS K0', 'EngineAI PM01',
                          'RoboParty Roboto Origin', 'RoboParty RP1'})
        self.assertEqual({r.name for r in groups['torques']},
                         {'Fourier N1', 'Menlo Asimov 1', 'Menlo Asimov 2', 'Booster T1', 'Booster K1 5 Ah',
                          'Booster K1 2 Ah', 'Unitree R1'})
        self.assertEqual(len(groups['shares']), 17)
        for members in groups.values():
            self.assertEqual([r.mass[1] for r in members], sorted(r.mass[1] for r in members))
        # The mark of joints drawn as a share of the robot's mass on the robots of that group and no other.
        self.assertTrue(all(r.shares == (r.group == 'shares') for r in rows))

    def test_each_robots_five_cells_say_where_its_numbers_come_from(self):
        by_name = {r.name: r for r in schemas.type_one_robots_rows(self.apart)}
        self.assertEqual(by_name['Unitree G1'].bases,
                         {'Mass': ('maker',), 'Joints': ('weighed', 'estimated'), 'Battery': ('third party',),
                          'Boards': ('weighed', 'estimated'), 'Frame and shell': ('estimated',)})
        self.assertEqual(by_name['AgiBot X1'].bases['Joints'], ('maker',))
        self.assertEqual(by_name['Booster T1'].bases['Battery'], ('third party', 'estimated'))
        self.assertEqual(by_name['Booster T1'].bases['Boards'], ('maker', 'estimated'))
        self.assertEqual(by_name['Noetix E1'].bases['Battery'], ('estimated',))
        self.assertEqual(by_name['Noetix E1'].bases['Mass'], ('third party',))
        self.assertEqual(by_name['PrimeBOT Q1'].bases['Mass'], ('estimated',))

    def test_a_rows_label_and_its_first_source_decide_where_it_comes_from(self):
        base = lambda label, sources: schemas.row_base(self.apart, label, sources)  # noqa: E731
        self.assertEqual(base('S', 'U1 "About 35kg"; U3 adds up to 35.115 kg'), 'maker')
        self.assertEqual(base('S', 'DM1: 1375 and 375 g'), 'maker')          # the actuators' maker
        self.assertEqual(base('S, E', 'HG "42 kg" ("Not Verified")'), 'third party')
        self.assertEqual(base('S', 'BT1b 10.5 Ah (third party)'), 'third party')   # BT1b, not the maker's BT1
        self.assertEqual(base('W, E', 'T1 slides 11, 43'), 'weighed')
        self.assertEqual(base('E', 'U1'), 'estimated')

    def test_every_stated_number_behind_the_cells_names_a_source_of_the_sheet_sources(self):
        for robot in self.apart.robots:
            for cell, rows in schemas.cell_rows(self.apart, robot).items():
                for label, sources in rows:
                    if label.startswith('S'):
                        self.assertIsNotNone(schemas.first_source(self.apart, sources), (robot, cell, sources))

    def test_the_table_pages_show_every_robot_once_and_no_text_overlaps_another(self):
        from matplotlib.text import Text
        pages = schemas.type_one_robots_pages(self.apart)
        shown = []
        for figure in pages:
            figure.canvas.draw()
            renderer = figure.canvas.get_renderer()
            texts = [t for t in figure.findobj(Text) if t.get_visible() and t.get_text().strip()]
            shown += [t.get_text() for t in texts if t.get_text() in self.apart.robots]
            boxes = [t.get_window_extent(renderer) for t in texts]
            page = figure.bbox
            for k, b in enumerate(boxes):
                self.assertTrue(page.x0 <= b.x0 and b.x1 <= page.x1 and page.y0 <= b.y0 and b.y1 <= page.y1,
                                texts[k].get_text())
                for j in range(k):
                    self.assertFalse(b.overlaps(boxes[j]), (texts[k].get_text(), texts[j].get_text()))
            schemas.plt.close(figure)
        self.assertEqual(sorted(shown), sorted(self.apart.robots))


class TestMaterialFaults(Small):
    """The checks of the sheet 'Material Composition per Cat', each seen to catch the fault it guards."""

    def test_a_materials_sum_row_that_disagrees(self):
        faults = self.faults(material_changes={(6, 'Max (kg)'): 8})
        self.assertOneFault(faults, 6, 'the maximum of the materials adds up to 7.5 kg, but the sum row says 8 kg')

    def test_a_share_that_disagrees(self):
        faults = self.faults(material_changes={(4, 'Share of Category Material Mode (%)'): 0.5})
        self.assertOneFault(faults, 4, 'Aluminum alloys: the share says 50.00%, but its mode is 60.00% of the sum')

    def test_an_unallocated_row_that_disagrees_with_the_parts(self):
        faults = self.faults(material_changes={(10, 'Mode (kg)'): 3})
        self.assertOneFault(faults, 10, "the unallocated row says 3 kg against the parts' sum")

    def test_an_unallocated_row_that_disagrees_with_the_stated_mass(self):
        faults = self.faults(material_changes={(10, 'Max (kg)'): 2})
        self.assertOneFault(faults, 10, 'the unallocated row says 2 kg against the stated mass')

    def test_a_material_out_of_order(self):
        faults = self.faults(material_changes={(5, 'Min (kg)'): 2.5, (6, 'Min (kg)'): 4.5})
        self.assertOneFault(faults, 5, 'Copper conductors: not in order')

    def test_a_material_coloured_against_its_evidence(self):
        faults = self.faults(material_changes={(4, 'Material Family'): ('fill', 'FFFFF2CC')})
        self.assertOneFault(faults, 4, 'its evidence, Estimate, is coloured FFFCE4D6')


class SmallLists(Small):
    """Builds the small workbook and its shared lists, and checks both."""

    def lists_faults(self, parts=None, materials=None, lists=None, splits=None, battery=None, cells=None,
                     rows=None):
        """The faults of the small workbook and its shared lists; with `battery`, also of its sheet Battery cells,
        checked against the small element recipes (with `cells` instead of CELLS when given); with `rows`, also of
        the sheet Rows of the recipes."""
        with tempfile.TemporaryDirectory() as folder:
            build(os.path.join(folder, 'small.xlsx'))
            build_lists(os.path.join(folder, 'lists.xlsx'), parts, materials, lists, splits, battery)
            build_recipes(os.path.join(folder, 'recipes.xlsx'), cells, recipe=RECIPE + NESTED, rows=rows)
            self.book = schemas.read_workbook(os.path.join(folder, 'small.xlsx'), Params())
            self.lists = schemas.read_shared_lists(os.path.join(folder, 'lists.xlsx'))
            self.recipes = schemas.read_element_recipes(os.path.join(folder, 'recipes.xlsx'))
        faults = schemas.check_workbook(self.book, Params()) + schemas.check_shared_lists(self.book, self.lists,
                                                                                          Params())
        if battery is not None:
            faults += schemas.check_battery_cells(self.book, self.lists, self.recipes, Params())
        if rows is not None:
            faults += schemas.check_material_recipes(self.book, self.lists, self.recipes, Params())
        return faults

    def assertFaultSays(self, faults, words):
        self.assertEqual(len(faults), 1, [f.text for f in faults])
        self.assertIn(words, faults[0].text)


class TestSharedLists(SmallLists):

    def test_the_small_lists_are_read_without_fault_and_add_up(self):
        self.assertEqual(self.lists_faults(), [])
        one, two = self.book.parts_lists[1], self.book.parts_lists[2]
        self.assertEqual(schemas.by_view(one, self.lists, 'group'), {'Joint drives': 4, 'Frame': 1})
        self.assertEqual(schemas.by_view(one, self.lists, 'region'), {'Legs': 4, 'Whole body': 1})
        self.assertEqual(schemas.by_material(self.book.materials_lists[1], self.lists), {'Aluminium': 3, 'Copper': 2})
        self.assertEqual(schemas.by_material(self.book.materials_lists[2], self.lists),
                         {'Battery materials': 4, 'Not allocated': 2})
        self.assertEqual(schemas.motors_of(one, self.lists), 2)
        self.assertEqual(schemas.motors_of(two, self.lists), 0)

    def test_a_split_counts_the_mass_of_one_unit_per_unit(self):
        parts = [(1, 4, 'Knee actuator', 'Joint drives', '', 'Legs 1; Torso 1', 2, 'Split.')] + PLACED_PARTS[1:]
        self.assertEqual(self.lists_faults(parts), [])
        regions = schemas.by_view(self.book.parts_lists[1], self.lists, 'region')
        self.assertEqual(regions, {'Legs': 2, 'Torso': 2, 'Whole body': 1})

    def test_a_part_not_placed(self):
        self.assertFaultSays(self.lists_faults(PLACED_PARTS[:1] + PLACED_PARTS[2:]), 'Frame (row 5')

    def test_a_part_placed_twice(self):
        self.assertFaultSays(self.lists_faults(PLACED_PARTS + PLACED_PARTS[:1]), 'placed a second time')

    def test_a_part_under_another_name(self):
        parts = [(1, 4, 'Knee motor', 'Joint drives', 'Legs', '', 2, '')] + PLACED_PARTS[1:]
        self.assertFaultSays(self.lists_faults(parts), "row 4 is 'Knee actuator' in the workbook, not 'Knee motor'")

    def test_a_part_under_another_robot_type(self):
        parts = [(2, 4, 'Knee actuator', 'Joint drives', 'Legs', '', 2, '')] + PLACED_PARTS[1:]
        self.assertFaultSays(self.lists_faults(parts), 'robot type 2, but the workbook has it in 1')

    def test_a_part_group_not_in_the_lists(self):
        parts = [(1, 4, 'Knee actuator', 'Motors', 'Legs', '', 2, '')] + PLACED_PARTS[1:]
        self.assertFaultSays(self.lists_faults(parts), "the part group 'Motors' is not in the sheet Lists")

    def test_a_body_region_not_in_the_lists(self):
        parts = [(1, 4, 'Knee actuator', 'Joint drives', 'Feet', '', 2, '')] + PLACED_PARTS[1:]
        self.assertFaultSays(self.lists_faults(parts), "the body region 'Feet' is not in the sheet Lists")

    def test_a_region_and_a_split_at_once(self):
        parts = [(1, 4, 'Knee actuator', 'Joint drives', 'Legs', 'Legs 2', 2, '')] + PLACED_PARTS[1:]
        self.assertFaultSays(self.lists_faults(parts), 'give either one body region or a split by units')

    def test_a_split_that_does_not_add_up(self):
        parts = [(1, 4, 'Knee actuator', 'Joint drives', '', 'Legs 1; Torso 2', 2, '')] + PLACED_PARTS[1:]
        self.assertFaultSays(self.lists_faults(parts), 'the split adds up to 3 units, but the part has 2')

    def test_more_motors_than_units(self):
        parts = [(1, 4, 'Knee actuator', 'Joint drives', 'Legs', '', 3, '')] + PLACED_PARTS[1:]
        self.assertFaultSays(self.lists_faults(parts), 'the number of motors must be a whole number from 0 to its 2')

    def test_a_material_name_not_placed(self):
        self.assertFaultSays(self.lists_faults(materials=PLACED_MATERIALS[:1] + PLACED_MATERIALS[2:]),
                             "'Copper conductors'")

    def test_a_shared_material_not_in_the_lists(self):
        materials = [('Aluminum alloys', 'Brass', '')] + PLACED_MATERIALS[1:]
        self.assertFaultSays(self.lists_faults(materials=materials), "the shared material 'Brass' is not")

    def test_a_material_placed_in_not_allocated(self):
        materials = [('Aluminum alloys', 'Not allocated', '')] + PLACED_MATERIALS[1:]
        self.assertFaultSays(self.lists_faults(materials=materials), "the shared material 'Not allocated'")

    def test_a_material_name_the_workbook_does_not_have(self):
        materials = PLACED_MATERIALS + [('Gold', 'Copper', '')]
        self.assertFaultSays(self.lists_faults(materials=materials), "'Gold' is not a material name of the workbook")

    def test_the_page_of_a_robot_type_is_drawn(self):
        self.lists_faults()
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, 'categories', '1_small_shared_lists.png')
            schemas.draw_shared_lists(self.book.categories[0], self.book.parts_lists[1],
                                      self.book.materials_lists[1], self.lists, [], path, Params())
            self.assertTrue(os.path.exists(path))


class TestSplits(SmallLists):
    """A material name split by an estimate (point 9 of SHARED_LISTS.md). In the small workbook, robot type 1's
    'Copper conductors' (1 / 2 / 3 kg) is split into electronics (10 / 20 / 30 %) and copper as the remainder."""

    LISTS = LISTS + [('Material', 'Electronics')]
    MATERIALS = [PLACED_MATERIALS[0], ('Copper conductors', 'Split', 'Point 9.'), PLACED_MATERIALS[2]]
    SPLIT = [(1, 'Copper conductors', 'Electronics', 0.1, 0.2, 0.3, '', 'E', 'Basis.', 'Decided.'),
             (1, 'Copper conductors', 'Copper', None, None, None, 'Yes', 'E', 'The rest.', 'Decided.')]

    def split_faults(self, splits=None, materials=None):
        return self.lists_faults(materials=self.MATERIALS if materials is None else materials, lists=self.LISTS,
                                 splits=self.SPLIT if splits is None else splits)

    def test_a_split_is_read_and_adds_back_to_the_material(self):
        self.assertEqual(self.split_faults(), [])
        materials = self.book.materials_lists[1]
        totals = schemas.by_material(materials, self.lists)
        self.assertEqual(set(totals), {'Aluminium', 'Copper', 'Electronics'})
        self.assertAlmostEqual(totals['Electronics'], 0.4)
        self.assertAlmostEqual(totals['Copper'], 1.6)
        self.assertAlmostEqual(sum(totals.values()), 5)
        rows = schemas.material_rows(materials, self.lists)
        for name, expected in (('Electronics', (0.1, 0.4, 0.9)), ('Copper', (0.7, 1.6, 2.7))):
            for got, want in zip(rows[name][0], expected):
                self.assertAlmostEqual(got, want, msg=name)

    def test_a_split_without_a_remainder(self):
        self.assertFaultSays(self.split_faults(self.SPLIT[:1]), 'exactly one remainder, not 0')

    def test_a_split_with_two_remainders(self):
        two = [self.SPLIT[1], (1, 'Copper conductors', 'Electronics', None, None, None, 'Yes', 'E', '', '')]
        self.assertFaultSays(self.split_faults(two), 'exactly one remainder, not 2')

    def test_shares_out_of_order(self):
        rows = [(1, 'Copper conductors', 'Electronics', 0.3, 0.2, 0.4, '', 'E', '', ''), self.SPLIT[1]]
        self.assertFaultSays(self.split_faults(rows), 'the shares must be in order between 0 and 1')

    def test_shares_that_leave_the_remainder_below_zero(self):
        rows = [(1, 'Copper conductors', 'Electronics', 0.5, 0.8, 1.0, '', 'E', '', ''),
                (1, 'Copper conductors', 'Aluminium', 0.1, 0.1, 0.2, '', 'E', '', ''), self.SPLIT[1]]
        self.assertFaultSays(self.split_faults(rows), 'the largest shares add up to 1.2')

    def test_a_split_not_labelled_as_an_estimate(self):
        rows = [(1, 'Copper conductors', 'Electronics', 0.1, 0.2, 0.3, '', 'Direct', '', ''), self.SPLIT[1]]
        self.assertFaultSays(self.split_faults(rows), "labelled 'E', not 'Direct'")

    def test_a_piece_in_mixed(self):
        rows = [(1, 'Copper conductors', 'Mixed', 0.1, 0.2, 0.3, '', 'E', '', ''), self.SPLIT[1]]
        self.assertFaultSays(self.split_faults(rows), "the piece 'Mixed' is not a material")

    def test_a_split_whose_material_row_does_not_say_split(self):
        self.assertFaultSays(self.split_faults(materials=PLACED_MATERIALS), "does not say 'Split'")

    def test_a_material_row_saying_split_without_a_split(self):
        self.assertFaultSays(self.split_faults(splits=[]), "it says 'Split', but the sheet Splits has no rows")

    def test_the_pages_show_the_split(self):
        self.split_faults()
        with tempfile.TemporaryDirectory() as folder:
            for draw, name in ((schemas.draw_shared_lists, 'lists.png'), (schemas.draw_robot, 'robot.png')):
                path = os.path.join(folder, name)
                args = (self.book.categories[0], self.book.parts_lists[1], self.book.materials_lists[1], self.lists)
                draw(*args, [], path, Params()) if draw is schemas.draw_shared_lists else draw(*args, path, Params())
                self.assertTrue(os.path.exists(path))


class TestElementRecipes(unittest.TestCase):
    """The element recipes (data/researched/element_recipes.xlsx): Matthias, 2026-10-26, "I agree with the element
    recipe of the NMC and LFP cylindrical cells as proposed in E1.7i"."""

    def recipe_faults(self, compounds=None, recipe=None):
        with tempfile.TemporaryDirectory() as folder:
            build_recipes(os.path.join(folder, 'recipes.xlsx'), compounds=compounds, recipe=recipe)
            self.recipes = schemas.read_element_recipes(os.path.join(folder, 'recipes.xlsx'))
        return schemas.check_element_recipes(self.recipes, Params())

    def assertFaultSays(self, faults, words):
        self.assertEqual(len(faults), 1, [f.text for f in faults])
        self.assertIn(words, faults[0].text)

    def test_a_formula_gives_the_mass_shares_of_its_elements(self):
        shares = schemas.composition('LiFePO4')
        self.assertAlmostEqual(shares['Fe'], 55.845 / (6.94 + 55.845 + 30.974 + 4 * 15.999))
        self.assertAlmostEqual(sum(shares.values()), 1)
        self.assertAlmostEqual(schemas.composition('SiO0.64')['Si'], 28.085 / (28.085 + 0.64 * 15.999))
        self.assertEqual(schemas.composition('C 0.9; H 0.1'), {'C': 0.9, 'H': 0.1})

    def test_the_small_recipe_is_read_without_fault_and_adds_up(self):
        self.assertEqual(self.recipe_faults(), [])
        iron = schemas.element_shares(self.recipes, 'Test recipe', {'Can': 'Iron can'})
        plated = schemas.element_shares(self.recipes, 'Test recipe', {'Can': 'Plated can'})
        lipf6 = schemas.composition('LiPF6')
        for shares in (iron, plated):
            self.assertAlmostEqual(sum(shares.values()), 1)
            self.assertAlmostEqual(shares['C'], 0.7)
            self.assertAlmostEqual(shares['F'], 0.1 * lipf6['F'])
        self.assertAlmostEqual(iron['Fe'], 0.2)
        self.assertAlmostEqual(plated['Ni'], 0.02)
        self.assertAlmostEqual(plated['Fe'], 0.18)

    def test_a_composition_that_does_not_add_up(self):
        compounds = COMPOUNDS[:4] + [('Half and half', 'C 0.5; H 0.4', '')]
        self.assertFaultSays(self.recipe_faults(compounds=compounds), 'add up to 0.9, not 1')

    def test_a_composition_that_cannot_be_read(self):
        compounds = COMPOUNDS[:4] + [('Half and half', 'carbon and hydrogen', '')]
        self.assertFaultSays(self.recipe_faults(compounds=compounds), 'cannot be read')

    def test_an_element_that_has_no_atomic_mass(self):
        compounds = COMPOUNDS[:4] + [('Half and half', 'Xy2', '')]
        self.assertFaultSays(self.recipe_faults(compounds=compounds), "'Xy'")

    def test_a_compound_not_in_the_sheet_compounds(self):
        recipe = RECIPE[:6] + [('Test recipe', '', 'Rest', None, None, None, 'Yes', None, 'Graphene', '', 'E', '')]
        self.assertFaultSays(self.recipe_faults(recipe=recipe), "the compound 'Graphene' is not in the sheet")

    def test_a_group_with_two_remainders(self):
        recipe = RECIPE[:5] + [('Test recipe', '', 'Salt', None, None, None, 'Yes', None, 'LiPF6', '', 'E', ''),
                               RECIPE[6]]
        self.assertFaultSays(self.recipe_faults(recipe=recipe), 'exactly one remainder, not 2')

    def test_alternatives_whose_probabilities_do_not_add_up_to_one(self):
        recipe = RECIPE[:1] + [RECIPE[1][:7] + (0.6,) + RECIPE[1][8:]] + RECIPE[2:]
        self.assertFaultSays(self.recipe_faults(recipe=recipe), 'probabilities add up to 1.1')

    def test_shares_out_of_order(self):
        recipe = RECIPE[:5] + [('Test recipe', '', 'Salt', 0.15, 0.1, 0.2, '', None, 'LiPF6', '', 'E', '')] + RECIPE[6:]
        self.assertFaultSays(self.recipe_faults(recipe=recipe), 'in order between 0 and 1')

    def test_shares_that_leave_the_remainder_below_zero(self):
        recipe = RECIPE[:5] + [('Test recipe', '', 'Salt', 0.5, 0.7, 0.9, '', None, 'LiPF6', '', 'E', '')] + RECIPE[6:]
        self.assertFaultSays(self.recipe_faults(recipe=recipe), 'the largest shares add up to 1.1')

    def test_an_item_within_an_item_that_does_not_exist(self):
        recipe = RECIPE + [('Test recipe', 'Lid', 'Rivet', 0.1, 0.1, 0.1, '', None, 'Iron', '', 'E', '')]
        self.assertFaultSays(self.recipe_faults(recipe=recipe), "within 'Lid', which is not an item")

    def test_a_recipe_within_a_recipe(self):
        self.assertEqual(self.recipe_faults(recipe=RECIPE + NESTED), [])
        shares = schemas.element_shares(self.recipes, 'Outer')
        self.assertAlmostEqual(shares['Fe'], 0.4)
        self.assertAlmostEqual(shares['C'], 0.6 * (0.5 + 0.5 * 0.5))
        self.assertAlmostEqual(shares['H'], 0.6 * 0.5 * 0.5)
        self.assertAlmostEqual(sum(shares.values()), 1)

    def test_recipes_that_contain_each_other(self):
        recipe = RECIPE + NESTED[:1] + [('Inner', '', 'Mix', None, None, None, 'Yes', None, 'Outer', '', 'E', '')] + \
            NESTED[2:]
        self.assertFaultSays(self.recipe_faults(recipe=recipe), 'contain themselves')

    def test_a_name_that_is_both_a_compound_and_a_recipe(self):
        recipe = RECIPE + [('Carbon', '', 'Coal', None, None, None, 'Yes', None, 'Iron', '', 'E', '')]
        self.assertFaultSays(self.recipe_faults(recipe=recipe), "'Carbon' is both a compound and a recipe")


class TestMaterialRecipes(SmallLists):
    """Which recipe each workbook row of a shared material follows (sheet Rows of element_recipes.xlsx): Matthias,
    2026-10-26, "Yes, enter the plastics recipe now as laid out"."""

    ROWS = [(1, 'Copper conductors', 'Copper', 'Outer', 'For the tests.')]
    REST = [battery_row('Rest', 'Copper', share=(0.1, 0.2, 0.3)), battery_row('Rest', 'Not allocated', remainder='Yes')]

    def test_every_piece_has_its_recipe(self):
        self.assertEqual(self.lists_faults(rows=self.ROWS), [])

    def test_a_piece_of_a_covered_material_without_a_row(self):
        faults = self.lists_faults(battery=BATTERY[:2] + self.REST, rows=self.ROWS)
        self.assertFaultSays(faults, "robot type 2: 'Battery materials' in Copper has no row in the sheet Rows")

    def test_a_row_for_a_piece_that_does_not_exist(self):
        rows = self.ROWS + [(1, 'Gold', 'Copper', 'Outer', '')]
        self.assertFaultSays(self.lists_faults(rows=rows), "robot type 1 has no piece 'Gold' in Copper")

    def test_a_row_naming_a_recipe_that_does_not_exist(self):
        rows = [(1, 'Copper conductors', 'Copper', 'Bronze', '')]
        self.assertFaultSays(self.lists_faults(rows=rows), "the recipe 'Bronze' is not in the sheet Recipe")

    def test_a_piece_with_two_rows(self):
        self.assertFaultSays(self.lists_faults(rows=self.ROWS * 2), 'has a second row')


class TestBatteryCells(SmallLists):
    """The battery as cells only, counted from each pack's voltage and capacity: Matthias, 2026-10-26, "The number of
    cells, which also defines the cell mass depends on the voltage, which these batteries have and the robots use"."""

    def test_the_cells_are_counted_and_the_rest_goes_to_not_allocated(self):
        self.assertEqual(self.lists_faults(battery=BATTERY), [])
        materials = self.book.materials_lists[2]
        cells = schemas.battery_cells(materials, self.lists, self.recipes)
        for got, want in zip(cells, (0.9, 0.95, 1.0)):
            self.assertAlmostEqual(got, want)
        totals = schemas.by_material(materials, self.lists, self.recipes)
        self.assertAlmostEqual(totals['Battery materials'], 0.95)
        self.assertAlmostEqual(totals['Not allocated'], 2 + 4 - 0.95)
        self.assertAlmostEqual(sum(totals.values()), 6)
        # Without the recipes, the workbook's row as it is.
        self.assertEqual(schemas.by_material(materials, self.lists), {'Battery materials': 4, 'Not allocated': 2})

    def test_a_rest_split_among_materials_adds_back_to_the_row(self):
        rest = [battery_row('Rest', 'Copper', share=(0.1, 0.2, 0.3)), battery_row('Rest', 'Aluminium', remainder='Yes')]
        self.assertEqual(self.lists_faults(battery=BATTERY[:2] + rest), [])
        totals = schemas.by_material(self.book.materials_lists[2], self.lists, self.recipes)
        self.assertAlmostEqual(totals['Copper'], 0.2 * (4 - 0.95))
        self.assertAlmostEqual(totals['Aluminium'], 0.8 * (4 - 0.95))
        self.assertAlmostEqual(sum(totals.values()), 6)
        rows = schemas.material_rows(self.book.materials_lists[2], self.lists, self.recipes)
        for got, want in zip(rows['Copper'][0], (0.1 * (3 - 1.0), 0.2 * (4 - 0.95), 0.3 * (5 - 0.9))):
            self.assertAlmostEqual(got, want)

    def test_a_cell_not_in_the_sheet_cells(self):
        battery = [battery_row('Build', 'Cell Z', 10, 2, 1, 36, 4)] + BATTERY[2:]
        self.assertFaultSays(self.lists_faults(battery=battery), "the cell 'Cell Z' is not in the sheet Cells")

    def test_cells_in_parallel_that_miss_the_capacity(self):
        battery = [battery_row('Build', 'Cell A', 10, 1, 1, 36, 4)] + BATTERY[2:]
        self.assertFaultSays(self.lists_faults(battery=battery), 'is 50 % off the pack\'s 4 Ah')

    def test_cells_in_series_that_miss_the_voltage(self):
        battery = [battery_row('Build', 'Cell A', 12, 2, 1, 36, 4)] + BATTERY[2:]
        self.assertFaultSays(self.lists_faults(battery=battery), "more than one cell's voltage from the pack's 36 V")

    def test_cells_heavier_than_the_pack_allows(self):
        battery = [battery_row('Build', 'Cell A', 10, 2, 1, 36, 4, pack_kg=1.0)] + BATTERY[2:]
        self.assertFaultSays(self.lists_faults(battery=battery), "more than 85 % of the pack's 1 kg")

    def test_a_build_with_neither_cells_in_series_nor_an_energy(self):
        battery = [battery_row('Build', 'Cell A')] + BATTERY[2:]
        self.assertFaultSays(self.lists_faults(battery=battery), 'give the cells in series, in parallel and the packs')

    def test_a_build_from_the_energy_counts_whole_cells(self):
        battery = [battery_row('Build', 'Cell A', kwh=0.1)] + BATTERY[2:]   # 0.1 kWh / 7.2 Wh = 13.9: 14 cells
        self.assertEqual(self.lists_faults(battery=battery), [])
        self.assertAlmostEqual(schemas.battery_cells(self.book.materials_lists[2], self.lists, self.recipes)[1], 0.7)

    def test_a_rest_without_a_remainder(self):
        battery = BATTERY[:2] + [battery_row('Rest', 'Not allocated')]
        self.assertFaultSays(self.lists_faults(battery=battery), 'exactly one remainder, not 0')

    def test_a_rest_to_a_material_that_can_fall_below_zero(self):
        cells = [('Cell A', 'NMC', '18650', 3.6, 2.0, 200, None, 'Test recipe', '', 'Test')] + CELLS[1:]
        rest = [battery_row('Rest', 'Copper', remainder='Yes')]
        self.assertFaultSays(self.lists_faults(battery=BATTERY[:2] + rest, cells=cells), 'can fall below zero')

    def test_a_battery_row_both_kept_and_built(self):
        battery = BATTERY + [battery_row('Kept', '')]
        self.assertFaultSays(self.lists_faults(battery=battery), 'both kept and built')

    def test_a_kept_battery_row_names_the_recipe_of_its_cells(self):
        # Matthias, 2026-10-26: "Robot types 2 and 6 both get NMC cells with a liquid electrolyte, using the car file's
        # NMC recipe" (EVIDENCE.md E1.7e). A kept row's cells follow the recipe it names, for 02_composition.py.
        self.assertEqual(self.lists_faults(battery=[battery_row('Kept', 'Test recipe')]), [])
        for name in ('', 'Bronze'):
            faults = self.lists_faults(battery=[battery_row('Kept', name)])
            self.assertFaultSays(faults, f"a kept battery row names the recipe its cells follow; '{name}' is not in "
                                         'the sheet Recipe')

    def test_a_battery_row_without_any_row_in_the_sheet(self):
        self.assertFaultSays(self.lists_faults(battery=[]), "no row in the sheet Battery cells")

    def test_a_build_not_labelled_as_an_estimate(self):
        battery = [battery_row('Build', 'Cell A', 10, 2, 1, 36, 4, label='Direct')] + BATTERY[2:]
        self.assertFaultSays(self.lists_faults(battery=battery), "labelled 'E', not 'Direct'")

    def test_the_robot_page_shows_the_cells(self):
        self.lists_faults(battery=BATTERY)
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, '2_robot.png')
            schemas.draw_robot(self.book.categories[1], self.book.parts_lists[2], self.book.materials_lists[2],
                               self.lists, path, Params(), recipes=self.recipes)
            self.assertTrue(os.path.exists(path))


class TestRobotPage(SmallLists):
    """The one-view schema of a robot type: Matthias, 2026-10-26, "I want theat one can understand the robot and
    its compostion and parts in one vew"."""

    def category(self, anchor, height_and_mass):
        return schemas.Category(1, 'Small', '', height_and_mass, '', anchor, 2)

    def test_the_height_of_the_anchor_robot_where_the_workbook_gives_it(self):
        height, note = schemas.robot_height(self.category('Unitree G1 (China)\n1.32 m | 35.0 kg', 'Height: 1–2 m'))
        self.assertEqual((height, note), (1.32, ''))

    def test_otherwise_the_middle_of_the_category_range_said_on_the_page(self):
        height, note = schemas.robot_height(self.category('Atlas (US)', 'Height: ~1.73–1.85 m\nMass: 73–90 kg'))
        self.assertAlmostEqual(height, 1.79)
        self.assertIn('middle of the category\'s height range, 1.73–1.85 m', note)

    def test_without_any_height_the_tallest(self):
        height, note = schemas.robot_height(self.category('Robot', 'Mass: 50 kg'))
        self.assertEqual(height, schemas.TALLEST)
        self.assertTrue(note)

    def test_labels_keep_their_order_and_their_distance(self):
        placed = schemas.spread([5.0, 1.0, 1.1, 1.15], 0.5, 1.0, 6.0)
        self.assertEqual(placed[1:], [1.0, 1.5, 2.0])
        self.assertEqual(placed[0], 5.0)
        crowded = schemas.spread([5.9, 5.95, 6.0], 0.5, 1.0, 6.0)
        self.assertEqual(crowded, [5.0, 5.5, 6.0])

    def test_the_motors_of_a_split_part_follow_its_units(self):
        parts = [(1, 4, 'Knee actuator', 'Joint drives', '', 'Legs 1; Torso 1', 2, 'Split.')] + PLACED_PARTS[1:]
        self.lists_faults(parts)
        self.assertEqual(schemas.motors_by_region(self.book.parts_lists[1], self.lists), {'Legs': 1, 'Torso': 1})

    def test_bounds_add_up_the_minimums_and_the_maximums(self):
        self.assertEqual(schemas.bounds([(1, 2, 3), (0.5, 1, 1.5), (None, 4, None)]), (1.5, 4.5))

    def test_the_page_is_drawn(self):
        self.lists_faults()
        with tempfile.TemporaryDirectory() as folder:
            for n in (1, 2):
                path = os.path.join(folder, 'categories', f'{n}_robot.png')
                schemas.draw_robot(self.book.categories[n - 1], self.book.parts_lists[n],
                                   self.book.materials_lists[n], self.lists, path, Params())
                self.assertTrue(os.path.exists(path))


class TestTheWorkbook(unittest.TestCase):
    """The real workbook, as 01 reads it; skipped where the workbook is not on disk."""

    def test_six_categories_68_parts_and_no_fault(self):
        params = Params()
        path = absolute(params.data.workbook)
        if not os.path.exists(path):
            self.skipTest(f'{params.data.workbook} is not on disk')
        book = schemas.read_workbook(path, params)
        self.assertEqual([c.number for c in book.categories], [1, 2, 3, 4, 5, 6])
        self.assertEqual(sum(len(pl.parts) for pl in book.parts_lists.values()), 68)
        self.assertEqual(schemas.check_workbook(book, params), [])
        # The two numbers the overview leans on, read by hand from the workbook (rows 77 and 78).
        heavy = book.parts_lists[5]
        self.assertEqual((heavy.sum_kg[1], heavy.stated_mass_kg), (103.4, 90.0))
        self.assertTrue(re.match(r'Notional', book.categories[5].anchor))
        self.assertEqual(sum(len(ml.materials) for ml in book.materials_lists.values()), 59)

    def test_the_shared_lists_place_everything_and_add_up(self):
        params = Params()
        if not (os.path.exists(absolute(params.data.workbook)) and os.path.exists(absolute(params.data.shared_lists))):
            self.skipTest('the workbook or the shared lists are not on disk')
        book = schemas.read_workbook(absolute(params.data.workbook), params)
        lists = schemas.read_shared_lists(absolute(params.data.shared_lists))
        self.assertEqual(schemas.check_shared_lists(book, lists, params), [])
        self.assertEqual((len(lists.parts), len(lists.materials)), (68, 35))
        for n, parts_list in book.parts_lists.items():
            mode = parts_list.sum_kg[1]
            for totals in (schemas.by_view(parts_list, lists, 'group'), schemas.by_view(parts_list, lists, 'region'),
                           schemas.by_material(book.materials_lists[n], lists)):
                self.assertAlmostEqual(sum(totals.values()), mode, places=6, msg=f'robot type {n}')
        self.assertEqual([schemas.motors_of(book.parts_lists[n], lists) for n in range(1, 7)], [23, 31, 52, 30, 28, 14])
        # Point 9: robot type 3's 'Electrical & magnetics group' (4.2 kg at the mode), split by the shares of robot
        # types 2, 4 and 5: magnets 30 %, electronics 20 %, copper the remaining 50 %.
        three = schemas.by_material(book.materials_lists[3], lists)
        for name, mode in (('Magnets (NdFeB)', 1.26), ('Electronics', 0.84), ('Copper', 2.1), ('Mixed', 2.3)):
            self.assertAlmostEqual(three[name], mode, places=6, msg=name)

    def test_the_element_recipes_reproduce_what_matthias_agreed(self):
        params = Params()
        if not os.path.exists(absolute(params.data.element_recipes)):
            self.skipTest(f'{params.data.element_recipes} is not on disk')
        recipes = schemas.read_element_recipes(absolute(params.data.element_recipes))
        self.assertEqual(schemas.check_element_recipes(recipes, params), [])
        # EVIDENCE.md E1.7i, the recipe at the modes, % of the cell, for each can.
        expected = {('NMC cylindrical', 'Nickel-plated steel'): {'C': 29.68, 'Ni': 19.34, 'O': 14.69, 'Fe': 16.45,
                                                                 'Cu': 8.30, 'Al': 3.68, 'Li': 2.57, 'Co': 1.35},
                    ('NMC cylindrical', 'Stainless steel 304'): {'Ni': 20.40, 'Fe': 12.39, 'Cr': 3.04},
                    ('LFP cylindrical', 'Nickel-plated steel'): {'Fe': 29.97, 'C': 25.40, 'O': 20.10, 'P': 7.83,
                                                                 'Li': 1.75, 'Ni': 0.30},
                    ('LFP cylindrical', 'Stainless steel 304'): {'Fe': 25.92, 'Ni': 1.36, 'Cr': 3.04}}
        # The table was computed from the car file's cathode shares as EVIDENCE.md prints them (three decimals); the
        # workbook holds them to five, which moves a share by up to 0.02 percentage points. The compounds written as
        # mass shares are rounded to five decimals, so the elements add up to 1 within 0.00001.
        for (recipe, can), elements in expected.items():
            shares = schemas.element_shares(recipes, recipe, {'Can and cap': can})
            self.assertAlmostEqual(sum(shares.values()), 1, places=5)
            for element, percent in elements.items():
                self.assertAlmostEqual(shares[element] * 100, percent, delta=0.03, msg=f'{recipe}, {can}, {element}')

    def test_the_battery_cells_reproduce_what_matthias_agreed(self):
        params = Params()
        paths = [absolute(p) for p in (params.data.workbook, params.data.shared_lists, params.data.element_recipes)]
        if not all(os.path.exists(p) for p in paths):
            self.skipTest('the workbook, the shared lists or the element recipes are not on disk')
        book = schemas.read_workbook(paths[0], params)
        lists = schemas.read_shared_lists(paths[1])
        recipes = schemas.read_element_recipes(paths[2])
        self.assertEqual(schemas.check_battery_cells(book, lists, recipes, params), [])
        # EVIDENCE.md E1.7j and E1.7k: min / mode (the median of the builds) / max, kg.
        for n, expected in ((1, (1.74, 1.81, 1.95)), (3, (7.20, 7.335, 7.335)), (4, (2.60, 2.755, 2.92)),
                            (5, (9.80, 10.39, 11.03))):
            for got, want in zip(schemas.battery_cells(book.materials_lists[n], lists, recipes), expected):
                self.assertAlmostEqual(got, want, delta=0.006, msg=f'robot type {n}')
        for n in (2, 6):
            self.assertIsNone(schemas.battery_cells(book.materials_lists[n], lists, recipes))
            # E1.7e: their kept rows' cells follow the NMC recipe.
            self.assertEqual([(b.kind, b.name) for b in lists.battery if b.robot == n], [('Kept', 'NMC cylindrical')])
        # Cells only change where the kilograms sit, never how many there are.
        for n, parts_list in book.parts_lists.items():
            totals = schemas.by_material(book.materials_lists[n], lists, recipes)
            self.assertAlmostEqual(sum(totals.values()), parts_list.sum_kg[1], places=6, msg=f'robot type {n}')
        one = schemas.by_material(book.materials_lists[1], lists, recipes)
        rest = 2.5 - 1.81
        self.assertAlmostEqual(one['Battery materials'], 1.81, delta=0.006)
        self.assertAlmostEqual(one['Electronics'], 0.5 + 0.165 * rest, delta=0.006)

    def test_the_plastics_recipes_reproduce_what_matthias_agreed(self):
        params = Params()
        paths = [absolute(p) for p in (params.data.workbook, params.data.shared_lists, params.data.element_recipes)]
        if not all(os.path.exists(p) for p in paths):
            self.skipTest('the workbook, the shared lists or the element recipes are not on disk')
        book = schemas.read_workbook(paths[0], params)
        lists = schemas.read_shared_lists(paths[1])
        recipes = schemas.read_element_recipes(paths[2])
        self.assertEqual(schemas.check_material_recipes(book, lists, recipes, params), [])
        # EVIDENCE.md E1.8c with the six values of E1.8d, % of each row at the modes (the TPU's soft segment taken as
        # the first alternative, polyester, as in the sandbox of V25). One decimal, so within 0.06.
        expected = {'Plastics, robot type 1, Engineering polymers': {'C': 75.2, 'O': 11.0, 'H': 7.8, 'N': 6.0},
                    'PC/ABS': {'C': 78.8, 'O': 13.2, 'H': 6.3, 'N': 1.7},
                    'Plastics, robot type 2, Engineering plastics': {'C': 71.8, 'N': 2.7, 'Si': 1.9, 'Ca': 1.1},
                    'Plastics, robot type 3, Engineering plastics': {'C': 61.1, 'O': 19.0, 'N': 5.1, 'P': 0.5},
                    'Plastics, robot type 4, Polymers & elastomers': {'C': 62.3, 'O': 21.0, 'N': 4.8, 'Si': 2.5},
                    'Plastics, robot type 5, Composites & polymers': {'C': 64.7, 'N': 4.6, 'Si': 3.7, 'Ca': 2.2},
                    'Plastics, robot type 5, Sealing elastomers': {'C': 44.7, 'O': 26.7, 'Si': 19.8, 'N': 2.0},
                    'Plastics, robot type 6, Reinforced polymers': {'C': 64.7, 'N': 4.6, 'Al': 1.1},
                    'Plastics, robot type 6, Protective ballistic fibers': {'C': 78.1, 'H': 9.3, 'O': 6.7, 'N': 5.9}}
        for recipe, elements in expected.items():
            shares = schemas.element_shares(recipes, recipe)
            self.assertAlmostEqual(sum(shares.values()), 1, places=6, msg=recipe)
            for element, percent in elements.items():
                self.assertAlmostEqual(shares[element] * 100, percent, delta=0.06, msg=f'{recipe}, {element}')

    def test_the_fibre_composite_recipes_reproduce_what_matthias_agreed(self):
        params = Params()
        paths = [absolute(p) for p in (params.data.workbook, params.data.shared_lists, params.data.element_recipes)]
        if not all(os.path.exists(p) for p in paths):
            self.skipTest('the workbook, the shared lists or the element recipes are not on disk')
        book = schemas.read_workbook(paths[0], params)
        lists = schemas.read_shared_lists(paths[1])
        recipes = schemas.read_element_recipes(paths[2])
        self.assertEqual(schemas.check_element_recipes(recipes, params), [])
        self.assertEqual(schemas.check_material_recipes(book, lists, recipes, params), [])
        rows = {(r.robot, r.family): r.recipe for r in recipes.rows if r.shared == 'Fibre composites'}
        self.assertEqual(set(rows), {(1, 'Carbon-fiber composite'), (2, 'Carbon-fiber composite'), (3, 'Fiber composites')})
        # EVIDENCE.md E1.9c, % of each row at the modes, with each of the two epoxies. One decimal, so within 0.06.
        expected = {('CFRP', 'Aromatic amine'): {'C': 87.2, 'O': 6.6, 'H': 2.5, 'N': 2.4, 'S': 1.3},
                    ('CFRP', 'Cycloaliphatic amine'): {'C': 89.0, 'O': 5.4, 'H': 3.3, 'N': 2.4},
                    (rows[(3, 'Fiber composites')], 'Aromatic amine'): {'C': 54.1, 'O': 22.5, 'Si': 8.7, 'Ca': 5.1,
                                                                         'S': 1.2},
                    (rows[(3, 'Fiber composites')], 'Cycloaliphatic amine'): {'C': 55.7, 'O': 21.4, 'Si': 8.7,
                                                                               'H': 2.9}}
        self.assertEqual(rows[(1, 'Carbon-fiber composite')], 'CFRP')
        for (recipe, epoxy), elements in expected.items():
            shares = schemas.element_shares(recipes, recipe, {'Epoxy system': epoxy})
            self.assertAlmostEqual(sum(shares.values()), 1, places=6, msg=recipe)
            self.assertEqual('S' in shares, epoxy == 'Aromatic amine', msg=recipe)
            for element, percent in elements.items():
                self.assertAlmostEqual(shares[element] * 100, percent, delta=0.06, msg=f'{recipe}, {epoxy}, {element}')

    def test_the_aluminium_recipes_reproduce_what_matthias_agreed(self):
        # Matthias, 2026-10-27: "I agree with the aluminium recipe as proposed in E1.1c, including its choices.", then,
        # with its modes set again by his rule (standards and commercial labs first): "I agree with the modes of E1.1e
        # and E1.2d, including their choices." and "Yes, enter both".
        params = Params()
        paths = [absolute(p) for p in (params.data.workbook, params.data.shared_lists, params.data.element_recipes)]
        if not all(os.path.exists(p) for p in paths):
            self.skipTest('the workbook, the shared lists or the element recipes are not on disk')
        book = schemas.read_workbook(paths[0], params)
        lists = schemas.read_shared_lists(paths[1])
        recipes = schemas.read_element_recipes(paths[2])
        self.assertEqual(schemas.check_element_recipes(recipes, params), [])
        self.assertEqual(schemas.check_material_recipes(book, lists, recipes, params), [])
        rows = {(r.robot, r.family): r.recipe for r in recipes.rows if r.shared == 'Aluminium'}
        self.assertEqual(set(rows), {(n, 'Aluminum alloys') for n in range(1, 7)})
        # EVIDENCE.md E1.1e with ADC12 from its combined analyses (E1.1f), % of each robot type's aluminium at the modes,
        # with each die-cast alloy. Two decimals, so within 0.006.
        expected = {(1, 'A380'): {'Al': 92.94, 'Si': 2.18, 'Zn': 1.60, 'Cu': 1.28, 'Mg': 1.20, 'Fe': 0.41},
                    (1, 'ADC12'): {'Al': 92.72, 'Si': 2.66, 'Zn': 1.65, 'Mg': 1.24, 'Cu': 0.98, 'Fe': 0.39},
                    (2, 'A380'): {'Al': 90.70, 'Zn': 3.19, 'Si': 2.05, 'Cu': 1.66, 'Mg': 1.64, 'Fe': 0.36},
                    (2, 'ADC12'): {'Al': 90.48, 'Zn': 3.24, 'Si': 2.53, 'Mg': 1.67, 'Cu': 1.36, 'Fe': 0.33},
                    (3, 'A380'): {'Al': 92.25, 'Si': 2.14, 'Zn': 2.08, 'Cu': 1.40, 'Mg': 1.34, 'Fe': 0.40},
                    (3, 'ADC12'): {'Al': 92.03, 'Si': 2.62, 'Zn': 2.13, 'Mg': 1.37, 'Cu': 1.10, 'Fe': 0.37},
                    (4, 'A380'): {'Al': 91.38, 'Zn': 2.70, 'Si': 2.09, 'Cu': 1.54, 'Mg': 1.50, 'Fe': 0.37},
                    (4, 'ADC12'): {'Al': 91.17, 'Zn': 2.75, 'Si': 2.57, 'Mg': 1.54, 'Cu': 1.24, 'Fe': 0.35},
                    (5, 'A380'): {'Al': 90.52, 'Zn': 3.32, 'Si': 2.04, 'Cu': 1.69, 'Mg': 1.67, 'Fe': 0.35},
                    (5, 'ADC12'): {'Al': 90.30, 'Zn': 3.37, 'Si': 2.52, 'Mg': 1.70, 'Cu': 1.39, 'Fe': 0.32},
                    (6, 'A380'): {'Al': 91.07, 'Zn': 2.92, 'Si': 2.08, 'Cu': 1.60, 'Mg': 1.56, 'Fe': 0.37},
                    (6, 'ADC12'): {'Al': 90.86, 'Zn': 2.97, 'Si': 2.55, 'Mg': 1.60, 'Cu': 1.30, 'Fe': 0.34}}
        for (n, alloy), elements in expected.items():
            shares = schemas.element_shares(recipes, rows[(n, 'Aluminum alloys')], {'Alloy': alloy})
            self.assertAlmostEqual(sum(shares.values()), 1, places=9, msg=f'robot type {n}, {alloy}')
            self.assertEqual('Pb' in shares, alloy == 'ADC12', msg=f'robot type {n}, {alloy}')
            for element, percent in elements.items():
                self.assertAlmostEqual(shares[element] * 100, percent, delta=0.006,
                                       msg=f'robot type {n}, {alloy}, {element}')

    def test_the_magnesium_recipes_reproduce_what_matthias_agreed(self):
        # Matthias, 2026-10-27: "I agree with the modes of E1.1e and E1.2d, including their choices." and "Yes, enter
        # both: ... the magnesium recipe of E1.2c (its AM share, AM60B or AM50A half and half, robot type 3 as AZ91D
        # alone) with the modes of E1.2d."
        params = Params()
        paths = [absolute(p) for p in (params.data.workbook, params.data.shared_lists, params.data.element_recipes)]
        if not all(os.path.exists(p) for p in paths):
            self.skipTest('the workbook, the shared lists or the element recipes are not on disk')
        book = schemas.read_workbook(paths[0], params)
        lists = schemas.read_shared_lists(paths[1])
        recipes = schemas.read_element_recipes(paths[2])
        self.assertEqual(schemas.check_element_recipes(recipes, params), [])
        self.assertEqual(schemas.check_material_recipes(book, lists, recipes, params), [])
        rows = {(r.robot, r.family): r.recipe for r in recipes.rows if r.shared == 'Magnesium'}
        self.assertEqual(set(rows), {(n, 'Magnesium alloys') for n in (2, 3, 4, 6)})
        self.assertEqual(rows[(3, 'Magnesium alloys')], 'Magnesium AZ91D')
        # EVIDENCE.md E1.2d with the university-only elements combined (E1.2e), % of each robot type's magnesium at the
        # modes, with each AM alloy. Within 0.006.
        expected = {(2, 'AM60B'): {'Mg': 91.38, 'Al': 7.71, 'Zn': 0.463, 'Mn': 0.386},
                    (2, 'AM50A'): {'Mg': 92.00, 'Al': 7.24, 'Zn': 0.416, 'Mn': 0.300},
                    (3, 'AM60B'): {'Mg': 89.94, 'Al': 9.00, 'Zn': 0.700, 'Mn': 0.300},
                    (4, 'AM60B'): {'Mg': 91.42, 'Al': 7.68, 'Zn': 0.458, 'Mn': 0.388},
                    (4, 'AM50A'): {'Mg': 92.04, 'Al': 7.20, 'Zn': 0.410, 'Mn': 0.300},
                    (6, 'AM60B'): {'Mg': 91.62, 'Al': 7.50, 'Zn': 0.425, 'Mn': 0.400},
                    (6, 'AM50A'): {'Mg': 92.33, 'Al': 6.96, 'Zn': 0.370, 'Mn': 0.300}}
        for (n, alloy), elements in expected.items():
            shares = schemas.element_shares(recipes, rows[(n, 'Magnesium alloys')], {'AM alloy': alloy})
            self.assertAlmostEqual(sum(shares.values()), 1, places=9, msg=f'robot type {n}, {alloy}')
            for element, percent in elements.items():
                self.assertAlmostEqual(shares[element] * 100, percent, delta=0.006,
                                       msg=f'robot type {n}, {alloy}, {element}')


# Robot type 1 taken apart (data/researched/parts_taken_apart.xlsx): Matthias, 2026-10-27, "Yes. Show me the steps and
# then we move step by step through them", step 3: 01 checks the file; then, on "Yes, build points 1 to 8 as
# specified, tests first" (EVIDENCE.md E2.8e), one file for every robot of the type: a sheet Robots and a column Robot.
# The headings of each sheet, as the file has them.
APART = {
    'Robots': ['Robot', 'Maker', 'Type', 'Height (m)', 'Stated mass (kg)', 'Sources'],
    'Joints': ['Robot', 'Number', 'Joint', 'Joints in the robot', 'Module', 'Torque limit (N·m)',
               'Module mass, min (kg)', 'mode (kg)', 'max (kg)', 'Share of the joint modules', 'Cooling', 'Heat group',
               'Body region', 'Sources', 'Basis'],
    'MagnetGrades': ['Heat group', 'Grade class', 'Probability', 'Label', 'Basis'],
    'JointParts': ['Component', 'Small joint, weighed (g)', 'Small joints not weighed: times min', 'times mode',
                   'times max', 'Material (one material only)', 'Grade', 'Large joint: factor min', 'factor mode',
                   'factor max', 'Why the factor', 'Label', 'Sources'],
    'Parts': ['Robot', 'Part', 'Number in the robot', "Kind (the figure's row)", 'Body region',
              'Mass of one, min (kg)', 'mode (kg)', 'max (kg)', 'Material (one material only)', 'Grade',
              'Remainder of, or computed from', 'Label', 'Sources'],
    'Shares': ['Part or component', 'Material', 'Grade', 'Share min', 'Share mode', 'Share max', 'Remainder', 'Label',
               'Basis'],
    'Boards': ['Robot', 'Board', 'Number in the robot', "Kind (the figure's row)", 'Body region',
               'Area of one, min (cm2)', 'mode (cm2)', 'max (cm2)', 'Mass from',
               'Car board category (RAWCLICVehicleElectronics 04_)', 'Label', 'Sources'],
    'Sensors': ['Robot', 'Sensor', 'Number in the robot', "Kind (the figure's row)", 'Body region',
                'Mass of one, min (g)', 'mode (g)', 'max (g)', 'Label', 'Sources'],
    'SensorElements': ['Sensor', 'Car sensor type (RAWCLICVehicleElectronics 07_, sheet Sensor Details)',
                       'Units in one sensor'],
    'Constants': ['Robot', 'Quantity', 'Min', 'Mode', 'Max', 'Unit', 'Label', 'Source'],
    'Sources': ['Code', 'Source', 'Link', 'What it is'],
}
# The small file: two robots of type 1. Robot A as the G1 (its joints follow the weighed joint parts, no module
# masses): one large joint kind (2 hips) and one small (2 ankles), four parts, one board, one sensor. Robot B as the X1
# (module masses given): 2 large hips and 2 small arm joints, its joint modules and frame, one board, one sensor. Every
# constant 01 needs, the shared ones once for all.
SMALL_APART = {
    'Robots': [('Robot A', 'Maker A', 1, 1.32, 35, 'U1'), ('Robot B', 'Maker B', 1, 1.30, 33, 'X1b')],
    'Joints': [('Robot A', 1, 'Hip', 2, 'large', 88, None, None, None, None, 'passive', 'Hips', 'Legs', 'U3',
                'For the tests.'),
               ('Robot A', 2, 'Ankle', 2, 'small', 35, None, None, None, None, 'passive', 'Ankles', 'Legs', 'U3',
                'For the tests.'),
               ('Robot B', 1, 'Hip', 2, 'large', 150, 1.22, 1.28, 1.34, None, '', 'Hips', 'Legs', 'X1b',
                'For the tests.'),
               ('Robot B', 2, 'Elbow', 2, 'small', 19, 0.43, 0.45, 0.47, None, '', 'Elbows and wrists', 'Arms', 'X1b',
                'For the tests.')],
    # The magnets' grade classes by heat group (Matthias, 2026-10-28, "Accept as proposed"; EVIDENCE.md E2.11).
    'MagnetGrades': [('Hips', 'SH', 0.5, 'E', ''), ('Hips', 'UH', 0.4, 'E', ''), ('Hips', 'EH', 0.1, 'E', ''),
                     ('Ankles', 'SH', 0.6, 'E', ''), ('Ankles', 'UH', 0.4, 'E', ''),
                     ('Elbows and wrists', 'H', 0.7, 'E', ''), ('Elbows and wrists', 'SH', 0.3, 'E', '')],
    'JointParts': [('Motor stator', 100, 0.5, 1, 1.5, '', '', 1.0, 1.78, 2.67, 'the area', 'W; E', 'T1'),
                   ('Housing', 80, 0.5, 1, 1.5, 'Aluminium', '', 1.0, 1.39, 2.09, 'the surface', 'W; E', 'T1')],
    'Parts': [('Robot A', 'Joint module, small', 2, 'Ankles', 'Legs', None, None, None, '', '',
               'the sum of JointParts', 'W, E', 'T1'),
              ('Robot A', 'Joint module, large', 2, 'Hips', 'Legs', None, None, None, '', '', 'the sum of JointParts',
               'E', 'T1'),
              ('Robot A', 'Screws', 1, 'Frame and shell', 'Whole body', 0.2, 0.5, 1.0, 'Steel', '42CrMo4', '', 'E',
               'no source'),
              ('Robot A', 'Frame', 1, 'Frame and shell', 'Whole body', None, None, None, 'Aluminium', '',
               'the robot (Constants), less every other part', 'E', 'T1'),
              ('Robot B', 'Joint module, small', 2, 'Arms', 'Arms', None, None, None, '', '',
               "the Joints' module masses", 'S', 'X1b'),
              ('Robot B', 'Joint module, large', 2, 'Hips', 'Legs', None, None, None, '', '',
               "the Joints' module masses", 'S', 'X1b'),
              ('Robot B', 'Frame', 1, 'Frame and shell', 'Whole body', None, None, None, 'Aluminium', '',
               'the robot (Constants), less every other part', 'E', 'X1a')],
    'Shares': [('Motor stator', 'Copper', '', 0.23, 0.36, 0.45, '', 'S', 'For the tests.'),
               ('Motor stator', 'Steel', 'M19 or M27', None, None, None, 'Yes', 'E', 'For the tests.')],
    'Boards': [('Robot A', 'Driver board, small joint', 2, 'Ankles', 'Legs', 11.5, 22.9, 28.3, 'JointParts', 'PE_HVS',
                'W, E', 'T1'),
               ('Robot B', 'Mini PC board', 1, 'Torso boards', 'Torso', 60, 120, 180,
                'area times Constants board mass per cm2', 'ITC', 'E', 'X1a')],
    'Sensors': [('Robot A', 'Motion sensor', 1, 'Torso boards', 'Torso', 1, 2, 5, 'E', 'U3'),
                ('Robot B', 'Motion sensor, YIS320', 2, 'Torso boards', 'Torso', 1, 2, 5, 'E', 'X1a')],
    'SensorElements': [('Motion sensor', 'IMU', 1), ('Motion sensor, YIS320', 'IMU', 1)],
    'Constants': [('all', 'Board mass per cm2 (boards not weighed)', 0.28, 0.56, 0.84, 'g/cm2', 'E', 'E2.8'),
                  ('all', "Copper share of a board's mass", 0.125, 0.2, 0.287, 'share', 'S', 'Bizzo'),
                  ('all', 'Power conductor cross-section (2 conductors)', 0.52, 0.82, 1.31, 'mm2', 'S', 'TME'),
                  ('all', 'Signal conductor cross-section (2 conductors)', 0.065, 0.13, 0.2, 'mm2', 'E', 'E2.8'),
                  ('all', 'Copper density', 8.96, 8.96, 8.96, 'g/cm3', 'S', 'copper'),
                  ('all', 'Insulation and jackets, per mass of copper', 0.5, 1.0, 1.5, 'times', 'E', 'E2.8'),
                  ('Robot A', 'Robot mass, with battery', 34, 35, 36, 'kg', 'S', 'U1'),
                  ('Robot A', 'Battery voltage', 46.8, 46.8, 46.8, 'V', 'S', 'T1'),
                  ('Robot A', 'Battery capacity', 8.55, 9.0, 9.45, 'Ah', 'S', 'T1'),
                  ('Robot A', 'Cable length (four conductors)', 3, 6, 10, 'm', 'E', 'T1'),
                  ('Robot B', 'Robot mass, with battery', 32, 33, 34, 'kg', 'S', 'X1b'),
                  ('Robot B', 'Battery voltage', 43.2, 44.4, 48, 'V', 'S', 'X1a'),
                  ('Robot B', 'Battery capacity', 8, 8, 8, 'Ah', 'S', 'X1a'),
                  ('Robot B', 'Cable length (four conductors)', 10, 11, 12, 'm', 'S', 'X1a'),
                  # Last, so that the rows the tests change by their number stay where they were (points 3c, 3b).
                  ('all', 'Sensor housing, aluminium share', 0.25, 0.5, 0.75, 'share', 'E', 'point 3c'),
                  ('all', 'Board rest, glass fibre share', 0.4, 0.55, 0.7, 'share', 'S, E', 'E5.21'),
                  ('all', 'Board rest, bromine share', 0.044, 0.088, 0.132, 'share', 'E', 'E5.21')],
    'Sources': [('T1', 'China Post Securities, teardown', 'https://example.org', 'A teardown')],
}
APART_LISTS = [('Body region', 'Legs'), ('Body region', 'Arms'), ('Body region', 'Torso'), ('Body region', 'Head'),
               ('Body region', 'Whole body'), ('Part group', 'Joint drives'), ('Material', 'Aluminium'),
               ('Material', 'Copper'), ('Material', 'Steel'), ('Material', 'Plastics'), ('Material', 'Mixed'),
               ('Material', 'Not allocated')]


def build_apart(path, changes=None, drop=None, one_robot_layout=False):
    """Write the small parts-taken-apart file to `path`. `changes` maps (sheet, row index, column index) to a new
    value, or (sheet, row index) to None to leave the row out or to a tuple to add it; `drop` names a sheet to leave
    out; `one_robot_layout` writes the sheet Joints without its column Robot, as the file was before it held more than
    one robot."""
    book = Workbook()
    book.active.title = 'About'
    book.active.append(['For the tests.'])
    for name, headings in APART.items():
        if name == drop:
            continue
        rows = [list(r) for r in SMALL_APART[name]]
        for key, value in (changes or {}).items():
            if key[0] != name:
                continue
            if len(key) == 3:
                rows[key[1]][key[2]] = value
            elif value is None:
                rows[key[1]] = None
            else:
                rows.append(list(value))
        sheet = book.create_sheet(name)
        sheet.append(headings)
        for row in rows:
            if row is not None:
                sheet.append(row)
        if one_robot_layout and name == 'Joints':
            sheet.delete_cols(1)
    book.save(path)


class TestPartsTakenApart(unittest.TestCase):
    """01's checks of the file of robot type 1 taken apart; each test breaks one thing in the small file."""

    def faults(self, changes=None):
        with tempfile.TemporaryDirectory() as folder:
            build_apart(os.path.join(folder, 'apart.xlsx'), changes)
            build_lists(os.path.join(folder, 'lists.xlsx'), parts=[], materials=[], lists=APART_LISTS)
            apart = schemas.read_parts_taken_apart(os.path.join(folder, 'apart.xlsx'))
            lists = schemas.read_shared_lists(os.path.join(folder, 'lists.xlsx'))
        return schemas.check_parts_taken_apart(apart, lists, Params())

    def assertFaultSays(self, faults, words):
        self.assertEqual(len(faults), 1, [f.text for f in faults])
        self.assertIn(words, faults[0].text)

    def test_the_small_file_is_read_without_fault(self):
        self.assertEqual([f.text for f in self.faults()], [])

    def test_a_missing_sheet_stops_the_reading(self):
        for sheet in ('Shares', 'Robots'):
            with tempfile.TemporaryDirectory() as folder:
                build_apart(os.path.join(folder, 'apart.xlsx'), drop=sheet)
                with self.assertRaises(schemas.WorkbookError, msg=sheet):
                    schemas.read_parts_taken_apart(os.path.join(folder, 'apart.xlsx'))

    def test_the_one_robot_layout_is_refused(self):
        with tempfile.TemporaryDirectory() as folder:
            build_apart(os.path.join(folder, 'apart.xlsx'), one_robot_layout=True)
            with self.assertRaises(schemas.WorkbookError):
                schemas.read_parts_taken_apart(os.path.join(folder, 'apart.xlsx'))

    def test_a_robot_not_in_the_sheet_robots(self):
        self.assertFaultSays(self.faults({('Parts', 2, 0): 'Robot Z'}), "'Robot Z'")

    def test_a_robot_without_its_rows_on_a_sheet(self):
        faults = self.faults({('Boards', 1): None})
        self.assertFaultSays(faults, 'Robot B')
        self.assertIn('Boards', faults[0].text)

    def test_a_minimum_above_the_mode(self):
        self.assertFaultSays(self.faults({('Parts', 2, 5): 0.6}), 'in order')

    def test_an_estimate_narrower_than_half_to_one_and_a_half_times_its_mode(self):
        faults = self.faults({('Parts', 2, 5): 0.4, ('Parts', 2, 7): 0.6})
        self.assertFaultSays(faults, 'half to 1.5 times')

    def test_a_large_joint_factor_below_one_is_a_fault(self):
        self.assertFaultSays(self.faults({('JointParts', 1, 7): 0.6}), 'below 1')

    def test_a_group_of_shares_without_its_remainder(self):
        self.assertFaultSays(self.faults({('Shares', 1, 6): ''}), 'exactly one remainder')

    def test_shares_that_leave_the_remainder_below_zero(self):
        faults = self.faults({('Shares', 2): ('Motor stator', 'Plastics', 'coating', 0.3, 0.6, 0.9, '', 'E', '')})
        self.assertFaultSays(faults, 'below zero')

    def test_a_material_not_in_the_shared_lists(self):
        self.assertFaultSays(self.faults({('Parts', 2, 8): 'Brass'}), "'Brass'")

    def test_shares_of_a_part_that_does_not_exist(self):
        faults = self.faults({('Shares', 2): ('Rotor', 'Steel', '', None, None, None, 'Yes', 'E', '')})
        self.assertFaultSays(faults, "'Rotor'")

    def test_a_part_with_neither_one_material_nor_shares(self):
        self.assertFaultSays(self.faults({('JointParts', 1, 5): ''}), "'Housing'")

    def test_joints_that_do_not_match_the_joint_modules(self):
        faults = self.faults({('Joints', 0, 3): 3})
        self.assertFaultSays(faults, 'joint modules')
        self.assertIn('Robot A', faults[0].text)

    def test_a_joint_number_given_twice_in_one_robot(self):
        self.assertFaultSays(self.faults({('Joints', 1, 1): 1}), 'number 1')

    def test_a_module_mass_given_in_part(self):
        self.assertFaultSays(self.faults({('Joints', 2, 6): None}), 'module mass')

    def test_a_module_mass_out_of_order(self):
        self.assertFaultSays(self.faults({('Joints', 2, 6): 1.5}), 'in order')

    def test_a_magnet_grade_that_is_not_a_class(self):
        self.assertFaultSays(self.faults({('MagnetGrades', 0, 1): 'XX'}), "'XX'")

    def test_a_joint_whose_heat_group_has_no_grades(self):
        self.assertFaultSays(self.faults({('Joints', 0, 11): 'Feet'}), "'Feet'")

    def test_a_heat_groups_probabilities_that_do_not_add_up_to_one(self):
        self.assertFaultSays(self.faults({('MagnetGrades', 0, 2): 0.6}), 'add up to 1')

    def test_a_grade_class_given_twice_in_one_heat_group(self):
        self.assertFaultSays(self.faults({('MagnetGrades', 2, 1): 'SH', ('MagnetGrades', 1, 2): 0.4}), 'twice')

    def test_a_heat_group_that_no_joint_has(self):
        self.assertFaultSays(self.faults({('MagnetGrades', 7): ('Neck', 'H', 1.0, 'E', '')}), "'Neck'")

    def test_the_old_grade_columns_are_refused(self):
        # The sheet Joints with its base grade class and its scenarios, as it was before 2026-10-28: the old goes
        # entirely, and its headings stop the reading.
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, 'apart.xlsx')
            build_apart(path)
            book = load_workbook(path)
            book['Joints'].cell(1, 12).value = 'Magnet grade class (base)'
            book.save(path)
            with self.assertRaises(schemas.WorkbookError):
                schemas.read_parts_taken_apart(path)

    def test_a_sensor_element_of_a_sensor_not_listed(self):
        self.assertFaultSays(self.faults({('SensorElements', 0, 0): 'Camera'}), "'Camera'")

    def test_a_shared_constant_that_is_missing(self):
        self.assertFaultSays(self.faults({('Constants', 4): None}), "'Copper density'")

    def test_a_robots_own_constant_that_is_missing(self):
        faults = self.faults({('Constants', 10): None})
        self.assertFaultSays(faults, "'Robot mass, with battery'")
        self.assertIn('Robot B', faults[0].text)

    # Matthias, 2026-10-29, "Yes, build it" (the entry of type 1's research, point 7: 01 checks the new rows, tests
    # first, a planted fault for each): the joints as shares of the robot's mass (E5.17 point 1), the battery's energy
    # per kg of robot (point 2), a pack drawn from its energy (E5.14 point 3), the shell a share of the robot (point 6),
    # the housing aluminium or plastics (E5.19 point 1).
    SHARES = {('Joints', 2, 6): None, ('Joints', 2, 7): None, ('Joints', 2, 8): None, ('Joints', 2, 9): 0.75,
              ('Joints', 3, 6): None, ('Joints', 3, 7): None, ('Joints', 3, 8): None, ('Joints', 3, 9): 0.25,
              ('Constants', 14): ('all', 'Joint modules, share of the robot mass', 0.245, 0.49, 0.735, 'share', 'E',
                                  'E5.17')}

    def test_a_robot_whose_joints_are_shares_of_its_mass(self):
        self.assertEqual([f.text for f in self.faults(self.SHARES)], [])

    def test_joint_shares_that_do_not_add_up_to_one(self):
        self.assertFaultSays(self.faults({**self.SHARES, ('Joints', 3, 9): 0.2}), 'add up to 0.95')

    def test_a_robot_with_shares_for_some_of_its_joints_only(self):
        changes = {**self.SHARES, ('Joints', 3, 6): 0.43, ('Joints', 3, 7): 0.45, ('Joints', 3, 8): 0.47,
                   ('Joints', 3, 9): None, ('Joints', 2, 9): 1.0}
        self.assertFaultSays(self.faults(changes), 'all or none')

    def test_a_joint_with_a_module_mass_and_a_share(self):
        faults = self.faults({**self.SHARES, ('Joints', 3, 6): 0.43, ('Joints', 3, 7): 0.45, ('Joints', 3, 8): 0.47})
        self.assertFaultSays(faults, 'not both')

    def test_joint_shares_without_the_shared_share(self):
        without = {key: value for key, value in self.SHARES.items() if key[0] != 'Constants'}
        self.assertFaultSays(self.faults(without), "'Joint modules, share of the robot mass'")

    def test_board_rest_shares_that_could_leave_no_epoxy(self):
        # 2026-10-29, "Yes, enter them" (point 3b): epoxy is the remainder of a board's rest, 1 less its glass and its
        # bromine; their maxima must leave it above zero (rows 15 and 16 of the small file).
        self.assertEqual([f.text for f in self.faults()], [])
        self.assertFaultSays(self.faults({('Constants', 15, 4): 0.9}), 'the epoxy, their remainder')

    def test_a_battery_voltage_without_its_capacity(self):
        faults = self.faults({('Constants', 12): None})
        self.assertFaultSays(faults, 'not only its battery voltage')
        self.assertIn('Robot B', faults[0].text)

    def test_a_battery_without_voltage_and_capacity_needs_the_energy_per_kg_of_robot(self):
        neither = {('Constants', 11): None, ('Constants', 12): None}
        self.assertFaultSays(self.faults(neither), "'Battery energy per kg of robot'")
        given = {**neither, ('Constants', 14): ('all', 'Battery energy per kg of robot', 6, 12, 18, 'Wh/kg', 'E', '')}
        self.assertEqual([f.text for f in self.faults(given)], [])

    def test_a_pack_with_no_mass_and_a_housing_needs_the_pack_energy_per_kg(self):
        pack = {('Parts', 7): ('Robot B', 'Battery pack', 1, 'Battery pack', 'Torso', None, None, None, '', '',
                               'its cells, its management board and its housing', 'S, E', ''),
                ('Parts', 8): ('Robot B', 'Battery housing', 1, 'Battery pack', 'Torso', None, None, None, 'Mixed', '',
                               'the battery pack, less its cells and its management board', 'E', '')}
        self.assertFaultSays(self.faults(pack), "'Battery pack, Wh per kg (packs not weighed)'")
        given = {**pack, ('Constants', 14): ('all', 'Battery pack, Wh per kg (packs not weighed)', 109.3, 164, 232.5,
                                             'Wh/kg', 'S, E', '')}
        self.assertEqual([f.text for f in self.faults(given)], [])

    def test_a_part_that_is_a_share_of_the_robot_needs_the_shell_share(self):
        shell = {('Parts', 7): ('Robot B', 'Shell', 1, 'Frame and shell', 'Whole body', None, None, None, 'Plastics',
                                '', "a share of the robot's mass: Constants' Shell, share of the robot mass", 'E', '')}
        self.assertFaultSays(self.faults(shell), "'Shell, share of the robot mass'")
        given = {**shell, ('Constants', 14): ('all', 'Shell, share of the robot mass', 0.0575, 0.115, 0.1725, 'share',
                                              'E', '')}
        self.assertEqual([f.text for f in self.faults(given)], [])

    def test_a_material_written_as_either_of_two(self):
        self.assertEqual([f.text for f in self.faults({('Shares', 1, 1): 'Steel or Aluminium'})], [])
        self.assertFaultSays(self.faults({('Shares', 1, 1): 'Steel or Brass'}), "'Brass'")

    def test_a_label_that_is_not_known(self):
        self.assertFaultSays(self.faults({('Parts', 2, 11): 'X'}), "'X'")

    def test_two_parts_that_take_the_robots_remainder(self):
        faults = self.faults({('Parts', 7): ('Robot A', 'Shell', 1, 'Frame and shell', 'Whole body', None, None, None,
                                             'Plastics', '', 'the robot (Constants), less every other part', 'E',
                                             'T1')})
        self.assertFaultSays(faults, 'remainder of the robot')

    def test_a_body_region_not_in_the_lists(self):
        self.assertFaultSays(self.faults({('Sensors', 0, 4): 'Tail'}), "'Tail'")

    def test_the_real_file_holds_together(self):
        params = Params()
        paths = [absolute(params.data.parts_taken_apart), absolute(params.data.shared_lists)]
        if not all(os.path.exists(p) for p in paths):
            self.skipTest('the parts taken apart or the shared lists are not on disk')
        apart = schemas.read_parts_taken_apart(paths[0])
        lists = schemas.read_shared_lists(paths[1])
        self.assertEqual([f.text for f in schemas.check_parts_taken_apart(apart, lists, params)], [])
        # Every robot of type 1 with rows of its own since 2026-10-29 ("Yes, build it"): 32, the K1 in two versions.
        # EVIDENCE.md E2.8 and E2.8e, as agreed by Matthias on 2026-10-27, for the three taken apart first. The G1: 23
        # joints in 12 kinds, 9 large;
        # the small joint's weighed parts add up to 525 g (China Post Securities, slide 43), the large joint's at the
        # factors' modes to 1,011.717 g (by hand: 154.1 g of motor x 1.78, 155.2 g of gears x 2.33, 12.8 g of board x
        # 3.56, 11 g of bearings x 1.78, 79.3 g of housing x 1.39, 23.7 + 50.8 + 38.1 g x 1.78). The X1: 29 joints
        # (R86-3 x 6, R86-2 x 9, R52 x 10, L28 x 4), 15 large, their modes adding up to 20.27 kg (AgiBot), and its two
        # grippers as small joints of 0.43 kg since 2026-10-29 (E5.19 point 2): 31 joints, 21.13 kg. The N1: 23
        # joints (90 N·m x 4, 54 N·m x 7, 30 N·m x 12), 11 large, their modes from the fit as E2.8e rounds them
        # (0.88, 0.70, 0.54 kg) adding up to 4 x 0.88 + 7 x 0.70 + 12 x 0.54 = 14.90 kg.
        self.assertEqual(len(apart.robots), 32)
        self.assertTrue({'AgiBot X1', 'Fourier N1', 'Unitree G1'} <= set(apart.robots))
        joints = lambda robot: [j for j in apart.joints if j.robot == robot]  # noqa: E731
        self.assertEqual((len(joints('Unitree G1')), sum(j.count for j in joints('Unitree G1'))), (12, 23))
        self.assertEqual(sum(j.count for j in joints('Unitree G1') if j.module == 'large'), 9)
        self.assertEqual(sum(j.count for j in joints('AgiBot X1')), 31)
        self.assertEqual(sum(j.count for j in joints('AgiBot X1') if j.module == 'large'), 15)
        self.assertAlmostEqual(sum(j.count * j.module_mass[1] for j in joints('AgiBot X1')), 21.13, places=6)
        self.assertEqual(sum(j.count for j in joints('Fourier N1')), 23)
        self.assertEqual(sum(j.count for j in joints('Fourier N1') if j.module == 'large'), 11)
        self.assertAlmostEqual(sum(j.count * j.module_mass[1] for j in joints('Fourier N1')), 14.90, places=6)
        self.assertAlmostEqual(sum(c.grams for c in apart.joint_parts), 525.0, places=6)
        self.assertAlmostEqual(sum(c.grams * c.factor[1] for c in apart.joint_parts), 1011.717, places=6)


# The sibling projects' files robot type 1 taken apart reads where they lie (step 4): small copies of their sheets.
# The car board file's ten metals, as RAWCLICVehicleElectronics' 04_ holds them for every category.
SIBLING_PCB = [(c, e, 1, 2, 3) for c in ('PE_HVS', 'ITC')
               for e in ('Cu', 'Ag', 'Au', 'Sn', 'Pb', 'Ni', 'Pd', 'Al', 'Zn', 'Fe')]
SIBLING_SENSORS = [('IMU', '6-axis MEMS'), ('angle sensor', 'Hall IC')]
SIBLING_TANTALUM = [('Cu', 0.1, 0.18, 0.28), ('Ta', 0, 0.0001, 0.00045)]
SIBLING_MAGNETS = [('N35H', 'max', 'H'), ('N42SH', 'max', 'SH'), ('N40UH', 'max', 'UH'), ('N38EH', 'max', 'EH')]


# RAWCLICVehicleTractionMotor's steels (its element workbook's sheets ElectricalSteel and CastFeSteel): a row of minima
# and of maxima per grade, the element written 'balance' the remainder.
SIBLING_STEELS = {'ElectricalSteel': (['Grade', 'Fe', 'Si', 'C', 'Range'],
                                      [('M19', 'balance', 0.025, 0, 'min'), ('M19', 'balance', 0.035, 0.0002, 'max'),
                                       ('M27', 'balance', 0.024, 0, 'min'), ('M27', 'balance', 0.031, 0.00025, 'max')]),
                  'CastFeSteel': (['Grade', 'Fe', 'C', 'Cr', 'Ni', 'Range'],
                                  [('Type 304', 'balance', 0, 0.18, 0.08, 'min'),
                                   ('Type 304', 'balance', 0.0008, 0.2, 0.105, 'max')])}


def build_steels(path, steels=None):
    """Write the small copy of the sibling's steel sheets to `path`."""
    book = Workbook()
    book.remove(book.active)
    for sheet_name, (header, rows) in (steels or SIBLING_STEELS).items():
        sheet = book.create_sheet(sheet_name)
        sheet.append(header)
        for row in rows:
            sheet.append(list(row))
    book.save(path)


def build_siblings(folder, pcb=None, sensors=None, tantalum=None, magnets=None):
    """Write small copies of the five sibling sheets into `folder` and return Params that name them."""
    paths = {'traction_motor_steels': os.path.join(folder, 'traction_motor_steels.xlsx')}
    build_steels(paths['traction_motor_steels'])
    for key, sheet_name, header, rows in (
            ('vehicle_pcb_composition', 'PCPComposition',
             ['PCBCategory', 'Element', 'mg_cm2_min', 'mg_cm2_mode', 'mg_cm2_max'], pcb or SIBLING_PCB),
            ('vehicle_sensor_composition', 'Sensor Details', ['SensorType', 'Description'],
             sensors or SIBLING_SENSORS),
            ('data_center_element_composition', 'BoardsAndComponents', ['Element', 'Min', 'Mode', 'Max'],
             tantalum or SIBLING_TANTALUM),
            ('traction_motor_magnets', 'PermanentMagnetNdFeB', ['Grade'] + [''] * 11 + ['Range', 'TempClass'],
             [(g,) + ('',) * 11 + (r, k) for g, r, k in (magnets or SIBLING_MAGNETS)])):
        book = Workbook()
        sheet = book.active
        sheet.title = sheet_name
        sheet.append(header)
        for row in rows:
            sheet.append(list(row))
        paths[key] = os.path.join(folder, key + '.xlsx')
        book.save(paths[key])
    params = Params()
    for key, path in paths.items():
        setattr(params.data, key, path)
    return params


class TestSiblingData(unittest.TestCase):
    """Step 4: the car board categories, car sensor types, tantalum row and magnet grade classes that robot type 1
    taken apart names are in the sibling projects' files; each test leaves one of them out."""

    def faults(self, apart_changes=None, **siblings):
        with tempfile.TemporaryDirectory() as folder:
            build_apart(os.path.join(folder, 'apart.xlsx'), apart_changes)
            apart = schemas.read_parts_taken_apart(os.path.join(folder, 'apart.xlsx'))
            params = build_siblings(folder, **siblings)
            return schemas.check_sibling_data(apart, params)

    def assertFaultSays(self, faults, words):
        self.assertEqual(len(faults), 1, [f.text for f in faults])
        self.assertIn(words, faults[0].text)

    def test_everything_named_is_there(self):
        self.assertEqual([f.text for f in self.faults()], [])

    def test_a_magnet_grade_class_a_part_names_that_the_traction_motors_do_not_have(self):
        # Point 4: a part of magnets naming its class; the small traction-motor file has H, SH, UH and EH only.
        faults = self.faults({('Parts', 2, 8): 'Magnets (NdFeB)', ('Parts', 2, 9): 'N'})
        self.assertFaultSays(faults, "the magnet grade class 'N' is not in")

    def test_a_board_category_the_car_file_does_not_have(self):
        self.assertFaultSays(self.faults({('Boards', 0, 9): 'XYZ'}), "'XYZ'")

    def test_a_board_category_without_one_of_its_metals(self):
        pcb = [r for r in SIBLING_PCB if not (r[0] == 'PE_HVS' and r[1] == 'Pd')]
        self.assertFaultSays(self.faults(pcb=pcb), 'Pd')

    def test_a_sensor_type_the_car_file_does_not_have(self):
        self.assertFaultSays(self.faults({('SensorElements', 0, 1): 'radar'}), "'radar'")

    def test_a_magnet_grade_class_the_traction_motors_do_not_have(self):
        magnets = [m for m in SIBLING_MAGNETS if m[2] != 'UH']
        self.assertFaultSays(self.faults(magnets=magnets), "'UH'")

    def test_no_tantalum_in_the_data_centre_file(self):
        self.assertFaultSays(self.faults(tantalum=[('Cu', 0.1, 0.18, 0.28)]), 'Ta')

    def test_a_sibling_file_that_is_not_there(self):
        with tempfile.TemporaryDirectory() as folder:
            build_apart(os.path.join(folder, 'apart.xlsx'))
            apart = schemas.read_parts_taken_apart(os.path.join(folder, 'apart.xlsx'))
            params = build_siblings(folder)
            params.data.vehicle_sensor_composition = os.path.join(folder, 'gone.xlsx')
            self.assertFaultSays(schemas.check_sibling_data(apart, params), 'data.vehicle_sensor_composition')

    def test_the_real_sibling_files_hold_what_robot_type_1_names(self):
        params = Params()
        paths = [params.data.vehicle_pcb_composition, params.data.vehicle_sensor_composition,
                 params.data.data_center_element_composition, params.data.traction_motor_magnets,
                 params.data.traction_motor_steels, absolute(params.data.parts_taken_apart)]
        if not all(os.path.exists(p) for p in paths):
            self.skipTest('a sibling file or the parts taken apart are not on disk')
        apart = schemas.read_parts_taken_apart(paths[-1])
        self.assertEqual([f.text for f in schemas.check_sibling_data(apart, params)], [])


    def test_the_steels_sheets_that_are_not_there(self):
        with tempfile.TemporaryDirectory() as folder:
            build_apart(os.path.join(folder, 'apart.xlsx'))
            apart = schemas.read_parts_taken_apart(os.path.join(folder, 'apart.xlsx'))
            params = build_siblings(folder)
            build_steels(params.data.traction_motor_steels, {'ElectricalSteel': SIBLING_STEELS['ElectricalSteel']})
            self.assertFaultSays(schemas.check_sibling_data(apart, params), "'CastFeSteel'")


class TestGrades(unittest.TestCase):
    """Matthias, 2026-10-29, "Yes, build it" (point 4, the grades of E5.18 and E5.19; point 7, a planted fault for
    each): every steel or titanium part has a grade, a copper part may, no other; every grade is a recipe or a steel of
    the sibling's workbook, and that steel's limits are in order."""

    RECIPES = ('42CrMo4', 'GCr15', 'Brass', 'E-glass', 'Epoxy')   # E-glass and Epoxy: the boards' rest (point 3b)

    def faults(self, changes=None, steels=None, recipes=None):
        with tempfile.TemporaryDirectory() as folder:
            build_apart(os.path.join(folder, 'apart.xlsx'), changes)
            apart = schemas.read_parts_taken_apart(os.path.join(folder, 'apart.xlsx'))
            params = Params()
            params.data.traction_motor_steels = os.path.join(folder, 'steels.xlsx')
            build_steels(params.data.traction_motor_steels, steels)
            recipes = types.SimpleNamespace(items=[types.SimpleNamespace(recipe=r) for r in recipes or self.RECIPES])
            return schemas.check_grades(apart, recipes, params)

    def assertFaultSays(self, faults, words):
        self.assertEqual(len(faults), 1, [f.text for f in faults])
        self.assertIn(words, faults[0].text)

    def test_the_small_files_grades_are_all_known(self):
        self.assertEqual([f.text for f in self.faults()], [])

    def test_a_steel_part_without_its_grade(self):
        self.assertFaultSays(self.faults({('Parts', 2, 9): ''}), 'give its grade')

    def test_a_grade_that_is_neither_a_recipe_nor_a_sibling_steel(self):
        self.assertFaultSays(self.faults({('Parts', 2, 9): 'S355'}), "'S355'")

    def test_either_of_two_grades_each_known(self):
        self.assertEqual([f.text for f in self.faults({('Parts', 2, 9): '42CrMo4 or GCr15'})], [])
        self.assertFaultSays(self.faults({('Shares', 1, 2): 'M19 or M99'}), "'M99'")

    def test_a_grade_on_a_material_that_has_none(self):
        self.assertFaultSays(self.faults({('JointParts', 1, 6): 'GCr15'}), "not for 'Aluminium'")

    def test_a_copper_part_may_name_its_recipe(self):
        self.assertEqual([f.text for f in self.faults({('Shares', 0, 2): 'Brass'})], [])

    def test_the_recipes_of_the_boards_rest_must_be_there(self):
        # 2026-10-29, "Yes, enter them", point 3b: the boards' glass fibre and epoxy by the recipes E-glass and Epoxy.
        without = tuple(r for r in self.RECIPES if r != 'Epoxy')
        self.assertFaultSays(self.faults(recipes=without), "the recipe 'Epoxy' of the boards' rest")

    def test_a_part_of_magnets_may_name_its_grade_class(self):
        # 2026-10-29, "Yes, build them", point 4: the RP1's N52 magnets in its covers, grade class N.
        magnets = {('Parts', 2, 8): 'Magnets (NdFeB)'}
        self.assertEqual([f.text for f in self.faults({**magnets, ('Parts', 2, 9): 'N'})], [])
        self.assertFaultSays(self.faults({**magnets, ('Parts', 2, 9): 'N52'}), 'a grade of magnets is a grade class')

    def test_a_joints_magnets_take_their_heat_groups_class(self):
        self.assertFaultSays(self.faults({('JointParts', 1, 5): 'Magnets (NdFeB)', ('JointParts', 1, 6): 'N'}),
                             "in a joint its heat group's")

    def test_a_sibling_steel_whose_maximum_is_below_its_minimum(self):
        # The fault found on 2026-10-29 in RAWCLICVehicleTractionMotor's sheet CastFeSteel, rows 16 and 17: Type 304's
        # chromium 0.18 at its minimum and 0.02 at its maximum.
        steels = dict(SIBLING_STEELS)
        steels['CastFeSteel'] = (['Grade', 'Fe', 'C', 'Cr', 'Ni', 'Range'],
                                 [('Type 304', 'balance', 0, 0.18, 0.08, 'min'),
                                  ('Type 304', 'balance', 0.0008, 0.02, 0.105, 'max')])
        self.assertFaultSays(self.faults({('Parts', 2, 9): 'Type 304'}, steels), 'Cr')

    def test_the_real_grades_can_all_be_drawn(self):
        params = Params()
        paths = [absolute(params.data.parts_taken_apart), absolute(params.data.element_recipes),
                 params.data.traction_motor_steels]
        if not all(os.path.exists(p) for p in paths):
            self.skipTest('the parts file, the element recipes or the sibling steels are not on disk')
        apart = schemas.read_parts_taken_apart(paths[0])
        recipes = schemas.read_element_recipes(paths[1])
        self.assertEqual([f.text for f in schemas.check_grades(apart, recipes, params)], [])


# --- The register of every known humanoid robot (data/researched/humanoid_register.xlsx) ------------------------------

REGISTER_HEADINGS = ['Number', 'Robot', 'Maker', 'Country', 'Status', 'Height (cm)', 'Mass (kg), stated',
                     'Mass min (kg)', 'Mass mode (kg)', 'Mass max (kg)', 'Mass basis', 'Types', 'Form', 'Entry',
                     'Taken apart as', 'Decided in', 'Source']
# A small register: (number, robot, mass min, mode, max, types, form, entry, taken apart as). Rows 2 to 6.
REGISTER_ROWS = [
    (1, 'Robot A', 34, 35, 36, '1', 'legs', '', 'Unitree G1'),
    (2, 'Robot B', 59, 60, 61, '3, 4', 'wheels', '', ''),
    (3, 'Robot B again', 59, 60, 61, '', 'wheels', 'the same robot as 2', ''),
    (4, 'Arm C', None, None, None, '', 'two arms', 'not assigned: arms only', ''),
    (5, 'Robot D', None, None, None, '6', 'not stated', '', ''),
]


def register_cells(n, robot, low, mode, high, types, form, entry, apart):
    """The cells of one robot of the small register, in the order of REGISTER_HEADINGS."""
    return [n, robot, 'Maker', 'China', 'Prototype', '150', '' if mode is None else str(mode), low, mode, high,
            'the stated mass +- 1 kg', types, form, entry, apart, 'E4.1', 'https://example.org']


def in_the_run(row) -> str:
    """What the sheet Type 1 says of a robot in 02's run: drawn with its own parts (taken apart), or not drawn (no
    mass stated)."""
    n, robot, low, mode, high, types, form, entry, apart = row
    return 'drawn: its own parts' if apart else 'not drawn: no mass stated'


def build_register(path, rows=None, headings=None, type_one=None, without_type_one=False):
    """A small register in the layout of data/researched/humanoid_register.xlsx: the sheet Robots, and the sheet Type 1
    with every robot of type 1 and what the run does with it. `type_one` replaces the rows of the sheet Type 1 (each a
    row of REGISTER_ROWS and its 'In the run'); `without_type_one` leaves the sheet out."""
    rows = REGISTER_ROWS if rows is None else rows
    book = Workbook()
    book.active.title = 'About'
    sheet = book.create_sheet('Robots')
    sheet.append(headings or REGISTER_HEADINGS)
    for row in rows:
        sheet.append(register_cells(*row))
    if not without_type_one:
        listed = book.create_sheet('Type 1')
        listed.append(REGISTER_HEADINGS + ['In the run'])
        of_type = [(r, in_the_run(r)) for r in rows if '1' in [t.strip() for t in r[5].split(',')]]
        for row, run in (of_type if type_one is None else type_one):
            listed.append(register_cells(*row) + [run])
    book.save(path)


class TestRegister(unittest.TestCase):
    """01 reads the register and reports each fault it guards against (EVIDENCE.md E4 to E4.11)."""

    def faults(self, rows=None, apart_names=('Unitree G1', 'AgiBot X1', 'Fourier N1')):
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, 'register.xlsx')
            build_register(path, rows)
            register = schemas.read_register(path)
        return schemas.check_register(register, Params(), apart_names)

    def assertFaultSays(self, faults, words):
        self.assertTrue(any(words in f.text for f in faults), msg=[f.text for f in faults])

    def changed(self, k, **values):
        rows = [list(r) for r in REGISTER_ROWS]
        keys = ['number', 'robot', 'low', 'mode', 'high', 'types', 'form', 'entry', 'apart']
        for key, value in values.items():
            rows[k][keys.index(key)] = value
        return [tuple(r) for r in rows]

    def test_a_register_that_holds_together(self):
        self.assertEqual([f.text for f in self.faults()], [])

    def test_it_is_read_robot_by_robot(self):
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, 'register.xlsx')
            build_register(path)
            register = schemas.read_register(path)
        self.assertEqual([r.number for r in register], [1, 2, 3, 4, 5])
        self.assertEqual(register[1].types, (3, 4))
        self.assertEqual(register[0].mass, (34, 35, 36))
        self.assertEqual(register[4].mass, (None, None, None))
        self.assertEqual(register[0].taken_apart, 'Unitree G1')

    def test_other_headings_are_refused(self):
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, 'register.xlsx')
            build_register(path, headings=REGISTER_HEADINGS[:-1] + ['Link'])
            with self.assertRaises(schemas.WorkbookError):
                schemas.read_register(path)

    def test_a_number_given_twice(self):
        self.assertFaultSays(self.faults(self.changed(1, number=1)), 'number 1')

    def test_three_types(self):
        self.assertFaultSays(self.faults(self.changed(1, types='3, 4, 5')), 'at most two')

    def test_a_type_that_does_not_exist(self):
        self.assertFaultSays(self.faults(self.changed(1, types='7')), "'7'")

    def test_the_same_type_twice(self):
        self.assertFaultSays(self.faults(self.changed(1, types='3, 3')), 'twice')

    def test_a_robot_without_a_type_nor_why(self):
        self.assertFaultSays(self.faults(self.changed(1, types='')), 'no type')

    def test_a_robot_left_out_that_has_types(self):
        self.assertFaultSays(self.faults(self.changed(3, types='3')), 'no types')

    def test_an_entry_that_is_not_known(self):
        self.assertFaultSays(self.faults(self.changed(3, entry='maybe later')), "'maybe later'")

    def test_the_same_robot_as_one_that_is_not_there(self):
        self.assertFaultSays(self.faults(self.changed(2, entry='the same robot as 9')), '9')

    def test_the_same_robot_as_a_second_entry(self):
        rows = self.changed(3, entry='the same robot as 3', types='', form='two arms')
        self.assertFaultSays(self.faults(rows), 'second entry')

    def test_a_form_that_is_not_known(self):
        self.assertFaultSays(self.faults(self.changed(1, form='hovers')), "'hovers'")

    def test_a_mass_out_of_order(self):
        self.assertFaultSays(self.faults(self.changed(1, low=70)), 'in order')

    def test_a_mass_given_in_part(self):
        self.assertFaultSays(self.faults(self.changed(1, high=None)), 'min, mode and max')

    def test_a_robot_of_type_1_below_its_lowest_mass(self):
        rows = self.changed(0, low=9, mode=10, high=11)
        self.assertFaultSays(self.faults(rows), '15')

    def test_a_robot_taken_apart_that_the_parts_file_does_not_have(self):
        self.assertFaultSays(self.faults(self.changed(0, apart='Unitree G2')), "'Unitree G2'")

    def test_a_robot_in_two_versions(self):
        # "Taken apart as" may name several robots of the parts file, each drawn as often (the K1's two batteries).
        self.assertEqual([f.text for f in self.faults(self.changed(0, apart='Unitree G1; AgiBot X1'))], [])
        self.assertFaultSays(self.faults(self.changed(0, apart='Unitree G1; Unitree G2')), "'Unitree G2'")

    def test_a_robot_taken_apart_that_is_not_type_1(self):
        self.assertFaultSays(self.faults(self.changed(1, apart='AgiBot X1')), 'type 1')

    def test_the_real_register_holds_together(self):
        params = Params()
        path, apart_path = absolute(params.data.register), absolute(params.data.parts_taken_apart)
        if not all(os.path.exists(p) for p in (path, apart_path)):
            self.skipTest('the register or the parts taken apart are not on disk')
        apart = schemas.read_parts_taken_apart(apart_path)
        register = schemas.read_register(path)
        self.assertEqual([f.text for f in schemas.check_register(register, params, list(apart.robots))], [])
        type_one = schemas.read_type_sheet(path, 1)
        self.assertEqual([f.text for f in schemas.check_type_one(register, type_one, params)], [])


class TestGradeRecipes(unittest.TestCase):
    """The steels, the titanium, the brass and the nickel of type 1 in the real element recipes (Matthias, 2026-10-29:
    "I agree with the steel and titanium proposal of E5.18 ..." and "I agree with the Mixed proposal of E5.19 ..."),
    at their agreed modes, each adding up with its remainder."""

    # Recipe -> {element: (min, mode, max)} as agreed (EVIDENCE.md E5.18, E5.19), mass fractions; the remainder apart.
    AGREED = {
        '20CrMnTi': {'C': (0.0017, 0.002, 0.0023), 'Cr': (0.01, 0.0115, 0.013), 'Ti': (0.0004, 0.0007, 0.001),
                     'Mn': (0.008, 0.0095, 0.011)},
        'GCr15': {'C': (0.0095, 0.01, 0.0105), 'Cr': (0.014, 0.01525, 0.0165), 'Ni': (0, 0.0015, 0.003)},
        '42CrMo4': {'Cr': (0.009, 0.0105, 0.012), 'Mo': (0.0015, 0.00225, 0.003)},
        'C45': {'C': (0.0042, 0.0046, 0.005), 'Ni': (0, 0.002, 0.004), 'Mo': (0, 0.0005, 0.001)},
        'P20': {'Cr': (0.014, 0.017, 0.02), 'Mo': (0.003, 0.00425, 0.0055)},
        '17-4PH': {'Cr': (0.0775, 0.155, 0.2325), 'Ni': (0.024, 0.048, 0.072), 'Cu': (0.017, 0.034, 0.051)},
        'Ti-6Al-4V': {'Al': (0.055, 0.06125, 0.0675), 'V': (0.035, 0.04, 0.045)},
        'Brass': {'Cu': (0.58, 0.62, 0.65)},
    }
    REMAINDER = {'20CrMnTi': 'Fe', 'GCr15': 'Fe', '42CrMo4': 'Fe', 'C45': 'Fe', 'P20': 'Fe', '17-4PH': 'Fe',
                 'Ti-6Al-4V': 'Ti', 'Brass': 'Zn', 'Nickel strip': 'Ni'}

    @classmethod
    def setUpClass(cls):
        path = absolute(Params().data.element_recipes)
        if not os.path.exists(path):
            raise unittest.SkipTest('the element recipes are not on disk')
        cls.recipes = schemas.read_element_recipes(path)

    def test_every_grade_has_its_recipe_at_the_agreed_values(self):
        items = {}
        for i in self.recipes.items:
            items.setdefault(i.recipe, []).append(i)
        for recipe, elements in self.AGREED.items():
            self.assertIn(recipe, items, msg=recipe)
            for element, bounds in elements.items():
                found = next(i for i in items[recipe] if i.item == element)
                self.assertEqual(found.share, bounds, msg=f'{recipe} {element}')
        for recipe, rest in self.REMAINDER.items():
            remainder = [i for i in items[recipe] if i.remainder]
            self.assertEqual([i.item for i in remainder], [rest], msg=recipe)

    def test_at_the_modes_each_adds_up_to_one(self):
        for recipe in self.REMAINDER:
            shares = schemas.element_shares(self.recipes, recipe)
            self.assertAlmostEqual(sum(shares.values()), 1.0, places=9, msg=recipe)


class TestTypeOneSheet(unittest.TestCase):
    """The sheet Type 1 of the register (Matthias, 2026-10-29: "There has to be an excel sheet where all the robots of
    the different types are fully listed. For now it has to be just all for type 1."): 01 reads it and reports each
    way it can fail to list every robot of type 1 of the sheet Robots, as it is there, with what the run does with
    it."""

    def faults(self, rows=None, type_one=None):
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, 'register.xlsx')
            build_register(path, rows, type_one=type_one)
            register, listed = schemas.read_register(path), schemas.read_type_sheet(path, 1)
        return schemas.check_type_one(register, listed, Params())

    def assertFaultSays(self, faults, words):
        self.assertTrue(any(words in f.text for f in faults), msg=[f.text for f in faults])

    # The small register with two more robots of type 1: one with a mass and its own rows in two versions, one
    # without a mass (rows 7 and 8).
    ROWS = REGISTER_ROWS + [(6, 'Robot E', 29, 30, 31, '4, 1', 'legs', '', 'AgiBot X1; Fourier N1'),
                            (7, 'Robot F', None, None, None, '1', 'legs', '', '')]

    def listed(self, rows=None):
        """The rows of the sheet Type 1 as the builder writes them, from `rows` (default ROWS)."""
        rows = self.ROWS if rows is None else rows
        return [(r, in_the_run(r)) for r in rows if '1' in [t.strip() for t in r[5].split(',')]]

    def test_a_sheet_that_lists_every_robot_of_type_1(self):
        self.assertEqual([f.text for f in self.faults(self.ROWS)], [])

    def test_it_is_read_robot_by_robot_with_what_the_run_does(self):
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, 'register.xlsx')
            build_register(path, self.ROWS)
            listed = schemas.read_type_sheet(path, 1)
        self.assertEqual([r.number for r in listed], [1, 6, 7])
        self.assertEqual([r.in_run for r in listed], ['drawn: its own parts', 'drawn: its own parts',
                                                      'not drawn: no mass stated'])
        self.assertEqual(listed[1].types, (4, 1))
        self.assertEqual(listed[1].mass, (29, 30, 31))

    def test_a_register_without_the_sheet_is_refused(self):
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, 'register.xlsx')
            build_register(path, without_type_one=True)
            with self.assertRaises(schemas.WorkbookError):
                schemas.read_type_sheet(path, 1)

    def test_a_robot_of_type_1_missing_from_the_sheet(self):
        self.assertFaultSays(self.faults(self.ROWS, self.listed()[:-1]), '7 Robot F')

    def test_a_robot_listed_that_is_not_of_type_1(self):
        other = (REGISTER_ROWS[1], 'drawn: its own parts')
        self.assertFaultSays(self.faults(self.ROWS, self.listed() + [other]), '2 Robot B')

    def test_a_robot_listed_twice(self):
        self.assertFaultSays(self.faults(self.ROWS, self.listed() + self.listed()[:1]), 'twice')

    def test_a_robot_listed_otherwise_than_in_the_sheet_robots(self):
        changed = [((6, 'Robot E', 29, 32, 35, '4, 1', 'legs', '', ''), self.listed()[1][1])]
        self.assertFaultSays(self.faults(self.ROWS, self.listed()[:1] + changed + self.listed()[2:]), 'Mass mode')

    def test_a_robot_without_a_mass_said_to_be_drawn(self):
        wrong = [(self.listed()[2][0], self.listed()[1][1])]
        self.assertFaultSays(self.faults(self.ROWS, self.listed()[:2] + wrong), 'not drawn')

    def test_a_robot_with_a_mass_said_not_to_be_drawn(self):
        wrong = [(self.listed()[1][0], 'not drawn: no mass stated')]
        self.assertFaultSays(self.faults(self.ROWS, self.listed()[:1] + wrong + self.listed()[2:]), 'its own parts')

    def test_a_robot_of_type_1_with_a_mass_and_no_rows_of_its_own(self):
        # Since 2026-10-29 (the entry, decision C) a robot of type 1 is drawn from its own rows only.
        rows = [list(r) for r in self.ROWS]
        rows[5][8] = ''
        rows = [tuple(r) for r in rows]
        listed = [(r, 'drawn: its own parts') for r in rows if '1' in [t.strip() for t in r[5].split(',')]]
        self.assertFaultSays(self.faults(rows, listed), 'no rows of its own')

    def test_a_robot_taken_apart_said_not_to_be_drawn(self):
        wrong = [(self.listed()[0][0], 'not drawn: no mass stated')]
        self.assertFaultSays(self.faults(self.ROWS, wrong + self.listed()[1:]), 'its own parts')


if __name__ == '__main__':
    unittest.main()
