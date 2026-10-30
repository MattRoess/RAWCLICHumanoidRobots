"""
Tests of the builders of the parts files: tools/1_compact_research_developer/make_parts_taken_apart.py (type 1's,
data/researched/1_compact_research_developer/parts_taken_apart.xlsx) and
tools/2_full_scale_general_purpose_ai_research/make_parts_taken_apart.py (type 2's), with the machinery they share,
tools/parts_taken_apart.py.

**Copyright notice:** Copyright © 2026 Empa, Matthias Roesslein
**Licence:** CC BY 4.0, Creative Commons Attribution 4.0 International; the full text is in `LICENSE`
at the root of the repository.

Matthias, 2026-10-27: "Yes, put the builder in tools/make_parts_taken_apart.py as proposed, with its test and
documents"; 2026-10-29: "Yes, build it" (every robot of type 1 its own rows). The builder is the master of the data
file: what it writes must pass every check 01 makes of the file, hold every robot of type 1 as agreed (EVIDENCE.md E5),
and, while nobody has changed it, equal the file on disk cell by cell. Type 2's builder (2026-10-30, "Yes, build C")
writes its batch 1 as agreed in EVIDENCE.md E6.8. Everything is written into temporary folders: press Run, nothing in
the project is changed.
"""
import importlib.util
import os
import sys
import tempfile
import unittest

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)


def load(name: str, path: str):
    """Import a file of the project whose name starts with a digit or lies outside a package."""
    spec = importlib.util.spec_from_file_location(name, os.path.join(ROOT, path))
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


schemas = load('schemas', '01_schemas.py')
builder = load('make_parts_taken_apart', os.path.join('tools', '1_compact_research_developer',
                                                       'make_parts_taken_apart.py'))
builder_two = load('make_parts_taken_apart_2', os.path.join('tools', '2_full_scale_general_purpose_ai_research',
                                                           'make_parts_taken_apart.py'))
from src.params_schema import Params, absolute  # noqa: E402


# The robots the register's sheet Type 1 names in its column "Taken apart as", every version once (its builder).
REGISTER = load('make_humanoid_register', os.path.join('tools', 'make_humanoid_register.py'))
TYPE_ONE = [name for names in REGISTER.TAKEN_APART[1].values() for name in schemas.versions(names)]
TYPE_TWO = [name for names in REGISTER.TAKEN_APART[2].values() for name in schemas.versions(names)]
# Type 2's batch 1 (EVIDENCE.md E6.8), the GR-1 in two versions; batch 2 since 2026-10-30 (E7.5), TestTypeTwoBatchTwo.
BATCH_ONE_TWO = ['Fourier GR-1 one camera', 'Fourier GR-1 two cameras', 'Fourier GR-2', 'Unitree H1-2', 'Unitree H2',
                 'Unitree H2 Plus']


def written(which=None):
    """What a builder writes (type 1's unless `which` is another), read with 01's reader, from a temporary folder."""
    with tempfile.TemporaryDirectory() as folder:
        path = os.path.join(folder, 'parts_taken_apart.xlsx')
        (which or builder).build(path)
        return schemas.read_parts_taken_apart(path), load_workbook(path)


class TestMakePartsTakenApart(unittest.TestCase):
    """What the builder writes holds together as 01 checks it, holds every robot of type 1 as agreed, and equals the
    data file on disk."""

    @classmethod
    def setUpClass(cls):
        cls.apart, cls.book = written()

    def test_what_it_writes_passes_every_check_of_01(self):
        params = Params()
        lists_path, recipes_path = absolute(params.data.shared_lists), absolute(params.data.element_recipes)
        if not all(os.path.exists(p) for p in (lists_path, recipes_path)):
            self.skipTest('the shared lists or the element recipes are not on disk')
        lists = schemas.read_shared_lists(lists_path)
        recipes = schemas.read_element_recipes(recipes_path)
        faults = (schemas.check_parts_taken_apart(self.apart, lists, params)
                  + schemas.check_sibling_data(self.apart, params)
                  + schemas.check_grades(self.apart, recipes, params))
        self.assertEqual([f.text for f in faults], [])

    def test_every_robot_of_type_1_has_its_own_rows(self):
        # Matthias, 2026-10-29: "Yes, build it" (the entry, point 2: every robot of the sheet Type 1 its own rows;
        # point 3: the K1 in its two versions).
        self.assertEqual(sorted(self.apart.robots), sorted(TYPE_ONE))
        self.assertEqual(len(self.apart.robots), 32)
        self.assertIn('Booster K1 5 Ah', self.apart.robots)
        self.assertIn('Booster K1 2 Ah', self.apart.robots)

    def test_the_k1_in_its_two_versions_differs_in_its_battery_and_computer_only(self):
        # E5.14 point 2: 5 Ah (Education, Professional) or 2 Ah (Geek), each drawn as often.
        five, two = 'Booster K1 5 Ah', 'Booster K1 2 Ah'
        joints = lambda robot: [(j.name, j.count, j.module, j.module_mass) for j in self.apart.joints  # noqa: E731
                                if j.robot == robot]
        self.assertEqual(joints(five), joints(two))
        self.assertEqual(self.apart.constants[five]['Battery capacity'][1], (4.75, 5, 5.25))
        self.assertEqual(self.apart.constants[two]['Battery capacity'][1], (1.9, 2, 2.1))

    def test_the_g1_plus_is_the_g1_with_two_head_joints(self):
        # E5.17 point 5: the G1's own parts as taken apart, and its two head joints as two small joints.
        g1 = [(j.name, j.count, j.module) for j in self.apart.joints if j.robot == 'Unitree G1']
        plus = [(j.name, j.count, j.module) for j in self.apart.joints if j.robot == 'Unitree G1+']
        self.assertEqual(plus, g1 + [('Head (yaw and pitch)', 2, 'small')])
        parts = {p.name: p for p in self.apart.parts if p.robot == 'Unitree G1+'}
        self.assertEqual(parts['Joint module, small'].count, 16)
        self.assertEqual({p.name for p in self.apart.parts if p.robot == 'Unitree G1'}, set(parts))

    def test_the_asimov_2_is_the_asimov_1_scaled_to_38_kg(self):
        # E5.17 point 7: every mass of the Asimov 1 times 38 / 35.
        one = {j.name: j.module_mass for j in self.apart.joints if j.robot == 'Menlo Asimov 1'}
        two = {j.name: j.module_mass for j in self.apart.joints if j.robot == 'Menlo Asimov 2'}
        self.assertEqual(set(one), set(two))
        for name in one:
            for a, b in zip(one[name], two[name]):
                self.assertAlmostEqual(b, a * 38 / 35, places=3, msg=name)
        self.assertEqual(self.apart.constants['Menlo Asimov 2']['Robot mass, with battery'][1], (37, 38, 39))

    def test_the_actuators_of_batch_1_at_their_modes_as_agreed(self):
        # E5.12, agreed 2026-10-29: the actuators' kg at the modes. The RP1's 18.8 kg holds the RMM8463-2E at its
        # agreed mode, 0.7 kg (point 3); E5.12's table gave 19.2 kg with it at its bound, 0.8 kg.
        agreed = {'RoboParty Roboto Origin': 17.6, 'RoboParty RP1': 18.8, 'EngineAI PM01': 15.6,
                  'ROBOTIS AI Sapiens K1': 14.5, 'Menlo Asimov 1': 18.6, 'Booster T1': 14.9, 'Booster K1 5 Ah': 9.2,
                  'Unitree R1': 14.6}
        for robot, kg in agreed.items():
            total = sum(j.count * j.module_mass[1] for j in self.apart.joints if j.robot == robot)
            self.assertAlmostEqual(total, kg, delta=0.05, msg=robot)

    def test_a_robot_without_joint_torques_has_its_joints_as_shares_adding_up_to_1(self):
        # E5.17 point 1: the joint modules a share of the robot's mass, split as batch 1's are on average.
        shared = [r for r in self.apart.robots
                  if any(j.share is not None for j in self.apart.joints if j.robot == r)]
        self.assertEqual(len(shared), 17)
        for robot in shared:
            shares = [j.share for j in self.apart.joints if j.robot == robot]
            self.assertNotIn(None, shares, msg=robot)
            self.assertAlmostEqual(sum(shares), 1, places=12, msg=robot)

    def test_every_steel_and_titanium_part_has_its_grade(self):
        # E5.18, agreed 2026-10-29: each steel part takes the grade of its use.
        grades = {c.name: c.grade for c in self.apart.joint_parts if c.material == 'Steel'}
        self.assertEqual(grades, {'Gear stage 1, sun gear': '20CrMnTi',
                                  'Gear stage 1, planets and carrier with stage 2 sun gear': '20CrMnTi',
                                  'Gear stage 1, ring gear': '20CrMnTi', 'Gear stage 2, ring gear': '20CrMnTi',
                                  'Two small bearings': 'GCr15', 'Output flange': '42CrMo4'})
        for p in self.apart.parts:
            if p.material in ('Steel', 'Titanium'):
                self.assertTrue(p.grade, msg=f'{p.robot}, {p.name}')
        for s in self.apart.shares:
            if s.material in ('Steel', 'Titanium'):
                self.assertTrue(s.grade, msg=s.item)
        stator = {s.material: s.grade for s in self.apart.shares if s.item == 'Motor stator'}
        self.assertEqual(stator['Steel'], 'M19 or M27')

    def test_what_it_writes_equals_the_data_file_on_disk(self):
        on_disk = builder.output_path()
        if not os.path.exists(on_disk):
            self.skipTest('the data file is not on disk')
        present = load_workbook(on_disk)
        self.assertEqual(self.book.sheetnames, present.sheetnames)
        for name in self.book.sheetnames:
            rows = [list(r) for r in self.book[name].iter_rows(values_only=True)]
            self.assertEqual(rows, [list(r) for r in present[name].iter_rows(values_only=True)], msg=name)

    def test_every_stated_battery_pack_holds_its_heaviest_cells_and_board(self):
        """Matthias, 2026-10-28, "Cut at the cells": a pack whose mass is given has at its lower bound room for its
        heaviest cells (its largest energy over the NMC cells' lowest energy per kg, 00's composition.nmc_wh_per_kg_min)
        and its largest management board, so that the battery housing, the pack's remainder, never falls below zero.
        Each robot is drawn at its own mass (no scaling since 2026-10-29). A pack drawn from its energy is never lighter
        than its cells: its highest energy per kg is the cells' lowest."""
        params = Params()
        g_per_cm2 = self.apart.constants[schemas.ALL_ROBOTS]['Board mass per cm2 (boards not weighed)'][1][2]
        for robot in self.apart.robots:
            pack = next(p for p in self.apart.parts if p.robot == robot and p.name == 'Battery pack')
            if None in pack.bounds:
                continue
            constants = self.apart.constants[robot]
            cells = constants['Battery voltage'][1][2] * constants['Battery capacity'][1][2] / \
                params.composition.nmc_wh_per_kg_min
            board = max(b.bounds[2] * b.count for b in self.apart.boards
                        if b.robot == robot and b.kind == 'Battery pack')
            with self.subTest(robot=robot):
                self.assertGreaterEqual(pack.bounds[0], cells + board * g_per_cm2 / 1000 - 1e-9)
        packs = {p.robot: p.bounds for p in self.apart.parts if p.name == 'Battery pack'}
        self.assertEqual((packs['AgiBot X1'], packs['Fourier N1']), ((1.75, 2.2, 3.3), (2.26, 2.9, 3.2)))
        pack_wh_per_kg = self.apart.constants[schemas.ALL_ROBOTS]['Battery pack, Wh per kg (packs not weighed)'][1]
        self.assertLessEqual(pack_wh_per_kg[2], params.composition.nmc_wh_per_kg_min)

    def test_the_g1s_hips_are_cooled_passively_with_air_from_the_hip_area(self):
        """Matthias, 2026-10-28: "Joints have no fans, just passive cooling"; Munro (T3): "The hip area uses two
        centrifugal fans" for the whole robot, blowing air around the actuators, not two per hip."""
        heading = [c.value for c in self.book['Joints'][1]].index('Cooling')
        rows = [r for r in self.book['Joints'].iter_rows(min_row=2, values_only=True) if r[0] == 'Unitree G1'
                and str(r[2]).startswith('Hip')]
        self.assertEqual(len(rows), 3)
        for row in rows:
            self.assertEqual(row[heading], 'passive, with air from two centrifugal fans in the hip area', msg=row[2])

    def test_the_magnets_grade_by_each_joints_heat(self):
        """Matthias, 2026-10-28: "Heat has to matter and so the distribution will be different depending on the joint",
        then "Accept as proposed" (EVIDENCE.md E2.11): every joint in a heat group by its name, each group a
        distribution over grade classes. The head joints and the X1's grippers are cool (the entry, point 5)."""
        grades = {}
        for g in self.apart.grades:
            grades.setdefault(g.group, {})[g.klass] = g.probability
        self.assertEqual(grades, {'Knees': {'SH': 0.3, 'UH': 0.5, 'EH': 0.2},
                                  'Hips': {'SH': 0.5, 'UH': 0.4, 'EH': 0.1},
                                  'Waist': {'SH': 0.6, 'UH': 0.4},
                                  'Ankles': {'SH': 0.6, 'UH': 0.4},
                                  'Shoulders': {'H': 0.5, 'SH': 0.5},
                                  'Elbows and wrists': {'H': 0.7, 'SH': 0.3}})
        group_of = {'hip': 'Hips', 'knee': 'Knees', 'ankle': 'Ankles', 'waist': 'Waist', 'shoulder': 'Shoulders',
                    'elbow': 'Elbows and wrists', 'wrist': 'Elbows and wrists', 'head': 'Elbows and wrists',
                    'neck': 'Elbows and wrists', 'gripper': 'Elbows and wrists'}
        for j in self.apart.joints:
            self.assertEqual(j.heat_group, group_of[j.name.split()[0].lower()], msg=f'{j.robot}, {j.name}')

    def test_it_takes_no_arguments_and_writes_where_00_says(self):
        # In its type's folder, the folder the builder sits in (2026-10-30, "Yes, build C", Q1).
        self.assertEqual(builder.output_path(), absolute(Params().data.parts_taken_apart.format(
            folder='1_compact_research_developer')))
        self.assertEqual(builder_two.output_path(), absolute(Params().data.parts_taken_apart.format(
            folder='2_full_scale_general_purpose_ai_research')))

    def test_it_carries_the_masses_and_parts_decided_on_2026_10_28(self):
        """Matthias, 2026-10-28, on the X1's and N1's masses and the G1's parts not carried over: "X1 max 35.3 kg",
        "Encoders", "XT30 connectors", "Main board cooler", then "Yes, build" (EVIDENCE.md E2.8e). Every robot has two
        encoders and two XT30 connectors per joint (a robot whose joints are shares: per joint its sources state); the
        G1, X1 and N1 a main board fan and heat sink of their own."""
        masses = {robot: self.apart.constants[robot]['Robot mass, with battery'][1]
                  for robot in ('Unitree G1', 'AgiBot X1', 'Fourier N1')}
        self.assertEqual(masses, {'Unitree G1': (34, 35, 36), 'AgiBot X1': (32, 33, 35.3),
                                  'Fourier N1': (38, 39, 40)})
        stated = {b['name']: b['joints_stated'] for b in vars(builder).values()
                  if isinstance(b, dict) and b.get('joints') == builder.AVERAGE}
        for robot in self.apart.robots:
            parts = {p.name: p for p in self.apart.parts if p.robot == robot}
            sensors = {p.name: p for p in self.apart.sensors if p.robot == robot}
            joints = stated.get(robot) or sum(j.count for j in self.apart.joints if j.robot == robot)
            with self.subTest(robot=robot):
                self.assertEqual(parts['Encoders (two per joint)'].count, 2 * joints)
                self.assertEqual(sensors['Encoder'].count, 2 * joints)
                xt30 = parts['XT30 (2+2) connectors (two per joint)']
                self.assertEqual(xt30.count, 2 * joints)
                # Each of the G1's mass; the Asimov 2's, as all of its composition, the Asimov 1's times 38 / 35.
                factor = 38 / 35 if robot == 'Menlo Asimov 2' else 1
                pair = (0.0019, 0.00212, 0.0023)   # one mated pair (EVIDENCE.md E5.22); the Asimov 2's to four decimals
                self.assertEqual(xt30.bounds, pair if factor == 1 else tuple(round(kg * factor, 4) for kg in pair))
        for robot in ('Unitree G1', 'AgiBot X1', 'Fourier N1'):
            parts = {p.name: p for p in self.apart.parts if p.robot == robot}
            self.assertEqual(parts['Main board fan'].bounds, (0.005, 0.01, 0.02))
        for robot in ('AgiBot X1', 'Fourier N1'):
            parts = {p.name: p for p in self.apart.parts if p.robot == robot}
            sink = parts['Main board heat sink: heat pipes and fins']
            self.assertEqual((sink.bounds, sink.material, sink.label), ((0.055, 0.12, 0.19), 'Copper', 'E'))



if __name__ == '__main__':
    unittest.main()


class TestPartsNotEntered(unittest.TestCase):
    """Matthias, 2026-10-29, "Yes, build them", point 4: the parts the research found and the entry left out
    (EVIDENCE.md E5.20, "Not entered") entered, each out of its robot's frame: the computer's cooler of every robot
    without one, by its computer; the RP1's 30 V auxiliary battery, its bearings outside the actuators and its 80
    magnets in the covers; the Sprout's four time-of-flight sensors."""

    # The robots whose computer has an Intel or AMD laptop-class processor: cooled as the X1's and N1's mini PCs.
    X86 = {'AKINROBOTICS AKINCI-5', 'JAKA π', 'YMBOT E3'}

    @classmethod
    def setUpClass(cls):
        cls.apart, _ = written()

    def parts(self, robot: str) -> dict:
        return {p.name: p for p in self.apart.parts if p.robot == robot}

    def test_every_robot_has_its_computers_cooler(self):
        pipes, sink = 'Main board heat sink: heat pipes and fins', 'Main board heat sink'
        for robot in self.apart.robots:
            parts = self.parts(robot)
            factor = 38 / 35 if robot == 'Menlo Asimov 2' else 1   # the Asimov 1's masses times 38 / 35
            with self.subTest(robot=robot):
                self.assertIn('Main board fan', parts)
                self.assertEqual(parts['Main board fan'].bounds, tuple(round(kg * factor, 4) for kg in
                                                                       (0.005, 0.01, 0.02)))
                if robot in self.X86 | {'AgiBot X1', 'Fourier N1'}:
                    self.assertEqual((parts[pipes].bounds, parts[pipes].material), ((0.055, 0.12, 0.19), 'Copper'))
                    self.assertNotIn(sink, parts)
                elif robot == 'Booster T1':   # a Jetson AGX Orin and an Express-i7: both coolers, two fans
                    self.assertEqual(parts['Main board fan'].count, 2)
                    self.assertEqual(parts[sink].material, 'Aluminium')
                    self.assertEqual(parts[pipes].material, 'Copper')
                else:
                    self.assertEqual((parts[sink].bounds, parts[sink].material),
                                     (tuple(round(kg * factor, 4) for kg in (0.05, 0.1, 0.2)), 'Aluminium'))
                    self.assertNotIn(pipes, parts)

    def test_the_rp1s_fans_its_auxiliary_battery_bearings_and_cover_magnets(self):
        parts = self.parts('RoboParty RP1')
        self.assertEqual(parts['Waist fan'].count, 2)            # its BOM: two waist fans and a main-board fan
        battery = parts['Auxiliary battery cells (30 V)']
        self.assertEqual((battery.bounds, battery.material, battery.kind), ((0.12, 0.25, 0.5), 'Battery materials',
                                                                            'Battery pack'))
        bearings = parts['Bearings outside the actuators (9)']
        self.assertEqual((bearings.count, bearings.material, bearings.grade), (9, 'Steel', 'GCr15'))
        self.assertAlmostEqual(9 * bearings.bounds[1], 0.8 / 21 * 9, places=3)   # each as the X1's average
        magnets = parts['Magnets in the covers, N52 (80)']
        self.assertEqual((magnets.count, magnets.material, magnets.grade), (80, 'Magnets (NdFeB)', 'N'))
        # 4 × 2 mm and 4 × 4 mm discs of NdFeB, 7.5 g/cm3: 0.188 g and 0.377 g each.
        self.assertEqual(magnets.bounds, (0.000188, 0.000283, 0.000377))

    def test_the_sprouts_time_of_flight_sensors(self):
        sensors = {e.name: e for e in self.apart.sensors if e.robot == 'Fauna Sprout'}
        tof = sensors['Time-of-flight sensor']
        self.assertEqual((tof.count, tof.bounds), (4, (5, 10, 20)))
        units = {(t, u) for _, sensor, t, u in self.apart.elements if sensor == 'Time-of-flight sensor'}
        self.assertEqual(units, {('laser diode array', 1), ('photodetector array', 1)})


class TestTheLastOpenParts(unittest.TestCase):
    """Matthias, 2026-10-29, "Yes, enter all" (EVIDENCE.md E5.22): the X1's parts by its newer parts list (its screws
    split, its 90 pins of 440C on their own, no longer counted twice), the parts found in its lists (brass inserts,
    sensor shield, silicone dampers, an XT60), the N1's network switch, and every robot's XT30 a mated pair of 2.12 g
    with brass contacts."""

    @classmethod
    def setUpClass(cls):
        cls.apart, _ = written()

    def parts(self, robot: str) -> dict:
        return {p.name: p for p in self.apart.parts if p.robot == robot}

    def test_the_x1s_screws_and_pins_by_its_newer_list(self):
        parts = self.parts('AgiBot X1')
        self.assertNotIn('Screws (about 980)', parts)
        self.assertNotIn('Stainless 17-4PH shafts and crosses; steel pins (about 115)', parts)
        carbon, stainless = parts['Screws, class 12.9 and 10.9 (677)'], parts['Screws, stainless 304 (150)']
        self.assertEqual((carbon.bounds, carbon.grade), ((0.207, 0.4145, 0.622), '42CrMo4'))
        self.assertEqual((stainless.bounds, stainless.grade), ((0.0459, 0.0918, 0.1377), 'Type 304'))
        self.assertAlmostEqual(carbon.bounds[1] + stainless.bounds[1], 0.6 * 827 / 980, places=3)   # by number
        pins = parts['Pins, SUS440C (90)']
        self.assertEqual((pins.bounds, pins.material, pins.grade, pins.label), ((0.105, 0.1086, 0.112), 'Steel', '440C',
                                                                               'S'))
        shafts = parts['Stainless 17-4PH shafts and crosses']
        self.assertEqual((shafts.bounds, shafts.material, shafts.grade), ((0.075, 0.15, 0.225), 'Steel', '17-4PH'))
        self.assertFalse([s for s in self.apart.shares if s.item.startswith('Stainless 17-4PH shafts and crosses;')])

    def test_the_x1s_parts_found_in_its_lists(self):
        parts = self.parts('AgiBot X1')
        inserts = parts['Brass threaded inserts, M3 (12) and M4 (4)']
        self.assertEqual((inserts.bounds, inserts.material, inserts.grade),
                         ((0.0065, 0.013, 0.0195), 'Copper', 'Brass'))
        shield = parts['Sensor shield, stainless 304']
        self.assertEqual((shield.bounds, shield.grade), ((0.01, 0.02, 0.03), 'Type 304'))
        dampers = parts['Silicone dampers (4)']
        self.assertEqual((dampers.count, dampers.bounds, dampers.material), (4, (0.001, 0.002, 0.003), 'Plastics'))
        xt60 = parts['XT60 connector, mated pair']
        self.assertEqual((xt60.count, xt60.bounds), (1, (0.0059, 0.0065, 0.0072)))

    def test_the_n1s_network_switch(self):
        boards = {b.name: b for b in self.apart.boards if b.robot == 'Fourier N1'}
        switch = boards['Network switch board']
        self.assertEqual((switch.bounds, switch.category, switch.label), ((30, 60, 90), 'ITC', 'E'))

    def test_every_connectors_contacts_are_brass(self):
        for item in ('XT30 (2+2) connectors (two per joint)', 'XT60 connector, mated pair'):
            contacts = [s for s in self.apart.shares if s.item == item and s.material == 'Copper']
            self.assertEqual([(s.grade, s.share) for s in contacts], [('Brass', (0.175, 0.35, 0.525))], msg=item)


class TestTypeTwo(unittest.TestCase):
    """Type 2's builder (Matthias, 2026-10-30: "Yes, agree E6.8", "Yes, build C"): what it writes passes every check of
    01, holds the five robots of batch 1 as the register names them (the GR-1 in two versions), and each number as
    E6.8 agreed it."""

    @classmethod
    def setUpClass(cls):
        cls.apart, cls.book = written(builder_two)

    def parts(self, robot: str) -> dict:
        return {p.name: p for p in self.apart.parts if p.robot == robot}

    def constant(self, robot: str, quantity: str) -> tuple:
        return self.apart.constants[robot][quantity][1]

    def test_what_it_writes_passes_every_check_of_01(self):
        params = Params()
        lists_path, recipes_path = absolute(params.data.shared_lists), absolute(params.data.element_recipes)
        if not all(os.path.exists(p) for p in (lists_path, recipes_path)):
            self.skipTest('the shared lists or the element recipes are not on disk')
        lists, recipes = schemas.read_shared_lists(lists_path), schemas.read_element_recipes(recipes_path)
        faults = (schemas.check_parts_taken_apart(self.apart, lists, params)
                  + schemas.check_sibling_data(self.apart, params) + schemas.check_grades(self.apart, recipes, params))
        self.assertEqual([f.text for f in faults], [])

    def test_its_robots_are_those_the_register_names(self):
        self.assertEqual(sorted(self.apart.robots), sorted(TYPE_TWO))
        self.assertEqual({r.type for r in self.apart.robots.values()}, {2})
        # Batch 1, five robots (the GR-1 in two versions), batch 2, 24 (the H2-X in two versions; E7.5), and batch 2b,
        # 24 (the Walker Tienkung in two versions; E8.3).
        self.assertEqual(len(TYPE_TWO), 56 + 46)   # since 2026-10-30, the robots on wheels (E9.6)
        self.assertTrue(set(BATCH_ONE_TWO) <= set(TYPE_TWO))

    def test_the_joints_masses_by_the_refit_and_the_m107(self):
        # E6.8 point 1: 0.0689 kg × torque^0.566, half to 1.5 times; the H1-2's knees the M107, 1.9 kg ± 5 %.
        joints = {(j.robot, j.name): j for j in self.apart.joints}
        hip = joints[('Unitree H2', 'Hip pitch')]
        mode = 0.0689 * 360 ** 0.566
        self.assertEqual(hip.module_mass, tuple(round(f * mode, 4) for f in (0.5, 1, 1.5)))
        self.assertAlmostEqual(hip.module_mass[1], 1.928, places=3)
        self.assertEqual(joints[('Unitree H1-2', 'Knee')].module_mass, (1.805, 1.9, 1.995))
        # Its joint modules E: 25 of its 27 joints from their torques, so 01's table and 02's notes say so.
        modules = {p.name: p.label for p in self.apart.parts if p.robot == 'Unitree H1-2'
                   and p.name.startswith('Joint module')}
        self.assertEqual(set(modules.values()), {'E'})
        self.assertIn('joint torques', schemas.estimate_kinds(self.apart, 'Unitree H1-2'))
        self.assertEqual(joints[('Unitree H1-2', 'Hip pitch')].torque, 220)           # its page's peak (point 2)
        self.assertEqual(joints[('Fourier GR-1 one camera', 'Knee')].torque, 225)     # GR1T2.urdf (decided)
        self.assertEqual(joints[('Fourier GR-1 one camera', 'Shoulder pitch')].torque, 38.46)   # gr1t2.urdf
        counts = {r: sum(j.count for j in self.apart.joints if j.robot == r) for r in BATCH_ONE_TWO}
        self.assertEqual(counts, {'Fourier GR-1 one camera': 34, 'Fourier GR-1 two cameras': 34, 'Fourier GR-2': 31,
                                  'Unitree H1-2': 27, 'Unitree H2': 31, 'Unitree H2 Plus': 33})

    def test_the_hands(self):
        # E6.8 point 4: the H2 Plus's two hands 1.65 kg for both (decided), the Fourier hands their model's 0.443 kg.
        hands = {j.robot: j for j in self.apart.joints if j.name.startswith('Hand') and j.robot in BATCH_ONE_TWO}
        self.assertEqual(set(hands), {'Unitree H2 Plus', 'Fourier GR-2', 'Fourier GR-1 one camera',
                                      'Fourier GR-1 two cameras'})
        self.assertEqual(hands['Unitree H2 Plus'].module_mass, (0.4125, 0.825, 1.2375))
        self.assertEqual(hands['Fourier GR-2'].module_mass, (0.2215, 0.443, 0.6645))
        self.assertEqual({j.module for j in hands.values()}, {'small'})

    def test_the_batteries_by_voltage_and_capacity(self):
        # E6.8 point 6: 972, 864, 950 and 483 Wh at the modes; the GR-2's pack cut at its stated 9 kg.
        wh = {r: self.constant(r, 'Battery voltage')[1] * self.constant(r, 'Battery capacity')[1]
              for r in BATCH_ONE_TWO}
        for robot, energy in (('Unitree H2', 972), ('Unitree H2 Plus', 972), ('Unitree H1-2', 864),
                              ('Fourier GR-2', 950), ('Fourier GR-1 one camera', 483)):
            self.assertAlmostEqual(wh[robot], energy, places=1, msg=robot)
        self.assertEqual(self.parts('Fourier GR-2')['Battery pack'].bounds, (4.33, 5.79, 9))

    def test_the_cables_by_height_and_the_fourier_robots_ethernet_line(self):
        # E6.8 point 9: 7 / 14 / 21 m (the GR-1 6.5 / 13 / 19.5 m); the Fourier robots two power conductors and an
        # Ethernet line of eight.
        for robot in BATCH_ONE_TWO:
            own = self.apart.constants[robot]
            cable = schemas.CABLE_ETHERNET if robot.startswith('Fourier') else schemas.CABLE_FOUR
            self.assertIn(cable, own, msg=robot)
            self.assertEqual(own[cable][1], (6.5, 13, 19.5) if 'GR-1' in robot else (7, 14, 21), msg=robot)
        self.assertEqual(self.apart.constants['all'][schemas.ETHERNET_CS][1], (0.1, 0.205, 0.326))

    def test_the_gr1s_two_versions_differ_in_their_cameras_only(self):
        one, two = 'Fourier GR-1 one camera', 'Fourier GR-1 two cameras'

        def rows(robot, items):
            return sorted((p.name, p.count, p.bounds) for p in items if p.robot == robot)

        self.assertEqual(rows(one, self.apart.parts), rows(two, self.apart.parts))
        self.assertEqual(rows(one, self.apart.boards), rows(two, self.apart.boards))
        cameras = {p.robot: p.count for p in self.apart.sensors if p.name.startswith('Depth camera')}
        self.assertEqual((cameras[one], cameras[two]), (1, 2))

    def test_screws_by_mass_and_the_h2s_titanium(self):
        # E6.8 points 11 and 12.
        self.assertEqual(self.parts('Unitree H2')['Screws'].bounds, (0.4, 1.0, 2.0))
        self.assertEqual(self.parts('Fourier GR-2')['Screws'].bounds, (0.36, 0.9, 1.8))
        titanium = self.parts('Unitree H2')['Titanium alloy parts']
        self.assertEqual((titanium.bounds, titanium.material, titanium.grade),
                         ((0.42, 0.84, 1.26), 'Titanium', 'Ti-6Al-4V'))
        self.assertNotIn('Titanium alloy parts', self.parts('Fourier GR-2'))

    def test_the_small_joints_driver_board_area_as_type_1s(self):
        # The G1's weighed board, the area every joint's board is scaled from: one copy, in both files.
        self.assertEqual(self.apart.constants['all'][schemas.DRIVER_AREA][1], (11.5, 22.9, 28.3))
        one, _ = written()
        g1 = next(b for b in one.boards if b.robot == 'Unitree G1' and b.name == 'Driver board, small joint')
        self.assertEqual(g1.bounds, (11.5, 22.9, 28.3))

    def test_what_it_writes_equals_the_data_file_on_disk(self):
        on_disk = builder_two.output_path()
        if not os.path.exists(on_disk):
            self.skipTest('the data file is not on disk')
        present = load_workbook(on_disk)
        for name in self.book.sheetnames:
            rows = [list(r) for r in self.book[name].iter_rows(values_only=True)]
            self.assertEqual(rows, [list(r) for r in present[name].iter_rows(values_only=True)], msg=name)


def fitted(torque: float) -> tuple:
    """Type 2's fit at a torque, half to 1.5 times, to four decimals (EVIDENCE.md E6.8 point 1)."""
    mode = 0.0689 * torque ** 0.566
    return tuple(round(f * mode, 4) for f in (0.5, 1, 1.5))


def half_to_one_and_a_half(mode: float) -> tuple:
    return tuple(round(f * mode, 4) for f in (0.5, 1, 1.5))


class TestTypeTwoBatchTwo(unittest.TestCase):
    """Type 2's batch 2 entered (Matthias, 2026-10-30: "Yes, all six" (E7.3), "Yes, agree E7.4", "Yes, agree E7.5",
    "Spec the entry", "Yes, build it"): its 24 robots as E7.5 agreed them, the H2-X in two versions; type 2's own
    template where a maker gives no joint torques (E7.4); the JingChu's stated knees beside the template's other
    joints (open point A); every joint named by its body part (D)."""

    # The seven robots whose joints give type 2's template (E7.4: every joint's torque; each robot once, the GR-1's
    # first version; their hands left out).
    SEVEN = ['Unitree H2', 'Unitree H1-2', 'Fourier GR-2', 'Fourier GR-1 one camera', 'Leju Kuavo 5', 'LimX Luna',
             'MagicLab MagicBot Gen1']
    # The robots whose joint modules are a share of their mass (E7.5 point 2), the H2-X in both versions.
    TEMPLATE = ['AeiROBOT ALICE 4', 'MirrorMe Bolt', 'Xiaomi CyberOne', 'Deep Robotics DR01', 'Figure AI Figure 03',
                'Unitree H2-X Core i7', 'Unitree H2-X Orin NX', 'HONOR D1', 'Simplexity I-Series',
                'ROBROS IGRIS-C', 'JCBOT JingChu', 'KAIST Humanoid', 'PAL Kangaroo', 'Kepler K1', 'Kepler K2',
                'LimX CL', 'LimX Oli', 'Lumos LUS2', 'VinMotion Motion 2', 'Zhejiang NAVAI-I3', 'Digit Nezha P01']
    GROUP = {'hip': 'Hip joints', 'knee': 'Knee joints', 'ankle': 'Ankle joints', 'waist': 'Waist joints',
             'shoulder': 'Shoulder joints', 'elbow': 'Elbow and wrist joints', 'wrist': 'Elbow and wrist joints',
             'head': 'Head joints'}
    LABEL = "as type 2's seven robots with every joint's torque on average"
    # Batch 2's robots (since batch 2b, 2026-10-30, the file holds more).
    BATCH_TWO = set(TEMPLATE) | {'DOBOT Atom', 'Leju Kuavo 5', 'LimX Luna', 'MagicLab MagicBot Gen1'}

    @classmethod
    def setUpClass(cls):
        cls.apart, cls.book = written(builder_two)

    def joints(self, robot: str) -> list:
        return [j for j in self.apart.joints if j.robot == robot]

    def constant(self, robot: str, quantity: str) -> tuple:
        return self.apart.constants[robot][quantity][1]

    def test_type_2s_template_numbers(self):
        # E7.4: the joint modules 20.2 / 40.4 / 60.6 % of the robot's mass; the battery 14.67 Wh per kg (type 2's
        # median), half to 1.5 times.
        shared = self.apart.constants['all']
        self.assertEqual(shared[schemas.JOINT_SHARE][1], (0.202, 0.404, 0.606))
        self.assertEqual(shared[schemas.ENERGY_PER_KG][1], (7.335, 14.67, 22.005))

    def test_the_template_split_is_the_seven_robots_without_their_hands(self):
        # Recomputed here from the seven robots' own rows: each group and class its share of the body's joint modules
        # at the modes, averaged over the seven; the rows named for it (open point B).
        expected: dict = {}
        for robot in self.SEVEN:
            body = [j for j in self.joints(robot) if not j.name.startswith('Hand')]
            total = sum(j.count * j.module_mass[1] for j in body)
            for j in body:
                key = (self.GROUP[j.name.split()[0].lower()], j.module)
                expected[key] = expected.get(key, 0) + j.count * j.module_mass[1] / total / len(self.SEVEN)
        rows = [j for j in self.joints('LimX CL')]
        got = {(j.name.split(',')[0], j.module): j.share for j in rows}
        self.assertEqual(set(got), set(expected))
        for key, share in expected.items():
            self.assertAlmostEqual(got[key], share, places=6, msg=key)
        self.assertAlmostEqual(sum(got.values()), 1, places=9)
        for j in rows:
            self.assertTrue(j.name.endswith(f'({self.LABEL})'), j.name)

    def test_which_robots_rest_on_the_template(self):
        shared = {j.robot for j in self.apart.joints if j.share is not None and j.robot in self.BATCH_TWO}
        self.assertEqual(shared, set(self.TEMPLATE))
        for robot in self.TEMPLATE:
            self.assertIn('joint shares', schemas.estimate_kinds(self.apart, robot), msg=robot)

    def test_the_jingchus_knees_stated_its_other_joints_the_template(self):
        # E7.5 point 2: "关节重量：2kg" each, ± 5 % (S); the template's knee rows left out, its others as for every robot
        # of the template; nothing rescaled.
        template = {(j.name, j.module): j.share for j in self.joints('LimX CL')}
        mine = self.joints('JCBOT JingChu')
        knees = [j for j in mine if j.name.split()[0] == 'Knee']
        self.assertEqual([(j.count, j.module, j.module_mass, j.share) for j in knees],
                         [(2, 'large', (1.9, 2.0, 2.1), None)])
        shares = {(j.name, j.module): j.share for j in mine if j.share is not None}
        knee_rows = {k: v for k, v in template.items() if k[0].startswith('Knee joints')}
        self.assertTrue(knee_rows)
        self.assertEqual(shares, {k: v for k, v in template.items() if k not in knee_rows})
        self.assertAlmostEqual(sum(shares.values()), 1 - sum(knee_rows.values()), places=9)
        hands = [j for j in mine if j.name.startswith('Hand')]
        self.assertEqual([(j.count, j.module_mass) for j in hands], [(2, (0.2215, 0.443, 0.6645))])

    def test_the_four_robots_with_every_joints_torque(self):
        # E7.5 point 1: type 2's fit at each joint's torque of the maker's model; the joints named by their axes (D).
        kuavo = {'Hip roll': 127, 'Hip yaw': 71, 'Hip pitch': 132, 'Knee': 280, 'Ankle pitch': 91.6,
                 'Ankle roll': 68.4, 'Waist yaw': 102, 'Shoulder pitch': 66, 'Shoulder roll': 75, 'Shoulder yaw': 57,
                 'Elbow': 75, 'Wrist yaw': 14.1, 'Wrist roll': 14.1, 'Wrist pitch': 14.1, 'Head yaw': 1.5,
                 'Head pitch': 12}
        atom = {'Hip pitch': 207.76, 'Hip roll': 241.42, 'Hip yaw': 104.16, 'Knee': 213.8, 'Ankle pitch': 89.896,
                'Ankle roll': 89.896, 'Waist yaw': 104.16, 'Shoulder pitch': 56, 'Shoulder roll': 36,
                'Shoulder yaw': 36, 'Elbow pitch': 36, 'Elbow roll': 18.6, 'Wrist pitch': 18.6, 'Wrist yaw': 18.6,
                'Head yaw': 12, 'Head pitch': 12}
        for robot, torques, body, modules in (('Leju Kuavo 5', kuavo, 29, 21.66), ('DOBOT Atom', atom, 29, 22.37),
                                              ('LimX Luna', None, 27, 23.24), ('MagicLab MagicBot Gen1', None, 30,
                                                                               32.59)):
            rows = [j for j in self.joints(robot) if not j.name.startswith('Hand')]
            with self.subTest(robot=robot):
                self.assertEqual(sum(j.count for j in rows), body)
                for j in rows:
                    self.assertEqual(j.module_mass, fitted(j.torque), msg=j.name)
                    if torques:
                        self.assertEqual(j.torque, torques[j.name], msg=j.name)
                self.assertAlmostEqual(sum(j.count * j.module_mass[1] for j in rows), modules, places=1)
        # The Atom's waist (DOBOT's file a placeholder of 1000 N·m) its own hip yaw's, E (E7.5 point 1).
        self.assertNotIn(1000, [j.torque for j in self.joints('DOBOT Atom')])

    def test_every_joint_named_by_its_body_part(self):
        for j in self.apart.joints:
            self.assertIn(j.name.split()[0].lower(), schemas.JOINT_KIND, msg=f'{j.robot}, {j.name}')

    def test_the_masses_as_agreed(self):
        # E7.5 point 5 and its table: the maker's where stated (E7.3, 2), else the register's.
        masses = {'AeiROBOT ALICE 4': (44, 45, 46), 'DOBOT Atom': (57.91, 58.91, 61.81), 'MirrorMe Bolt': (74, 75, 76),
                  'Xiaomi CyberOne': (51, 52, 53), 'Deep Robotics DR01': (79, 80, 81),
                  'Figure AI Figure 03': (60, 61, 62), 'Unitree H2-X Core i7': (65, 66, 70),
                  'Unitree H2-X Orin NX': (65, 66, 70), 'HONOR D1': (44, 45, 46),
                  'Simplexity I-Series': (59, 60, 61), 'ROBROS IGRIS-C': (55, 56, 57), 'JCBOT JingChu': (64, 65, 66),
                  'KAIST Humanoid': (74, 75, 76), 'PAL Kangaroo': (50, 57.5, 65), 'Kepler K1': (84, 85, 86),
                  'Kepler K2': (74, 75, 76), 'Leju Kuavo 5': (62.5, 63.5, 69.44), 'LimX CL': (44, 45, 46),
                  'LimX Oli': (54, 55, 55), 'LimX Luna': (55, 56, 57), 'Lumos LUS2': (56, 57, 58),
                  'MagicLab MagicBot Gen1': (66.4, 70, 71), 'VinMotion Motion 2': (60, 67.5, 75),
                  'Zhejiang NAVAI-I3': (69, 70, 71), 'Digit Nezha P01': (64, 65, 66)}
        for robot, mass in masses.items():
            self.assertEqual(self.constant(robot, 'Robot mass, with battery'), mass, msg=robot)

    def test_the_batteries(self):
        # E7.5 point 6: the maker's energy where stated (voltage × capacity at the modes); else no voltage, the energy
        # per kg of the template.
        stated = {'Unitree H2-X Core i7': 1152, 'Unitree H2-X Orin NX': 1152, 'ROBROS IGRIS-C': 816,
                  'JCBOT JingChu': 960, 'PAL Kangaroo': 976, 'LimX Oli': 478.8, 'LimX Luna': 1008, 'Lumos LUS2': 540,
                  'MagicLab MagicBot Gen1': 1350}
        for robot in set(self.TEMPLATE) | {'DOBOT Atom', 'Leju Kuavo 5', 'LimX Luna', 'MagicLab MagicBot Gen1'}:
            own = self.apart.constants[robot]
            with self.subTest(robot=robot):
                if robot in stated:
                    wh = own['Battery voltage'][1][1] * own['Battery capacity'][1][1]
                    self.assertAlmostEqual(wh, stated[robot], places=1)
                else:
                    self.assertNotIn('Battery voltage', own)
        cells = {p.robot: p.label for p in self.apart.parts if p.name == 'Battery cells'}
        self.assertEqual(cells['Lumos LUS2'], 'S')   # "Ternary lithium battery"

    def test_the_h2x_two_versions_differ_in_pc2_only(self):
        intel, orin = 'Unitree H2-X Core i7', 'Unitree H2-X Orin NX'

        def rows(robot, items, leave=()):
            return sorted((p.name, p.count, p.bounds) for p in items if p.robot == robot and p.name not in leave)

        coolers = ('Main board heat sink', 'Main board heat sink: heat pipes and fins')
        self.assertEqual(rows(intel, self.apart.parts, coolers), rows(orin, self.apart.parts, coolers))
        self.assertEqual(rows(intel, self.apart.sensors), rows(orin, self.apart.sensors))
        boards = {r: {b.name: b.bounds for b in self.apart.boards if b.robot == r} for r in (intel, orin)}
        self.assertEqual(boards[intel]['Computer board PC2, Intel Core i7'], (60, 120, 180))
        self.assertEqual(boards[orin]['Compute module PC2, NVIDIA Jetson Orin NX'], (29.74, 31.3, 32.87))
        self.assertEqual(boards[orin]['Carrier board of the Orin NX'], (60, 120, 180))
        self.assertEqual({k: v for k, v in boards[intel].items() if 'PC2' not in k},
                         {k: v for k, v in boards[orin].items() if 'PC2' not in k and 'Carrier' not in k})

    def test_the_cables_by_height_and_the_kangaroos_ethernet_line(self):
        # E7.5 point 9: 8 m per metre of height, to the half metre, half to 1.5 times; the Kangaroo's EtherCAT.
        modes = {'AeiROBOT ALICE 4': 13, 'DOBOT Atom': 13.5, 'MirrorMe Bolt': 14, 'Xiaomi CyberOne': 14,
                 'Deep Robotics DR01': 13.5, 'Figure AI Figure 03': 14, 'Unitree H2-X Core i7': 14,
                 'Unitree H2-X Orin NX': 14, 'HONOR D1': 13.5, 'Simplexity I-Series': 14, 'ROBROS IGRIS-C': 12.5,
                 'JCBOT JingChu': 13.5, 'KAIST Humanoid': 13, 'PAL Kangaroo': 13, 'Kepler K1': 14, 'Kepler K2': 14,
                 'Leju Kuavo 5': 14, 'LimX CL': 13, 'LimX Oli': 13, 'LimX Luna': 13, 'Lumos LUS2': 13,
                 'MagicLab MagicBot Gen1': 14, 'VinMotion Motion 2': 13, 'Zhejiang NAVAI-I3': 14.5,
                 'Digit Nezha P01': 11.5}
        for robot, mode in modes.items():
            own = self.apart.constants[robot]
            cable = schemas.CABLE_ETHERNET if robot == 'PAL Kangaroo' else schemas.CABLE_FOUR
            self.assertEqual(own[cable][1], half_to_one_and_a_half(mode), msg=robot)

    def test_the_hands_and_grippers_of_batch_2(self):
        # E7.5 point 4.
        six = half_to_one_and_a_half(0.443)
        eleven = half_to_one_and_a_half(0.443 * 11 / 6)
        expected = {'Leju Kuavo 5': six, 'ROBROS IGRIS-C': six, 'JCBOT JingChu': six, 'Kepler K1': eleven,
                    'Kepler K2': eleven, 'MagicLab MagicBot Gen1': half_to_one_and_a_half(1.178 / 2),
                    'Figure AI Figure 03': (0.4125, 0.825, 1.2375), 'LimX Oli': (0.41, 0.43, 0.45)}
        got = {}
        for j in self.apart.joints:
            if j.robot in self.BATCH_TWO and j.name.split()[0] in ('Hand', 'Gripper'):
                got[j.robot] = (j.count, j.module, j.module_mass)
        self.assertEqual(got, {r: (2, 'small', m) for r, m in expected.items()})

    def test_one_motion_sensor_where_the_maker_names_no_sensor(self):
        # E7.5 point 8.
        for robot in ('DOBOT Atom', 'MirrorMe Bolt', 'Deep Robotics DR01', 'HONOR D1', 'Simplexity I-Series',
                      'KAIST Humanoid', 'LimX CL', 'VinMotion Motion 2', 'Zhejiang NAVAI-I3', 'Digit Nezha P01'):
            sensors = [(p.name, p.count) for p in self.apart.sensors if p.robot == robot and p.name != 'Encoder']
            self.assertEqual(sensors, [('Motion sensor (IMU)', 1)], msg=robot)

    def test_a_compute_module_where_the_maker_names_only_its_power(self):
        # E7.5 point 7: an Orin NX-sized module and its carrier, estimates, so half to 1.5 times the Orin NX's 31.3 cm2.
        for robot in ('JCBOT JingChu', 'Kepler K1', 'Kepler K2', 'Leju Kuavo 5', 'MagicLab MagicBot Gen1'):
            boards = [b for b in self.apart.boards if b.robot == robot and b.kind == 'Boards']
            modules = [b for b in boards if b.name.startswith('Compute module')]
            carriers = [b for b in boards if b.name.startswith('Carrier board')]
            self.assertEqual([b.bounds for b in modules], [(15.65, 31.3, 46.95)], msg=robot)
            self.assertEqual([b.bounds for b in carriers], [(60, 120, 180)], msg=robot)


class TestTypeTwoBatchTwoB(unittest.TestCase):
    """Type 2's batch 2b entered (Matthias, 2026-10-30: "Yes, all six" (E8.2), the Omni "has to be included", "Yes, agree
    E8.3", "Yes, build it" with the spec's open points A to C): its 24 robots as E8.3 agreed them, the Walker Tienkung in
    two versions; the template of E7.4 unchanged (A); Qinglong's and TALOS's joints named by their place and axis (B)."""

    TORQUES = {'Unitree H1': (19, 21.67), 'OpenLoong Qinglong': (31, 30.11), 'X-Humanoid Tien Kung 2.0': (30, 24.79),
               'Walker Tienkung Vision': (21, 21.00), 'Walker Tienkung Embodied': (30, 24.06), 'PAL TALOS': (30, 26.09)}
    TEMPLATE = ['X-Humanoid Omni', 'Tesla Optimus Gen 3', 'Tesla Optimus Gen 2', 'Tesla Optimus V2.5',
                'Cyan Robotics Orca', 'Sanctuary Phoenix', 'PHYBOT M1', 'PAL REEM-C', 'Robotera L7', 'EngineAI SE01',
                'Zeroth Sean', 'Siasun Songxing', 'Unitree Superman', 'Westwood THEMIS', 'X-Humanoid TianGong Ultra',
                'Tlibot Tianzhen Z1', 'X-Humanoid Tien Kung 3.0', 'Tlibot T1', 'DroidUp Walker II']

    @classmethod
    def setUpClass(cls):
        cls.apart, _ = written(builder_two)

    def joints(self, robot: str) -> list:
        return [j for j in self.apart.joints if j.robot == robot]

    def test_the_template_unchanged(self):
        # Open point A: E7.4's numbers and its seven robots; batch 2b's robots of the template take its rows as they are.
        shared = self.apart.constants['all']
        self.assertEqual(shared[schemas.JOINT_SHARE][1], (0.202, 0.404, 0.606))
        rows = {(j.name, j.share) for j in self.joints('LimX CL')}
        for robot in self.TEMPLATE:
            self.assertEqual({(j.name, j.share) for j in self.joints(robot) if j.share is not None}, rows, msg=robot)

    def test_the_robots_with_every_joints_torque(self):
        for robot, (count, modules) in self.TORQUES.items():
            body = [j for j in self.joints(robot) if j.name.split()[0] not in ('Hand', 'Gripper')]
            with self.subTest(robot=robot):
                self.assertEqual(sum(j.count for j in body), count)
                self.assertAlmostEqual(sum(j.count * j.module_mass[1] for j in body), modules, places=1)
                self.assertTrue(all(j.share is None for j in self.joints(robot)))
        knees = [j for j in self.joints('Unitree H1') if j.name.startswith('Knee')]
        self.assertEqual([(j.count, j.module_mass) for j in knees], [(2, (1.805, 1.9, 1.995))])   # the M107 (E8.3, 1)
        self.assertEqual(max(j.torque for j in self.joints('PAL TALOS') if j.torque), 300)          # LAAS's model
        waist = [j for j in self.joints('Walker Tienkung Vision') if j.name.startswith('Waist')]
        self.assertEqual([(j.torque, j.module_mass) for j in waist], [(180, fitted(180))])       # E: its hip roll's

    def test_every_joint_named_by_its_body_part(self):
        for j in self.apart.joints:
            self.assertIn(j.name.split()[0].lower(), schemas.JOINT_KIND, msg=f'{j.robot}, {j.name}')

    def test_the_masses_as_agreed(self):
        masses = {'X-Humanoid Omni': (38.5, 39.5, 40.5), 'Tesla Optimus Gen 3': (51, 52, 53),
                  'Tesla Optimus Gen 2': (55, 56, 57), 'Tesla Optimus V2.5': (54, 55, 56), 'Cyan Robotics Orca': (49, 50, 51),
                  'Sanctuary Phoenix': (69, 70, 71), 'PHYBOT M1': (59, 60, 61), 'OpenLoong Qinglong': (85.14, 86.14, 87.14),
                  'PAL REEM-C': (79, 80, 81), 'Robotera L7': (66.2, 67.2, 68.2), 'EngineAI SE01': (79, 80, 81),
                  'Zeroth Sean': (44, 45, 46), 'Siasun Songxing': (84, 85, 86), 'Unitree Superman': (44, 45, 46),
                  'PAL TALOS': (93.34, 95, 96), 'Westwood THEMIS': (36.92, 37.92, 38.92),
                  'X-Humanoid TianGong Ultra': (54, 55, 56), 'Tlibot Tianzhen Z1': (44, 45, 46),
                  'X-Humanoid Tien Kung 2.0': (66.96, 67.96, 68.96), 'X-Humanoid Tien Kung 3.0': (63.5, 64.5, 65.5),
                  'Tlibot T1': (42, 43, 44), 'DroidUp Walker II': (31, 32, 33), 'Walker Tienkung Vision': (60, 61, 62),
                  'Walker Tienkung Embodied': (72, 73, 74), 'Unitree H1': (46, 47, 59.34)}
        for robot, mass in masses.items():
            self.assertEqual(self.apart.constants[robot]['Robot mass, with battery'][1], mass, msg=robot)

    def test_the_batteries(self):
        stated = {'Unitree H1': 864, 'Robotera L7': 900, 'Walker Tienkung Vision': 1584, 'Walker Tienkung Embodied': 1584,
                  'Tlibot T1': 1110, 'X-Humanoid Tien Kung 3.0': 864}
        for robot in self.TEMPLATE + list(self.TORQUES):
            own = self.apart.constants[robot]
            with self.subTest(robot=robot):
                if robot in stated:
                    self.assertAlmostEqual(own['Battery voltage'][1][1] * own['Battery capacity'][1][1], stated[robot],
                                           places=1)
                else:
                    self.assertNotIn('Battery voltage', own)

    def test_hands_and_grippers(self):
        expected = {'Robotera L7': (1.045, 1.1, 1.155), 'EngineAI SE01': half_to_one_and_a_half(0.443),
                    'Walker Tienkung Embodied': half_to_one_and_a_half(0.443),
                    'Westwood THEMIS': half_to_one_and_a_half(0.463), 'PAL TALOS': half_to_one_and_a_half(1.094)}
        got = {j.robot: j.module_mass for j in self.apart.joints if j.robot in self.TEMPLATE + list(self.TORQUES)
               and j.name.split()[0] in ('Hand', 'Gripper')}
        self.assertEqual(got, expected)

    def test_the_cables_by_height_and_talos_ethernet_line(self):
        modes = {'X-Humanoid Omni': 11, 'Unitree H1': 14.5, 'PAL TALOS': 14, 'Walker Tienkung Vision': 14,
                 'Robotera L7': 13.5, 'Zeroth Sean': 12.5}
        for robot, mode in modes.items():
            cable = schemas.CABLE_ETHERNET if robot == 'PAL TALOS' else schemas.CABLE_FOUR
            self.assertEqual(self.apart.constants[robot][cable][1], half_to_one_and_a_half(mode), msg=robot)

    def test_the_walker_tienkungs_two_versions(self):
        vision, embodied = 'Walker Tienkung Vision', 'Walker Tienkung Embodied'
        boards = {r: {b.name: b.count for b in self.apart.boards if b.robot == r and b.kind == 'Boards'}
                  for r in (vision, embodied)}
        self.assertEqual(boards[vision]['Compute module, NVIDIA Jetson AGX Orin'], 1)
        self.assertEqual(boards[embodied]['Compute module, NVIDIA Jetson AGX Orin'], 2)
        extra = {j.name for j in self.joints(embodied)} - {j.name for j in self.joints(vision)}
        self.assertEqual({n.split()[0] for n in extra}, {'Head', 'Wrist', 'Hand'})

    def test_the_compute_modules_read_from_their_power(self):
        # E8.3 point 7: a "2070" controller as the H2 Plus's Thor, the L7's "275 TOPS" Orin as an AGX Orin (E).
        for robot in ('X-Humanoid Omni', 'X-Humanoid Tien Kung 3.0', 'Robotera L7'):
            modules = [b.bounds for b in self.apart.boards if b.robot == robot and b.name.startswith('Compute module')]
            self.assertEqual(modules, [(43.5, 87, 130.5)], msg=robot)

    def test_one_motion_sensor_where_the_maker_names_none(self):
        for robot in ('Tesla Optimus Gen 3', 'Cyan Robotics Orca', 'Sanctuary Phoenix', 'PHYBOT M1', 'PAL REEM-C',
                      'Zeroth Sean', 'Siasun Songxing', 'Unitree Superman', 'Westwood THEMIS',
                      'X-Humanoid TianGong Ultra', 'Tlibot Tianzhen Z1', 'Tlibot T1', 'OpenLoong Qinglong',
                      'X-Humanoid Tien Kung 2.0'):
            sensors = [(p.name, p.count) for p in self.apart.sensors if p.robot == robot and p.name != 'Encoder']
            self.assertEqual(sensors, [('Motion sensor (IMU)', 1)], msg=robot)

# Type 2's robots on wheels (EVIDENCE.md E9.6, agreed 2026-10-30, "Yes, agree E9.6"; entered on "Yes, build it", E9.7):
# each as (its mass, its driven wheels and the hub motor's mode in kg, its steered wheels, its casters, its lift column
# ('stated', a share of its mass, or none), its battery's energy in Wh where stated or derived, else None).
ON_WHEELS = {
    'AgiBot A2-W': ((229, 230, 231), 2, 8.5, 0, 2, None, None),
    'AgiBot G1': ((149, 150, 151), 2, 3.75, 0, 2, None, None),
    'ROBOTIS AI Worker': ((89, 90, 94.42), 3, 3.8, 3, 0, 'stated', 2040),
    'Flexiv AICO 2-4': ((280, 281, 282), 2, 8.5, 0, 2, None, None),
    'Flexiv AICO 2-10': ((394, 395, 396), 2, 8.5, 0, 2, None, None),
    'PAL ARI': ((59, 60, 61), 2, 1.4, 0, 2, None, 960),
    'KIT ARMAR-7': ((119, 120, 121), 3, 3.75, 0, 0, None, None),
    'Astribot S1': ((79, 80, 81), 2, 3.75, 0, 2, None, None),
    'Astribot T1': ((65, 66, 67), 2, 1.4, 0, 2, None, None),
    'ELU.AI AstroDroid AD-01': ((119, 120, 121), 2, 3.75, 0, 2, None, None),
    'Honda Avatar': ((109, 110, 111), 2, 3.75, 0, 2, None, None),
    'CASBOT W1': ((259, 260, 261), 2, 8.5, 0, 4, 'share', 1200),
    'Toyota CUE7': ((73, 74, 75), 2, 3.75, 0, 0, None, None),
    'RealMan Dual-Arm Lift': ((156, 157, 158), 2, 4.85, 0, 2, 'share', None),
    'Tangible Robots Eggie': ((44, 45, 46), 2, 1.4, 0, 2, None, None),
    'Feather Robotics Feather': ((104, 105, 106), 3, 1.4, 3, 0, 'share', 2400),
    'Unitree G1-D Ultimate': ((89, 90, 91), 2, 3.75, 0, 2, 'share', 1872),
    'AgiBot G2': ((184, 185, 186), 4, 3.75, 4, 0, None, 1652),
    'Unitree H2-D Flagship': ((119, 120, 121), 2, 3.75, 0, 2, 'share', 1440),
    'MagicLab MagicBot D1': ((149, 150, 151), 4, 3.75, 4, 0, None, 1652),
    'Blue Worm Mantis Standard': ((49, 50, 51), 4, 1.4, 0, 0, 'share', None),
    'Spirit AI Moz1': ((69, 70, 71), 4, 1.4, 0, 0, None, None),
    'Sanctuary Phoenix Gen 8': ((69, 70, 71), 2, 1.4, 0, 2, None, None),
    'Galaxea R1 Pro': ((125, 126, 127), 3, 3.75, 3, 0, None, 1680),
    'Unitree R1-A5-D': ((29, 30, 31), 4, 1.4, 0, 0, None, None),
    'Unitree R1-A7-D': ((31, 32, 33), 4, 1.4, 0, 0, None, None),
    'Rainbow Robotics RB-Y1 A': ((130, 131, 132.22), 2, 3.75, 0, 2, None, 1250),
    'Rainbow Robotics RB-Y1 M': ((169, 170, 171), 4, 3.75, 0, 0, None, 2500),
    'Pollen Robotics Reachy 2': ((49, 50, 51), 3, 2.9, 0, 0, None, 840),
    'RealMan RealBot-01': ((124, 125, 126), 2, 3.75, 0, 2, None, 1920),
    'RealMan RealBot-L2': ((119, 120, 121), 2, 3.75, 0, 2, 'share', 960),
    'RealMan RealBot-S2': ((124, 125, 126), 4, 1.4, 4, 0, None, 960),
    'PAL TIAGo Pro': ((79, 80, 95), 4, 1.4, 0, 0, 'share', 720),
    'Tokyo Robotics Torobo': ((119, 120, 121), 4, 1.4, 0, 0, None, None),
    'Lanxin VB1-I': ((60, 65, 70), 2, 1.4, 0, 2, None, None),
    'Lanxin VB2': ((65, 70, 75), 4, 1.4, 0, 0, None, None),
    'Dexmate Vega 1 Pro': ((134, 135, 141.0), 2, 3.7, 2, 2, None, None),
    'Lanxin VersaBot VB-1': ((74, 75, 76), 4, 1.4, 0, 0, None, None),
    'Yue Quan Bionics W-Bot': ((59, 60, 61), 2, 1.4, 0, 2, None, 2400),
    'Tlibot Y1': ((149, 150, 151), 2, 3.75, 0, 2, None, None),
    'Tlibot Y2-B': ((84, 85, 86), 2, 3.75, 0, 2, None, None),
    'Tlibot Y2-P': ((94, 95, 96), 2, 3.75, 0, 2, None, None),
    'Yunmu YMBot D': ((99, 100, 101), 2, 3.75, 0, 2, None, None),
    'Yunmu YMBot D2': ((149, 150, 151), 2, 3.75, 0, 2, None, None),
    'Yunmu YMBot D4': ((99, 100, 101), 2, 3.75, 0, 2, None, None),
    'PsiBot ψ-SynRobot': ((64, 65, 66), 2, 1.4, 0, 2, None, None),
}
# The four robots on wheels whose makers give every joint (E9.6 point 2), their body joints' modules at the modes (kg),
# the template's source; and the W-Bot, by its page's torques.
OWN_ON_WHEELS = {'Dexmate Vega 1 Pro': 17.83, 'Rainbow Robotics RB-Y1 A': 18.17, 'Rainbow Robotics RB-Y1 M': 18.17,
                 'ROBOTIS AI Worker': 12.867, 'Pollen Robotics Reachy 2': 10.45, 'Yue Quan Bionics W-Bot': 18.72}
OUTSIDE_THE_BODY = ('Wheel', 'Steering', 'Lift', 'Hand', 'Gripper')


class TestTypeTwoOnWheels(unittest.TestCase):
    """Type 2's 43 robots on wheels entered (Matthias, 2026-10-30: "Yes, agree E9.4", "Yes, agree E9.6", "Spec the
    entry", "Yes, build it" with the spec's open points A to D as recommended): 46 blocks with the AICO 2, the RB-Y1 and
    the R1-D in two versions; their drive wheels, steering, casters and lift columns; the wheeled template; the CASBOT
    W1's LFP cells; the kind "Torso boards" renamed "Boards"."""

    @classmethod
    def setUpClass(cls):
        cls.apart, _ = written(builder_two)

    def joints(self, robot: str) -> list:
        return [j for j in self.apart.joints if j.robot == robot]

    def test_the_46_are_the_registers(self):
        self.assertTrue(set(ON_WHEELS) <= set(self.apart.robots))
        self.assertTrue(set(ON_WHEELS) <= set(TYPE_TWO))
        self.assertEqual(len(TYPE_TWO), 56 + 46)

    def test_the_masses_as_agreed(self):
        for robot, (mass, *_rest) in ON_WHEELS.items():
            self.assertEqual(self.apart.constants[robot]['Robot mass, with battery'][1], mass, msg=robot)

    def test_drive_wheels_steering_and_casters(self):
        for robot, (_m, driven, hub, steered, casters, _l, _w) in ON_WHEELS.items():
            with self.subTest(robot=robot):
                wheels = [j for j in self.joints(robot) if j.name.startswith('Wheel')]
                self.assertEqual(sum(j.count for j in wheels), driven)
                self.assertEqual({j.module_mass[1] for j in wheels}, {hub})
                self.assertEqual({(j.heat_group, j.region) for j in wheels}, {('Hips', 'Base')})   # open point A
                steering = [j for j in self.joints(robot) if j.name.startswith('Steering')]
                self.assertEqual(sum(j.count for j in steering), steered)
                self.assertTrue(all(j.heat_group == 'Waist' for j in steering))
                found = [p for p in self.apart.parts if p.robot == robot and p.name == 'Casters']
                self.assertEqual(sum(p.count for p in found), casters)
                self.assertTrue(all(p.kind == 'Wheeled base' and p.region == 'Base' for p in found))

    def test_the_lift_columns(self):
        shared = self.apart.constants['all']
        self.assertEqual(shared['Lift column, share of the robot mass, on wheels'][1], (0.0995, 0.199, 0.2985))
        for robot, (*_x, lift, _w) in ON_WHEELS.items():
            with self.subTest(robot=robot):
                column = [p for p in self.apart.parts if p.robot == robot and p.kind == 'Lift column']
                motor = [j for j in self.joints(robot) if j.name.startswith('Lift')]
                self.assertEqual(len(column), 0 if lift is None else 1)
                self.assertEqual(sum(j.count for j in motor), 0 if lift is None else 1)
                if lift == 'stated':
                    self.assertEqual(column[0].bounds, half_to_one_and_a_half(17.01))
                elif lift == 'share':
                    self.assertTrue(column[0].computed.startswith(schemas.ROBOT_SHARE))
                    self.assertIn('Lift column, share of the robot mass, on wheels', column[0].computed)

    def test_the_wheeled_template(self):
        shared = self.apart.constants['all']
        self.assertEqual(shared['Joint modules, share of the robot mass, on wheels'][1], (0.078, 0.156, 0.234))
        self.assertEqual(shared['Battery energy per kg of robot, on wheels'][1], (4.62, 10.28, 22.86))
        template = [r for r in ON_WHEELS if r not in OWN_ON_WHEELS]
        first = {(j.name, j.share) for j in self.joints(template[0]) if j.share is not None}
        groups: dict = {}
        for name, share in first:
            groups[name.split(',')[0]] = groups.get(name.split(',')[0], 0) + share
            self.assertIn("type 2's four robots on wheels", name)
        self.assertAlmostEqual(sum(groups.values()), 1, places=9)
        for group, share in {'Waist joints': 0.205, 'Shoulder joints': 0.383, 'Elbow and wrist joints': 0.336,
                             'Head joints': 0.076}.items():
            self.assertAlmostEqual(groups[group], share, delta=0.002, msg=group)
        for robot in template:
            self.assertEqual({(j.name, j.share) for j in self.joints(robot) if j.share is not None}, first, msg=robot)

    def test_the_robots_with_every_joint(self):
        for robot, modules in OWN_ON_WHEELS.items():
            body = [j for j in self.joints(robot) if not j.name.startswith(OUTSIDE_THE_BODY)]
            with self.subTest(robot=robot):
                self.assertTrue(all(j.share is None for j in self.joints(robot)))
                self.assertAlmostEqual(sum(j.count * j.module_mass[1] for j in body), modules, delta=0.01)

    def test_the_batteries(self):
        for robot, (*_x, wh) in ON_WHEELS.items():
            own = self.apart.constants[robot]
            with self.subTest(robot=robot):
                if wh is None:
                    self.assertNotIn('Battery voltage', own)
                else:
                    self.assertAlmostEqual(own['Battery voltage'][1][1] * own['Battery capacity'][1][1], wh, places=0)
                packs = [p for p in self.apart.parts if p.robot == robot and p.kind == 'Battery pack']
                self.assertEqual({p.region for p in packs}, {'Base'})
        reachy = [p for p in self.apart.parts if p.robot == 'Pollen Robotics Reachy 2' and p.name == 'Battery pack']
        self.assertEqual(reachy[0].bounds, (6.175, 6.5, 6.825))   # the 2021 base's "6.5kg", ± 5 %

    def test_encoders_count_each_hand_and_wheel(self):
        # Two encoders per joint, its hands or grippers one each, as batch 1's (template()): the AgiBot G1's "20 DoF", two
        # hands and two drive wheels; the Honda Avatar's 22 body joints, two hands and two wheels; the CASBOT W1's 17
        # body joints, two grippers, two drive wheels and its lift.
        for robot, joints in (('AgiBot G1', 20 + 2 + 2), ('Honda Avatar', 22 + 2 + 2), ('CASBOT W1', 17 + 2 + 2 + 1)):
            encoders = [p for p in self.apart.parts if p.robot == robot and p.name == 'Encoders (two per joint)']
            self.assertEqual(encoders[0].count, 2 * joints, msg=robot)

    def test_lfp_cells_for_the_casbot_w1_only(self):
        cells = {p.robot: p.grade for p in self.apart.parts if p.name == 'Battery cells'}
        self.assertEqual({r for r, g in cells.items() if g}, {'CASBOT W1'})
        self.assertEqual(cells['CASBOT W1'], 'LFP cylindrical')
        # Its rows say how 02 draws them (Matthias, 2026-10-30, "Ratio, as recommended", E9.8): its cells over a cell of
        # their recipe, its pack its cells times the NMC packs' ratio; every other robot's as before.
        rows = {(p.robot, p.name): p.computed for p in self.apart.parts}
        self.assertIn('a cell of its recipe', rows[('CASBOT W1', 'Battery cells')])
        self.assertIn("the NMC packs' pack-to-cell ratio", rows[('CASBOT W1', 'Battery pack')])
        others = {(robot, name): text for (robot, name), text in rows.items() if robot != 'CASBOT W1'}
        self.assertTrue(all("NMC cell's Wh/kg" in text for (_, name), text in others.items() if name == 'Battery cells'))
        self.assertFalse(any('pack-to-cell' in text for text in others.values()))

    def test_no_kind_torso_boards_left(self):
        for which in (builder, builder_two):
            apart, _ = written(which)
            kinds = {p.kind for p in apart.parts + apart.boards + apart.sensors}
            self.assertNotIn('Torso boards', kinds)
            self.assertIn('Boards', kinds)

    def test_every_joint_goes_to_a_kind(self):
        for j in self.apart.joints:
            self.assertIn(j.name.split()[0].lower(), schemas.JOINT_KIND, msg=f'{j.robot}, {j.name}')


# Types 3 and 4's builders (Matthias, 2026-10-30, "Yes, build it": EVIDENCE.md E11.7, open points A to C as
# recommended): batch 1's 27 robots, each in the file of its first type among 3 and 4.
builder_three = load('make_parts_taken_apart_3', os.path.join('tools', '3_industrial_logistics_warehousing',
                                                             'make_parts_taken_apart.py'))
builder_four = load('make_parts_taken_apart_4', os.path.join('tools', '4_service_healthcare_interaction',
                                                            'make_parts_taken_apart.py'))
TYPE_THREE_BATCH_1 = {
    'NEURA 4NE1 Gen 3': (79, 80, 81), 'NEURA 4NE1 Gen 3.5': (79, 80, 81), 'NEURA 4NE1 Gen 4': (63, 64, 65),
    'AgiBot A2 Max': (84, 85, 86), 'PNDbotics Adam Lite': (58.48, 60, 61), 'PNDbotics Adam SP': (60.25, 62, 63),
    'Hexagon AEON': (59, 60, 61), 'Agile Robots Agile One': (68, 69, 70), 'AIDOL Lab AIDOL': (94, 95, 96),
    'AeiROBOT ALICE M1': (96, 97, 98), 'TeknTrash ALPHA': (89, 90, 91), 'AI² Robotics AlphaBot 2': (134, 135, 136),
    'Anigon': (119, 120, 121), 'Apptronik Apollo': (71, 72, 73), 'Apptronik Apollo 2': (74, 75, 76),
    'Boston Dynamics Atlas': (89, 90, 91), 'NEUBILITY Billi': (109, 110, 111), 'Borg Robotics Borg 01': (61, 62, 63),
    'Wandercraft CALVIN-40': (79, 80, 81)}
TYPE_FOUR_BATCH_1 = {
    '1X NEO': (28.9, 29.9, 30.9), 'AgiBot A2': (68, 69, 70), 'AgiBot A3 Ultra': (59, 60, 61),
    'Andromeda Abi': (49, 50, 51), 'AgiBot A3': (54, 55, 56), 'AgiBot X2': (34, 35, 36),
    'AgiBot X2 Ultra': (38, 39, 40), 'DOBOT Atom 153 cm': (57.91, 58.91, 61.81), 'Vbot ATOM': (49, 50, 51)}
TYPES_34_ON_WHEELS = {'Andromeda Abi', 'AeiROBOT ALICE M1', 'AI² Robotics AlphaBot 2', 'Anigon', 'NEUBILITY Billi'}


class TestTypesThreeAndFour(unittest.TestCase):
    """Types 3 and 4's batch 1 entered (E11.6 agreed, "Yes, agree E11.6"; E11.7, "Yes, build it")."""

    @classmethod
    def setUpClass(cls):
        cls.three, _ = written(builder_three)
        cls.four, _ = written(builder_four)

    def test_each_robot_in_the_file_of_its_first_type_among_3_and_4(self):
        self.assertEqual(list(self.three.robots), list(TYPE_THREE_BATCH_1))
        self.assertEqual(list(self.four.robots), list(TYPE_FOUR_BATCH_1))
        self.assertEqual({r.type for r in self.three.robots.values()}, {3})
        self.assertEqual({r.type for r in self.four.robots.values()}, {4})
        self.assertFalse(set(self.three.robots) & set(self.four.robots))
        two, _ = written(builder_two)
        self.assertFalse((set(self.three.robots) | set(self.four.robots))
                         & (set(two.robots) | set(written()[0].robots)))

    def test_what_they_write_passes_every_check_of_01(self):
        lists = schemas.read_shared_lists(absolute(Params().data.shared_lists))
        for apart in (self.three, self.four):
            self.assertEqual([f.text for f in schemas.check_parts_taken_apart(apart, lists, Params())], [])

    def test_the_masses_as_decided(self):
        for apart, masses in ((self.three, TYPE_THREE_BATCH_1), (self.four, TYPE_FOUR_BATCH_1)):
            for robot, mass in masses.items():
                self.assertEqual(apart.constants[robot]['Robot mass, with battery'][1], mass, msg=robot)
        self.assertEqual(self.three.constants['NEURA 4NE1 Gen 4']['Robot mass, with battery'][3], 'S, E')

    def test_their_templates_under_names_of_their_own(self):
        for apart in (self.three, self.four):
            shared = apart.constants['all']
            self.assertEqual(shared['Joint modules, share of the robot mass, types 3 and 4'][1],
                             (0.2235, 0.447, 0.6705))
            self.assertEqual(shared['Joint modules, share of the robot mass, types 3 and 4, on wheels'][1],
                             (0.0695, 0.139, 0.2085))
            self.assertEqual(shared['Battery energy per kg of robot, types 3 and 4'][1], (7.335, 14.67, 22.005))
            self.assertEqual(shared['Battery energy per kg of robot, types 3 and 4, on wheels'][1],
                             (5.42, 10.84, 16.26))
            self.assertNotIn('Joint modules, share of the robot mass', shared)
        # The median of the seven robots on legs, at the modes, is the template's mode (E11.6 point 1).
        module = sys.modules[builder_three.SHARED.__name__]
        self.assertAlmostEqual(module.median_share(module.SEVEN), 0.447, delta=0.0005)

    def test_the_adam_by_its_actuators_and_the_atom_as_type_twos(self):
        knee = [j for j in self.three.joints if j.robot == 'PNDbotics Adam Lite' and j.name == 'Knee']
        self.assertEqual(knee[0].module_mass, (3.8, 4.0, 4.2))
        shoulder = [j for j in self.three.joints if j.robot == 'PNDbotics Adam Lite' and j.name == 'Shoulder pitch']
        self.assertAlmostEqual(shoulder[0].module_mass[1], 0.0689 * 60 ** 0.566, places=3)
        mine = [(j.name, j.count, j.torque) for j in self.four.joints if j.robot == 'DOBOT Atom 153 cm']
        two, _ = written(builder_two)
        self.assertEqual(mine, [(j.name, j.count, j.torque) for j in two.joints if j.robot == 'DOBOT Atom'])

    def test_the_batteries(self):
        def wh(apart, robot):
            own = apart.constants[robot]
            return own['Battery voltage'][1][1] * own['Battery capacity'][1][1]
        for apart, robot, energy in ((self.four, '1X NEO', 842), (self.four, 'AgiBot A2', 700),
                                     (self.four, 'AgiBot A3', 1152), (self.four, 'AgiBot A3 Ultra', 1152),
                                     (self.four, 'AgiBot X2', 500), (self.three, 'PNDbotics Adam Lite', 1172),
                                     (self.three, 'AI² Robotics AlphaBot 2', 1440)):
            self.assertAlmostEqual(wh(apart, robot), energy, delta=1, msg=robot)
        self.assertNotIn('Battery voltage', self.three.constants['Boston Dynamics Atlas'])

    def test_on_wheels_and_the_magnesium_frames(self):
        both = {**{n: self.three for n in self.three.robots}, **{n: self.four for n in self.four.robots}}
        wheeled = {n for n, a in both.items() if schemas.on_wheels(a, n)}
        self.assertEqual(wheeled, TYPES_34_ON_WHEELS)
        self.assertTrue(any(p.robot == 'Anigon' and p.kind == 'Lift column' for p in self.three.parts))
        frames = {p.robot: p.material for p in self.four.parts if p.name == 'Frame'}
        self.assertEqual((frames['AgiBot A3'], frames['AgiBot A3 Ultra']), ('Magnesium', 'Magnesium'))
