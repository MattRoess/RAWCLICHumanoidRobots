"""
Tests of tools/make_humanoid_register.py, the builder of data/researched/humanoid_register.xlsx.

**Copyright notice:** Copyright © 2026 Empa, Matthias Roesslein
**Licence:** CC BY 4.0, Creative Commons Attribution 4.0 International; the full text is in `LICENSE`
at the root of the repository.

Matthias, 2026-10-28: "Yes, code it in 02", on the plan whose first step is the register as a data file built in
tools/ with its test. The builder is the master of the register: what it writes must pass every check 01 makes of it,
hold the decisions of EVIDENCE.md E4 to E4.11 as counted there (VERIFICATION.md V60), and, while nobody has changed it,
equal the file on disk cell by cell. Everything is written into temporary folders: press Run, nothing in the project
is changed.
"""
import importlib.util
import os
import sys
import tempfile
import unittest
from collections import Counter

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)


def load(name: str, path: str):
    """Import a file of the project whose name starts with a digit or lies outside a package."""
    spec = importlib.util.spec_from_file_location(name, os.path.join(ROOT, path))
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


schemas = load('schemas', '01_schemas.py')
builder = load('make_humanoid_register', os.path.join('tools', 'make_humanoid_register.py'))
# The robots of each type's parts file, as its builder writes them: the names the register's "Taken apart as" must
# use (since 2026-10-30 one file and builder per type, each in its type's folder).
BUILDERS = {1: load('make_parts_taken_apart_1', os.path.join('tools', '1_compact_research_developer',
                                                            'make_parts_taken_apart.py')),
            2: load('make_parts_taken_apart_2', os.path.join('tools', '2_full_scale_general_purpose_ai_research',
                                                            'make_parts_taken_apart.py')),
            3: load('make_parts_taken_apart_3', os.path.join('tools', '3_industrial_logistics_warehousing',
                                                            'make_parts_taken_apart.py')),
            4: load('make_parts_taken_apart_4', os.path.join('tools', '4_service_healthcare_interaction',
                                                            'make_parts_taken_apart.py'))}
APART = {n: [row[0] for row in b.ROBOTS_ROWS] for n, b in BUILDERS.items()}
from src.params_schema import Params, absolute  # noqa: E402


class TestMakeHumanoidRegister(unittest.TestCase):
    """What the builder writes holds together as 01 checks it, counts as decided, and equals the file on disk."""

    @classmethod
    def setUpClass(cls):
        params = Params()
        if not os.path.exists(absolute(params.data.humanoid_guide)):
            raise unittest.SkipTest('the catalogue is not on disk')
        cls.folder = tempfile.TemporaryDirectory()
        cls.path = os.path.join(cls.folder.name, 'humanoid_register.xlsx')
        builder.build(cls.path, params)
        cls.register = schemas.read_register(cls.path)

    @classmethod
    def tearDownClass(cls):
        cls.folder.cleanup()

    def test_what_it_writes_passes_every_check_of_01(self):
        self.assertEqual([f.text for f in schemas.check_register(self.register, Params(), APART)], [])
        self.assertEqual([f.text for f in schemas.check_type_sheets(self.path, self.register, Params(), APART)], [])

    def test_every_robot_once_as_decided(self):
        # EVIDENCE.md E4.9 to E4.11 and V60: 349 entries of the catalogue and 5 added; 10 second entries of one robot;
        # 18 not assigned, the FF Nova below 15 kg (decided 2026-10-28) and the TRON 2, sold in no form with legs and
        # arms (E5.17, agreed 2026-10-29); 324 robots with a type, 137 with two.
        self.assertEqual([r.number for r in self.register], list(range(1, 355)))
        typed = [r for r in self.register if r.types]
        self.assertEqual(len(typed), 324)
        self.assertEqual(sum(len(r.types) == 2 for r in typed), 137)
        self.assertEqual(sum(r.entry.startswith(schemas.SAME_ROBOT) for r in self.register), 10)
        self.assertEqual(sum(r.entry.startswith(schemas.NOT_ASSIGNED) for r in self.register), 20)
        tron = next(r for r in self.register if r.number == 309)
        self.assertEqual(tron.entry, 'not assigned: no form with legs and arms')

    def test_the_robots_in_each_type(self):
        # V60's recount with the five added (1: 33 ...), less the FF Nova and the TRON 2, which leave type 1.
        counts = {t: sum(t in r.types for r in self.register) for t in range(1, 7)}
        self.assertEqual(counts, {1: 31, 2: 97, 3: 181, 4: 135, 5: 12, 6: 5})

    def test_every_robot_of_type_1_has_its_own_rows_in_the_parts_file(self):
        # Matthias, 2026-10-29: "Yes, build it" (the entry, point 2): every robot of type 1 its own rows; the K1 in its
        # two versions, each drawn as often (point 3). Since 2026-10-30 also the five of type 2's batch 1 and the 24 of
        # its batch 2 (E7.5), the H2-X in two versions.
        apart = {r.number: r.taken_apart for r in self.register if r.taken_apart}
        batch_two = {25, 47, 53, 69, 80, 98, 123, 132, 134, 135, 144, 151, 153, 157, 158, 163, 170, 171, 176, 177,
                     180, 199, 203, 210}
        batch_two_b = {216, 221, 222, 223, 224, 229, 232, 244, 256, 265, 271, 272, 276, 285, 290, 294, 297, 298, 299,
                       300, 302, 323, 326, 352}
        on_wheels = {13, 15, 20, 21, 34, 35, 40, 41, 42, 49, 60, 68, 82, 86, 92, 104, 105, 122, 179, 184, 200, 230, 247,
                     248, 251, 252, 253, 254, 255, 295, 307, 312, 313, 314, 315, 318, 335, 336, 337, 340, 341, 342, 349}
        # Since 2026-10-30 ("Yes, build it", EVIDENCE.md E11.7) batch 1 of types 3 and 4: 19 in type 3's file, 8 in
        # type 4's.
        types_34_batch_one = {3, 4, 5, 7, 10, 11, 12, 18, 22, 26, 28, 29, 31, 32, 33, 43, 51, 56, 58} | {
            1, 6, 8, 9, 14, 17, 45, 46}
        self.assertEqual(set(apart), {r.number for r in self.register if 1 in r.types} | {116, 117, 119, 120, 121}
                         | batch_two | batch_two_b | on_wheels         # since 2026-10-30 the 43 on wheels (E9.6)
                         | types_34_batch_one)
        self.assertEqual(apart[326], 'Walker Tienkung Vision; Walker Tienkung Embodied')
        self.assertEqual(apart[116], 'Fourier GR-1 one camera; Fourier GR-1 two cameras')
        self.assertEqual(apart[123], 'Unitree H2-X Core i7; Unitree H2-X Orin NX')
        self.assertEqual(apart[102], 'Unitree G1')
        self.assertEqual(apart[148], 'Booster K1 5 Ah; Booster K1 2 Ah')
        names = [n for v in apart.values() for n in v.split('; ')]
        self.assertEqual(len(names), len(set(names)))
        for name in names:
            self.assertFalse(set('(),') & set(name), msg=name)   # 04 reads the names between brackets, by commas

    def test_masses_as_stated(self):
        by_number = {r.number: r for r in self.register}
        self.assertEqual(by_number[102].mass, (34, 35, 36))          # G1: "35", +- 1 kg
        self.assertEqual(by_number[246].mass, (20, 25, 30))          # QUESTER1: "~20–30", the range
        self.assertEqual(by_number[325].mass, (72, 73, 74))          # Walker S2: "43, 73", the larger
        self.assertEqual(by_number[236].mass, (9.75, 19.5, 29.25))   # PrimeBOT Q1: none stated, E5.17 point 6

    def test_the_sheet_type_1_lists_every_robot_of_type_1_as_in_the_sheet_robots(self):
        # Matthias, 2026-10-29: "There has to be an excel sheet where all the robots of the different types are fully
        # listed. For now it has to be just all for type 1." Every robot whose types include 1, leading or second, in
        # the order of the register, every cell as in the sheet Robots, and 01's check of it finds nothing.
        listed = schemas.read_type_sheet(self.path, 1)
        self.assertEqual([r.number for r in listed], [r.number for r in self.register if 1 in r.types])
        self.assertEqual(len(listed), 31)
        self.assertIn(54, [r.number for r in listed])    # the Booster T1, 1.18 m: kept in type 1 (2026-10-29)
        self.assertIn(236, [r.number for r in listed])   # the PrimeBOT Q1, its mass an estimate (E5.17)
        self.assertEqual([f.text for f in schemas.check_type_sheet(self.register, listed, 1, Params(),
                                                                   set(APART[1]))], [])

    def test_the_sheet_type_2_lists_every_robot_of_type_2_batches_1_and_2_taken_apart(self):
        # Matthias, 2026-10-29, "Yes, as proposed" (type 2 as type 1, step 5: its robots all listed): every robot whose
        # types include 2, leading or second. Since 2026-10-30 ("Yes, build C"): its batch 1 taken apart, the others not
        # yet (Q4), the Xialan S0 without a mass; 02 draws type 2 for stock-and-flow from the workbook until all are.
        listed = schemas.read_type_sheet(self.path, 2)
        self.assertEqual([r.number for r in listed], [r.number for r in self.register if 2 in r.types])
        self.assertEqual(len(listed), 97)
        runs = {r.number: r.in_run for r in listed}
        # Batch 1 (E6.8) and, since 2026-10-30 ("Yes, build it"), batch 2's 24 robots (E7.5) and batch 2b's 24 (E8.3);
        # the 43 left all on wheels.
        batch_two = {25, 47, 53, 69, 80, 98, 123, 132, 134, 135, 144, 151, 153, 157, 158, 163, 170, 171, 176, 177,
                     180, 199, 203, 210}
        batch_two_b = {216, 221, 222, 223, 224, 229, 232, 244, 256, 265, 271, 272, 276, 285, 290, 294, 297, 298, 299,
                       300, 302, 323, 326, 352}
        # Since 2026-10-30 ("Yes, agree E9.6", "Yes, build it"): the 43 on wheels too; none with a mass left.
        on_wheels = {13, 15, 20, 21, 34, 35, 40, 41, 42, 49, 60, 68, 82, 86, 92, 104, 105, 122, 179, 184, 200, 230, 247,
                     248, 251, 252, 253, 254, 255, 295, 307, 312, 313, 314, 315, 318, 335, 336, 337, 340, 341, 342, 349}
        self.assertEqual(len(on_wheels), 43)
        self.assertEqual({n for n, run in runs.items() if run == schemas.IN_RUN_OWN},
                         {116, 117, 119, 120, 121} | batch_two | batch_two_b | on_wheels)
        self.assertEqual({n for n, run in runs.items() if run == schemas.IN_RUN_NOT}, {329})
        self.assertEqual({n for n, run in runs.items() if run == schemas.IN_RUN_LATER}, set())
        forms = {r.number: r.form for r in listed}
        self.assertEqual({forms[n] for n in on_wheels}, {'wheels'})
        self.assertEqual([f.text for f in schemas.check_type_sheet(self.register, listed, 2, Params(),
                                                                   set(APART[2]))], [])

    def test_the_sheets_type_3_and_type_4(self):
        # Matthias, 2026-10-30: "Yes, together", "Yes, build it" (EVIDENCE.md E10.1; since the entry of batch 1, E11.7
        # point 5): every robot whose types include 3 or 4, as in the sheet Robots; drawn with its own parts where they
        # are in a parts file of one of its types (types 1 to 4), not drawn where it has no mass ("If there is no mass
        # then it should not be taken into the evaluation"), else not yet taken apart; 01 finds nothing.
        for robot_type, count, own, without in ((3, 181, 72, 12), (4, 135, 33, 4)):
            listed = schemas.read_type_sheet(self.path, robot_type)
            self.assertEqual([r.number for r in listed], [r.number for r in self.register if robot_type in r.types])
            self.assertEqual(len(listed), count)
            runs = Counter(r.in_run for r in listed)
            self.assertEqual((runs[schemas.IN_RUN_OWN], runs[schemas.IN_RUN_NOT], runs[schemas.IN_RUN_LATER]),
                             (own, without, count - own - without))
            named = set().union(*APART.values())
            self.assertEqual([f.text for f in schemas.check_type_sheet(self.register, listed, robot_type, Params(),
                                                                       named)], [])
        self.assertEqual(schemas.type_sheet_numbers(self.path), [1, 2, 3, 4])
        self.assertEqual(schemas.types_taken_apart(self.path), [1, 2, 3, 4])

    def test_the_wheeled_robots_of_batch_2_on_wheels(self):
        # Matthias, 2026-10-30, "Yes, all six" (EVIDENCE.md E7.3, question 1): the Astribot S1 and the Moz1 stand on
        # wheels by their makers, the CUE7 by the press; to the wheeled batch, not taken apart yet.
        by_number = {r.number: r for r in self.register}
        for number in (40, 68, 200):
            self.assertEqual(by_number[number].form, 'wheels', msg=number)
            cells = dict(zip(schemas.REGISTER_HEADINGS, by_number[number].cells))
            self.assertIn('E7.3', cells['Decided in'], msg=number)
        runs = {r.number: r.in_run for r in schemas.read_type_sheet(self.path, 2)}
        self.assertEqual({runs[n] for n in (40, 68, 200)}, {schemas.IN_RUN_OWN})   # taken apart since E9.6

    def test_the_wheeled_robots_of_batch_2b_on_wheels(self):
        # Matthias, 2026-10-30, "Yes, all six" (EVIDENCE.md E8.2, question 1): Torobo, VB1-I and VB2 on wheels by their
        # makers; to the wheeled batch.
        by_number = {r.number: r for r in self.register}
        for number in (307, 312, 313):
            self.assertEqual(by_number[number].form, 'wheels', msg=number)
            cells = dict(zip(schemas.REGISTER_HEADINGS, by_number[number].cells))
            self.assertIn('E8.2', cells['Decided in'], msg=number)

    def test_the_sheet_type_1_says_what_the_run_does_with_each_robot(self):
        listed = {r.number: r.in_run for r in schemas.read_type_sheet(self.path, 1)}
        self.assertEqual(set(listed.values()), {'drawn: its own parts'})

    def test_it_writes_where_00_says(self):
        self.assertEqual(builder.output_path(Params()), absolute(Params().data.register))

    def test_what_it_writes_equals_the_register_on_disk(self):
        on_disk = absolute(Params().data.register)
        if not os.path.exists(on_disk):
            self.skipTest('the register is not on disk')
        written, present = load_workbook(self.path), load_workbook(on_disk)
        self.assertEqual(written.sheetnames, present.sheetnames)
        for name in written.sheetnames:
            rows = [list(r) for r in written[name].iter_rows(values_only=True)]
            self.assertEqual(rows, [list(r) for r in present[name].iter_rows(values_only=True)], msg=name)


if __name__ == '__main__':
    unittest.main()
